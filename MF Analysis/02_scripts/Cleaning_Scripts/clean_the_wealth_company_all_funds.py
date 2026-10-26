"""
Clean all downloaded monthly portfolio files for The Wealth Company Mutual Fund (The Wealth Company MF).

Process:
    - Reads all files in 01_raw_files/The_Wealth_Company_MF/<YYYY>/<MM>/
    - Parses files for the 8 target funds:
        1. The Wealth Company Multi Cap Fund
        2. The Wealth Company Mid Cap Fund
        3. The Wealth Company Large & Mid Cap Fund
        4. The Wealth Company Flexicap Fund
        5. The Wealth Company Multi Asset Allocation Fund
        6. The Wealth Company Ethical Fund
        7. The Wealth Company Balanced Advantage Fund
        8. The Wealth Company Small Cap Fund
    - Dynamic header detection for ISIN, Name of Instrument, Rating/Industry, and Quantity
    - Stops immediately upon encountering 'Sub Total', '(b) Unlisted', 'Total', etc.
    - Filters for valid listed Indian equities (ISIN starts with INE, length 12, quantity > 0)
    - Generates 03_clean_data/The_Wealth_Company_MF/The_Wealth_Company_MF_All_Funds_Cleaned.xlsx
"""

import sys
import os
import calendar
import warnings
from datetime import datetime
from pathlib import Path
import xlrd
import openpyxl
import pandas as pd

# Force UTF-8 output
try:
    sys.stdout.reconfigure(encoding="utf-8")
    sys.stderr.reconfigure(encoding="utf-8")
except (AttributeError, ValueError):
    pass

PROJECT_ROOT = Path(__file__).resolve().parents[2]
RAW_DIR = PROJECT_ROOT / "01_raw_files" / "The_Wealth_Company_MF"
CLEAN_DIR = PROJECT_ROOT / "03_clean_data" / "The_Wealth_Company_MF"
CLEAN_DIR.mkdir(parents=True, exist_ok=True)

OUTPUT_FILE = CLEAN_DIR / "The_Wealth_Company_MF_All_Funds_Cleaned.xlsx"

AMC_NAME = "The Wealth Company MF"

FUND_CODE_TO_NAME = {
    "Multi_Cap": "The Wealth Company Multi Cap Fund",
    "Mid_Cap": "The Wealth Company Mid Cap Fund",
    "Large_Mid_Cap": "The Wealth Company Large & Mid Cap Fund",
    "Flexicap": "The Wealth Company Flexicap Fund",
    "Multi_Asset": "The Wealth Company Multi Asset Allocation Fund",
    "Ethical": "The Wealth Company Ethical Fund",
    "Balanced_Advantage": "The Wealth Company Balanced Advantage Fund",
    "Small_Cap": "The Wealth Company Small Cap Fund",
}

STOP_TERMS = [
    "sub total",
    "subtotal",
    "(b) unlisted",
    "(b)unlisted",
    "total",
    "debt instruments",
    "money market",
    "treasury bill",
    "cash & cash equivalents",
    "net current assets",
    "real est inv trust",
]

STANDARD_COLUMNS = [
    "AMC",
    "Fund_Name",
    "Portfolio_Date",
    "Month",
    "Security_Name",
    "ISIN",
    "Industry_Rating",
    "Quantity",
]


def clean_xlsx_file(file_path, fund_name, portfolio_date, month_str):
    """Parses an openpyxl .xlsx file."""
    holdings = []
    wb = openpyxl.load_workbook(file_path, data_only=True)

    target_sheet = None
    for sname in wb.sheetnames:
        if "metadata" not in sname.lower():
            target_sheet = wb[sname]
            break
    if target_sheet is None:
        target_sheet = wb.active

    ws = target_sheet
    header_row = None
    col_map = {}

    for r in range(1, min(ws.max_row + 1, 35)):
        row_vals = [str(ws.cell(r, c).value or "").strip().lower() for c in range(1, min(ws.max_column + 1, 15))]
        if any("isin" in v for v in row_vals) and any("instrument" in v or "name" in v for v in row_vals):
            header_row = r
            for c, v in enumerate(row_vals, 1):
                if "isin" in v and "isin" not in col_map:
                    col_map["isin"] = c
                elif ("name" in v or "instrument" in v) and "name" not in col_map:
                    col_map["name"] = c
                elif ("rating" in v or "industry" in v) and "industry" not in col_map:
                    col_map["industry"] = c
                elif ("quantity" in v or "qty" in v) and "qty" not in col_map:
                    col_map["qty"] = c
            break

    if not header_row or "isin" not in col_map or "qty" not in col_map:
        return []

    if "name" not in col_map:
        col_map["name"] = 2
    if "industry" not in col_map:
        col_map["industry"] = 3

    in_eq = False
    for r in range(header_row + 1, ws.max_row + 1):
        c0 = str(ws.cell(r, 1).value or "").strip().lower()
        c1 = str(ws.cell(r, col_map["name"]).value or "").strip().lower()
        comb = f"{c0} {c1}"

        if "(a) listed" in comb or "listed / awaiting" in comb:
            in_eq = True
            continue

        if in_eq:
            if any(term in comb for term in STOP_TERMS):
                in_eq = False
                break

            isin_val = str(ws.cell(r, col_map["isin"]).value or "").strip().upper()
            name_val = str(ws.cell(r, col_map["name"]).value or "").strip()
            ind_val = str(ws.cell(r, col_map["industry"]).value or "").strip()
            raw_qty = ws.cell(r, col_map["qty"]).value

            try:
                qty_num = float(raw_qty)
                if qty_num <= 0:
                    continue
                qty_clean = int(qty_num) if qty_num.is_integer() else qty_num
            except (ValueError, TypeError):
                continue

            if isin_val.startswith("INE") and len(isin_val) == 12:
                holdings.append({
                    "AMC": AMC_NAME,
                    "Fund_Name": fund_name,
                    "Portfolio_Date": portfolio_date,
                    "Month": month_str,
                    "Security_Name": name_val,
                    "ISIN": isin_val,
                    "Industry_Rating": ind_val,
                    "Quantity": qty_clean,
                })

    return holdings


def clean_xls_file(file_path, fund_name, portfolio_date, month_str):
    """Parses an xlrd .xls file."""
    holdings = []
    wb = xlrd.open_workbook(file_path)

    sheet = None
    for sname in wb.sheet_names():
        if "metadata" not in sname.lower():
            sheet = wb.sheet_by_name(sname)
            break
    if sheet is None:
        sheet = wb.sheet_by_index(0)

    header_row = None
    col_map = {}

    for r in range(min(sheet.nrows, 35)):
        row_vals = [str(sheet.cell_value(r, c)).strip().lower() for c in range(sheet.ncols)]
        if any("isin" in v for v in row_vals) and any("instrument" in v or "name" in v for v in row_vals):
            header_row = r
            for c, v in enumerate(row_vals):
                if "isin" in v and "isin" not in col_map:
                    col_map["isin"] = c
                elif ("name" in v or "instrument" in v) and "name" not in col_map:
                    col_map["name"] = c
                elif ("rating" in v or "industry" in v) and "industry" not in col_map:
                    col_map["industry"] = c
                elif ("quantity" in v or "qty" in v) and "qty" not in col_map:
                    col_map["qty"] = c
            break

    if not header_row or "isin" not in col_map or "qty" not in col_map:
        return []

    if "name" not in col_map:
        col_map["name"] = 1
    if "industry" not in col_map:
        col_map["industry"] = 2

    in_eq = False
    for r in range(header_row + 1, sheet.nrows):
        c0 = str(sheet.cell_value(r, 0)).strip().lower()
        c1 = str(sheet.cell_value(r, col_map["name"])).strip().lower()
        comb = f"{c0} {c1}"

        if "(a) listed" in comb or "listed / awaiting" in comb:
            in_eq = True
            continue

        if in_eq:
            if any(term in comb for term in STOP_TERMS):
                in_eq = False
                break

            isin_val = str(sheet.cell_value(r, col_map["isin"])).strip().upper()
            name_val = str(sheet.cell_value(r, col_map["name"])).strip()
            ind_val = str(sheet.cell_value(r, col_map["industry"])).strip()
            raw_qty = sheet.cell_value(r, col_map["qty"])

            try:
                qty_num = float(raw_qty)
                if qty_num <= 0:
                    continue
                qty_clean = int(qty_num) if qty_num.is_integer() else qty_num
            except (ValueError, TypeError):
                continue

            if isin_val.startswith("INE") and len(isin_val) == 12:
                holdings.append({
                    "AMC": AMC_NAME,
                    "Fund_Name": fund_name,
                    "Portfolio_Date": portfolio_date,
                    "Month": month_str,
                    "Security_Name": name_val,
                    "ISIN": isin_val,
                    "Industry_Rating": ind_val,
                    "Quantity": qty_clean,
                })

    return holdings


def process_all_files():
    """Iterate over all raw files in The_Wealth_Company_MF, extract holdings for target funds, and write clean Excel."""
    print("=" * 80)
    print("Starting The Wealth Company MF Portfolio Cleaner")
    print(f"Reading from: {RAW_DIR}")
    print(f"Output target: {OUTPUT_FILE}")
    print("=" * 80)

    all_records = []
    file_count = 0

    for year_dir in sorted(RAW_DIR.iterdir()):
        if not year_dir.is_dir() or not year_dir.name.isdigit():
            continue
        year = int(year_dir.name)

        for month_dir in sorted(year_dir.iterdir()):
            if not month_dir.is_dir() or not month_dir.name.isdigit():
                continue
            month = int(month_dir.name)

            last_day = calendar.monthrange(year, month)[1]
            portfolio_date = datetime(year, month, last_day)
            month_str = f"{year}-{month:02d}-01"

            for file_path in sorted(list(month_dir.glob("*.xls")) + list(month_dir.glob("*.xlsx"))):
                fund_name = None
                for code, name in FUND_CODE_TO_NAME.items():
                    if file_path.name.startswith(f"{code}_"):
                        fund_name = name
                        break

                if not fund_name:
                    continue

                file_count += 1
                h = []
                if file_path.suffix.lower() == ".xlsx":
                    try:
                        h = clean_xlsx_file(file_path, fund_name, portfolio_date, month_str)
                    except Exception as e:
                        print(f"Error parsing .xlsx file {file_path.name}: {e}")
                elif file_path.suffix.lower() == ".xls":
                    try:
                        h = clean_xls_file(file_path, fund_name, portfolio_date, month_str)
                    except Exception as e:
                        print(f"Error parsing .xls file {file_path.name}: {e}")

                all_records.extend(h)
                print(f"[{year}-{month:02d}] {file_path.name}: {len(h)} holdings extracted.")

    print("=" * 80)
    print(f"Total processed files: {file_count}")
    print(f"Total holdings extracted: {len(all_records)}")

    if not all_records:
        print("ERROR: No holdings extracted! Aborting.")
        return False

    df = pd.DataFrame(all_records, columns=STANDARD_COLUMNS)

    # Validate dataframe
    print("\nSummary by Fund:")
    fund_summary = df.groupby("Fund_Name")["Quantity"].agg(["count", "nunique"]).reset_index()
    fund_summary.columns = ["Fund Name", "Total Records", "Unique Quantities"]
    print(fund_summary.to_string(index=False))

    print("\nSummary by Month:")
    month_summary = df.groupby("Month")["Fund_Name"].agg(["count", "nunique"]).reset_index()
    month_summary.columns = ["Month", "Total Records", "Funds Present"]
    print(month_summary.to_string(index=False))

    # Save to Excel
    print(f"\nWriting cleaned dataset to: {OUTPUT_FILE}")
    with pd.ExcelWriter(OUTPUT_FILE, engine="openpyxl") as writer:
        df.to_excel(writer, index=False, sheet_name="Cleaned_Holdings")

    print(f"Successfully generated: {OUTPUT_FILE} ({OUTPUT_FILE.stat().st_size / 1024:.1f} KB)")
    print("=" * 80)
    return True


if __name__ == "__main__":
    success = process_all_files()
    sys.exit(0 if success else 1)
