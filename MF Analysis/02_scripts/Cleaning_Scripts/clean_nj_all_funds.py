"""
Clean all downloaded monthly portfolio files for NJ Mutual Fund (NJ MF).

Process:
    - Reads all files in 01_raw_files/NJ_MF/<YYYY>/<MM>/
    - Supports both consolidated multi-scheme workbooks (Oct 2024 - Aug 2025)
      and individual scheme workbooks (Sep 2025 - Sep 2026)
    - Processes target schemes:
        1. NJ Flexi Cap Fund (NJFCP)
        2. NJ Balanced Advantage Fund (NJBAF)
        3. NJ Value Fund (NJVAF)
    - Dynamic header detection for Name, ISIN, Industry/Rating, and Quantity
    - Stops immediately upon encountering 'Sub Total', '(b)Unlisted', '(b) Unlisted', or 'Total'
    - Filters for valid listed Indian equities (ISIN starts with INE, length 12, quantity > 0)
    - Generates 03_clean_data/NJ_MF/NJ_MF_All_Funds_Cleaned.xlsx
"""

import sys
import os
import calendar
import warnings
from datetime import datetime
from pathlib import Path
import openpyxl
import pandas as pd

warnings.filterwarnings("ignore", category=UserWarning, module="openpyxl")

# Windows consoles default to cp1252, force UTF-8 output
try:
    sys.stdout.reconfigure(encoding="utf-8")
    sys.stderr.reconfigure(encoding="utf-8")
except (AttributeError, ValueError):
    pass

PROJECT_ROOT = Path(__file__).resolve().parents[2]
RAW_DIR = PROJECT_ROOT / "01_raw_files" / "NJ_MF"
CLEAN_DIR = PROJECT_ROOT / "03_clean_data" / "NJ_MF"
CLEAN_DIR.mkdir(parents=True, exist_ok=True)

OUTPUT_FILE = CLEAN_DIR / "NJ_MF_All_Funds_Cleaned.xlsx"

CODE_TO_FUND = {
    "NJFCP": "NJ Flexi Cap Fund",
    "NJBAF": "NJ Balanced Advantage Fund",
    "NJVAF": "NJ Value Fund",
}

STOP_TERMS = ["sub total", "(b)unlisted", "(b) unlisted", "total"]

STANDARD_COLUMNS = [
    "AMC",
    "Fund_Name",
    "Portfolio_Date",
    "Month",
    "Security_Name",
    "ISIN",
    "Industry_Rating",
    "Quantity",
]


def clean_sheet(ws, fund_name, portfolio_date, month_str):
    """Parses a single worksheet of an NJ portfolio workbook."""
    holdings = []

    # Locate header row and column mapping
    header_row = None
    col_map = {}
    for r in range(1, 16):
        row_vals = [str(ws.cell(row=r, column=c).value or "").strip() for c in range(1, ws.max_column + 1)]
        lower_vals = [v.lower() for v in row_vals]
        if any("isin" in v for v in lower_vals) and any("instrument" in v or "name" in v for v in lower_vals):
            header_row = r
            for c_idx, val in enumerate(lower_vals, 1):
                if "isin" in val and "isin" not in col_map:
                    col_map["isin"] = c_idx
                elif ("instrument" in val or "name" in val) and "yield" not in val and "name" not in col_map:
                    col_map["name"] = c_idx
                elif ("rating" in val or "industry" in val) and "rating" not in col_map:
                    col_map["rating"] = c_idx
                elif ("quantity" in val or "qty" in val) and "qty" not in col_map:
                    col_map["qty"] = c_idx
            break

    if not header_row or "isin" not in col_map or "name" not in col_map or "qty" not in col_map:
        return []

    for r in range(header_row + 1, ws.max_row + 1):
        name_val = str(ws.cell(row=r, column=col_map["name"]).value or "").strip()
        isin_val = str(ws.cell(row=r, column=col_map["isin"]).value or "").strip()

        name_lower = name_val.lower()
        isin_lower = isin_val.lower()

        # Immediate stop condition
        if any(term in name_lower or term in isin_lower for term in STOP_TERMS):
            break

        # Filter for listed equity
        if not (isin_val.startswith("INE") and len(isin_val) == 12):
            continue

        qty_val = ws.cell(row=r, column=col_map["qty"]).value
        try:
            qty = float(qty_val)
        except (ValueError, TypeError):
            continue

        if qty <= 0:
            continue

        rating_val = ""
        if "rating" in col_map:
            rating_val = str(ws.cell(row=r, column=col_map["rating"]).value or "").strip()

        holdings.append({
            "AMC": "NJ MF",
            "Fund_Name": fund_name,
            "Portfolio_Date": portfolio_date,
            "Month": month_str,
            "Security_Name": name_val,
            "ISIN": isin_val,
            "Industry_Rating": rating_val,
            "Quantity": qty,
        })

    return holdings


def clean_file(file_path):
    """Processes an NJ workbook (single-scheme or multi-scheme)."""
    file_holdings = []
    try:
        month = int(file_path.parent.name)
        year = int(file_path.parent.parent.name)

        last_day = calendar.monthrange(year, month)[1]
        portfolio_date = f"{year}-{month:02d}-{last_day:02d}"
        month_str = datetime.strptime(f"{year}-{month:02d}-01", "%Y-%m-%d").strftime("%b-%Y")

        wb = openpyxl.load_workbook(file_path, data_only=True)

        for sname in wb.sheetnames:
            target_fund = None
            sname_upper = sname.upper().strip()

            if sname_upper in CODE_TO_FUND:
                target_fund = CODE_TO_FUND[sname_upper]
            else:
                # Check top rows of sheet for scheme title
                ws_check = wb[sname]
                header_text = ""
                for r in range(1, 5):
                    for c in range(1, 6):
                        val = str(ws_check.cell(row=r, column=c).value or "")
                        if val:
                            header_text += " " + val
                h_lower = header_text.lower()
                if "flexi cap" in h_lower:
                    target_fund = "NJ Flexi Cap Fund"
                elif "balanced advantage" in h_lower:
                    target_fund = "NJ Balanced Advantage Fund"
                elif "value fund" in h_lower:
                    target_fund = "NJ Value Fund"

            if target_fund:
                sheet_holdings = clean_sheet(wb[sname], target_fund, portfolio_date, month_str)
                file_holdings.extend(sheet_holdings)

    except Exception as e:
        print(f"  [ERROR] {file_path.name}: {e}")

    return file_holdings


def main():
    print("=" * 80)
    print("NJ Mutual Fund (NJ MF) - Portfolio Cleaning Pipeline")
    print("=" * 80)

    files = sorted(list(RAW_DIR.rglob("*.xls*")))
    print(f"Found {len(files)} portfolio workbooks in {RAW_DIR}")

    all_holdings = []
    processed_count = 0
    funds_seen = set()

    for idx, f in enumerate(files, 1):
        file_holdings = clean_file(f)
        all_holdings.extend(file_holdings)
        processed_count += 1
        for h in file_holdings:
            funds_seen.add(h["Fund_Name"])

        if idx % 10 == 0 or idx == len(files):
            print(f"  Processed {idx}/{len(files)} workbooks | Total holdings so far: {len(all_holdings)}")

    print("-" * 80)
    print(f"Extracted {len(all_holdings)} listed equity holdings across {len(funds_seen)} funds.")

    if not all_holdings:
        print("[WARN] No holdings extracted! Please check raw files and column mappings.")
        return

    df = pd.DataFrame(all_holdings, columns=STANDARD_COLUMNS)

    # Sort deterministically
    df.sort_values(by=["Fund_Name", "Portfolio_Date", "Security_Name"], inplace=True)

    print(f"Writing cleaned dataset to: {OUTPUT_FILE} ...")
    with pd.ExcelWriter(OUTPUT_FILE, engine="openpyxl") as writer:
        df.to_excel(writer, index=False, sheet_name="Cleaned_Data")

    print(f"SUCCESS: Saved {len(df):,} holdings to {OUTPUT_FILE}")
    print("\nSummary by Fund:")
    fund_summary = df.groupby("Fund_Name").agg(
        Total_Holdings=("ISIN", "count"),
        Unique_ISINs=("ISIN", "nunique"),
        Months=("Month", "nunique")
    ).reset_index()
    for _, row in fund_summary.iterrows():
        print(f"  - {row['Fund_Name']:<30}: {row['Total_Holdings']:5d} holdings | {row['Unique_ISINs']:3d} unique stocks | {row['Months']:2d} months")

    print("=" * 80)


if __name__ == "__main__":
    main()
