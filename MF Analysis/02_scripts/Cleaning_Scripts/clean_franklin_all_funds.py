"""
Cleaning script for Franklin Templeton Mutual Fund.
Supports both multi-scheme master workbooks (.xlsx) and scheme-specific portfolio files (.xls / .xlsx).

AMC Name: Franklin Templeton MF
Output: 03_clean_data/FRANKLIN/Franklin_Templeton_MF_All_Funds_Cleaned.xlsx
"""

import sys
import re
from pathlib import Path
from datetime import datetime
import openpyxl
import xlrd
import pandas as pd
import numpy as np

# Ensure UTF-8 output
try:
    sys.stdout.reconfigure(encoding="utf-8")
    sys.stderr.reconfigure(encoding="utf-8")
except Exception:
    pass

PROJECT_ROOT = Path(__file__).resolve().parents[2]

RAW_FOLDER = PROJECT_ROOT / "01_raw_files" / "FRANKLIN"
OUTPUT_FOLDER = PROJECT_ROOT / "03_clean_data" / "FRANKLIN"
OUTPUT_FOLDER.mkdir(parents=True, exist_ok=True)

OUTPUT_FILE = OUTPUT_FOLDER / "Franklin_Templeton_MF_All_Funds_Cleaned.xlsx"

AMC_NAME = "Franklin Templeton MF"

# Recognized scheme sheet codes (Modern & Legacy)
TARGET_SHEETS = [
    # Modern scheme codes (from Jul 2025 onwards)
    "FIRF",    # Franklin India Retirement Fund
    "FICHF",   # Franklin India Conservative Hybrid Fund
    "FIMAAF",  # Franklin India Multi Asset Allocation Fund
    "FIESF",   # Franklin India Equity Savings Fund
    "FIBAF",   # Franklin India Balanced Advantage Fund
    "FIAHF",   # Franklin India Aggressive Hybrid Fund
    "TIVF",    # Templeton India Value Fund
    "FITF",    # Franklin India Technology Fund
    "FIOF",    # Franklin India Opportunities Fund
    "FIMCF",   # Franklin India Multi Cap Fund
    "FITX",    # Franklin India ELSS Tax Saver Fund
    "FIMICF",  # Franklin India Mid Cap Fund
    "FILCF",   # Franklin India Large Cap Fund
    "FIEF",    # Franklin India Flexi Cap Fund
    "FBIF",    # Franklin Build India Fund
    "FISCF",   # Franklin India Small Cap Fund
    "FILMCF",  # Franklin India Large & Mid Cap Fund
    "FIFEF",   # Franklin India Focused Equity Fund
    # Legacy scheme codes (Oct 2024 to Jun 2025)
    "FIPP",    # Franklin India Pension Plan -> Franklin India Retirement Fund
    "FIDHY",   # Franklin India Debt Hybrid Fund -> Franklin India Conservative Hybrid Fund
    "FIEHF",   # Franklin India Equity Hybrid Fund -> Franklin India Aggressive Hybrid Fund
    "FIPF",    # Franklin India Prima Fund -> Franklin India Mid Cap Fund
    "FIBF",    # Franklin India Bluechip Fund -> Franklin India Large Cap Fund
    "FIEAF",   # Franklin India Equity Advantage Fund -> Franklin India Large & Mid Cap Fund
]

LEGACY_NAME_MAPPING = {
    "Franklin India Prima Fund": "Franklin India Mid Cap Fund",
    "Franklin India Bluechip Fund": "Franklin India Large Cap Fund",
    "Franklin India Equity Advantage Fund": "Franklin India Large & Mid Cap Fund",
    "Franklin India Equity Hybrid Fund": "Franklin India Aggressive Hybrid Fund",
    "Franklin India Debt Hybrid Fund": "Franklin India Conservative Hybrid Fund",
    "Franklin India Pension Plan": "Franklin India Retirement Fund",
    "Franklin India Smaller Companies Fund": "Franklin India Small Cap Fund",
    "Franklin India Equity Fund": "Franklin India Flexi Cap Fund",
    "Franklin India Taxshield": "Franklin India ELSS Tax Saver Fund",
}

STANDARD_COLUMNS = [
    "AMC",
    "Fund_Name",
    "Portfolio_Date",
    "Month",
    "Security_Name",
    "ISIN",
    "Industry_Rating",
    "Quantity",
]

STOP_MARKERS = [
    "sub total",
    "subtotal",
    "sub-total",
    "total",
    "grand total",
    "(b) unlisted",
    "unlisted",
    "(b) units of real estate investment trusts",
    "(b) units",
    "reit",
    "debt instruments",
    "money market instruments",
    "treasury bill",
    "mutual fund units",
    "foreign equity securities",
    "foreign mutual fund units",
]


def is_valid_isin(val):
    if not val or not isinstance(val, str):
        return False
    val = val.strip()
    return len(val) == 12 and (val.startswith("INE") or val.startswith("INF") or val.startswith("IN9"))


def clean_sec_name(s):
    if not s or pd.isna(s):
        return ""
    # Remove trailing symbols like #, $, ^, ~, !, *, @, %, etc.
    return re.sub(r"[\s#\$\^~\!\*@%]+$", "", str(s)).strip()


def clean_fund_name(raw_name):
    if not raw_name or pd.isna(raw_name):
        return ""
    # Remove former name in parentheses e.g. "(formerly known as ...)" with optional inner whitespace
    clean = re.sub(r"\s*\(\s*(?:formerly|Formerly)\s+known\s+as\s+[^\)]+\)", "", str(raw_name), flags=re.IGNORECASE)
    # Remove trailing footnote markers
    clean = re.sub(r"[\s\^#\*~]+$", "", clean).strip()
    
    # Map legacy scheme names to standard names
    for old_k, std_name in LEGACY_NAME_MAPPING.items():
        if clean.lower() == old_k.lower() or clean.lower().startswith(old_k.lower()):
            return std_name
            
    return clean


def extract_portfolio_date(text, filename=""):
    # Master file filename pattern: Monthly-Portfolio-ISIN-31-Oct-2024.xlsx
    m_master = re.search(r"Monthly-Portfolio-ISIN-\d{1,2}-([a-zA-Z]+)-(\d{4})", str(filename))
    if m_master:
        mon, yr = m_master.group(1), m_master.group(2)
        dt = pd.to_datetime(f"{mon} 1, {yr}", errors="coerce")
        if pd.notna(dt):
            return dt.replace(day=1)

    m = re.search(r"([a-zA-Z]+ \d{1,2}, \d{4})", str(text))
    if m:
        dt = pd.to_datetime(m.group(1), errors="coerce")
        if pd.notna(dt):
            return dt.replace(day=1)

    m2 = re.search(r"(\d{1,2}[ -][a-zA-Z]+[ -]\d{4})", str(text))
    if m2:
        dt = pd.to_datetime(m2.group(1), errors="coerce")
        if pd.notna(dt):
            return dt.replace(day=1)

    # Fallback to filename (e.g. Aug_26 or Aug-2026)
    m3 = re.search(r"([a-zA-Z]{3})[_-](\d{2,4})", str(filename))
    if m3:
        mon, yr = m3.group(1), m3.group(2)
        if len(yr) == 2:
            yr = "20" + yr
        dt = pd.to_datetime(f"{mon} 1, {yr}", errors="coerce")
        if pd.notna(dt):
            return dt.replace(day=1)

    return None


def clean_openpyxl_workbook(workbook_path: Path) -> list:
    print(f"\nProcessing .xlsx workbook: {workbook_path.name}")
    wb = openpyxl.load_workbook(workbook_path, data_only=True)
    all_rows = []

    for sheet_name in wb.sheetnames:
        if sheet_name not in TARGET_SHEETS:
            continue

        ws = wb[sheet_name]
        raw_fund_name = str(ws.cell(row=1, column=1).value or "").strip()

        fund_name = clean_fund_name(raw_fund_name)
        date_text = str(ws.cell(row=3, column=1).value or "").strip()
        portfolio_dt = extract_portfolio_date(date_text, workbook_path.name)

        if portfolio_dt is None:
            portfolio_dt = pd.to_datetime("2026-08-01")

        p_date_str = portfolio_dt.strftime("%Y-%m-%d")
        month_str = portfolio_dt.strftime("%b-%Y")

        sheet_rows = 0
        for r in range(7, ws.max_row + 1):
            c1 = str(ws.cell(row=r, column=1).value or "").strip()
            c2 = str(ws.cell(row=r, column=2).value or "").strip()
            c3 = str(ws.cell(row=r, column=3).value or "").strip()
            val4 = ws.cell(row=r, column=4).value

            c1_low = c1.lower()
            c2_low = c2.lower()
            combined_low = f"{c1_low} {c2_low}".strip()

            # Unified stop boundary check
            if (
                any(c1_low.startswith(marker) for marker in STOP_MARKERS)
                or any(c2_low.startswith(marker) for marker in STOP_MARKERS)
                or any(combined_low.startswith(marker) for marker in STOP_MARKERS)
                or "reit" in combined_low
            ):
                break

            if is_valid_isin(c1):
                try:
                    qty = int(float(val4))
                except (ValueError, TypeError):
                    continue

                if qty > 0:
                    sec_name = clean_sec_name(c2)
                    all_rows.append({
                        "AMC": AMC_NAME,
                        "Fund_Name": fund_name,
                        "Portfolio_Date": p_date_str,
                        "Month": month_str,
                        "Security_Name": sec_name,
                        "ISIN": c1,
                        "Industry_Rating": c3,
                        "Quantity": qty,
                    })
                    sheet_rows += 1

        print(f"   ✓ [{sheet_name}] {fund_name:<45} | Rows: {sheet_rows}")

    wb.close()
    return all_rows


def clean_xlrd_workbook(workbook_path: Path) -> list:
    print(f"\nProcessing .xls workbook: {workbook_path.name}")
    wb = xlrd.open_workbook(workbook_path)
    all_rows = []

    for sheet_name in wb.sheet_names():
        ws = wb.sheet_by_name(sheet_name)
        if ws.nrows < 6:
            continue

        raw_fund_name = str(ws.cell_value(0, 0) or "").strip()
        is_target = (sheet_name in TARGET_SHEETS) or (wb.nsheets == 1)

        if not is_target:
            continue

        fund_name = clean_fund_name(raw_fund_name)
        date_text = str(ws.cell_value(2, 0) or "").strip()
        portfolio_dt = extract_portfolio_date(date_text, workbook_path.name)

        if portfolio_dt is None:
            portfolio_dt = pd.to_datetime("2026-08-01")

        p_date_str = portfolio_dt.strftime("%Y-%m-%d")
        month_str = portfolio_dt.strftime("%b-%Y")

        sheet_rows = 0
        for r in range(6, ws.nrows):
            c1 = str(ws.cell_value(r, 0) or "").strip()
            c2 = str(ws.cell_value(r, 1) or "").strip()
            c3 = str(ws.cell_value(r, 2) or "").strip()
            val4 = ws.cell_value(r, 3)

            c1_low = c1.lower()
            c2_low = c2.lower()
            combined_low = f"{c1_low} {c2_low}".strip()

            # Unified stop boundary check
            if (
                any(c1_low.startswith(marker) for marker in STOP_MARKERS)
                or any(c2_low.startswith(marker) for marker in STOP_MARKERS)
                or any(combined_low.startswith(marker) for marker in STOP_MARKERS)
                or "reit" in combined_low
            ):
                break

            if is_valid_isin(c1):
                try:
                    qty = int(float(val4))
                except (ValueError, TypeError):
                    continue

                if qty > 0:
                    sec_name = clean_sec_name(c2)
                    all_rows.append({
                        "AMC": AMC_NAME,
                        "Fund_Name": fund_name,
                        "Portfolio_Date": p_date_str,
                        "Month": month_str,
                        "Security_Name": sec_name,
                        "ISIN": c1,
                        "Industry_Rating": c3,
                        "Quantity": qty,
                    })
                    sheet_rows += 1

        print(f"   ✓ [{sheet_name}] {fund_name:<45} | Rows: {sheet_rows}")

    return all_rows


def main():
    print("=" * 80)
    print(f"Cleaning {AMC_NAME} Monthly Disclosures")
    print("=" * 80)

    workbooks = sorted(list(RAW_FOLDER.glob("*.xlsx")) + list(RAW_FOLDER.glob("*.xls")))
    workbooks = [w for w in workbooks if not w.name.startswith("~$")]

    if not workbooks:
        print(f"❌ No workbooks found in {RAW_FOLDER}")
        return

    print(f"Found {len(workbooks)} workbook(s) in {RAW_FOLDER.name}")

    all_data = []
    for wb_path in workbooks:
        if wb_path.suffix.lower() == ".xlsx":
            rows = clean_openpyxl_workbook(wb_path)
        elif wb_path.suffix.lower() == ".xls":
            rows = clean_xlrd_workbook(wb_path)
        else:
            continue
        all_data.extend(rows)

    if not all_data:
        print("❌ No cleaned records produced.")
        return

    df = pd.DataFrame(all_data)

    # Clean data types & formatting
    df["Portfolio_Date"] = pd.to_datetime(df["Portfolio_Date"], errors="coerce")
    df = df[df["Portfolio_Date"].notna()]
    df = df[df["Portfolio_Date"] >= "2024-10-01"].copy()

    # Drop exact duplicate holding rows across files
    df = df.drop_duplicates(subset=["Fund_Name", "Portfolio_Date", "ISIN", "Security_Name"]).reset_index(drop=True)
    df = df.sort_values(by=["Portfolio_Date", "Fund_Name", "Security_Name"]).reset_index(drop=True)

    # Format Portfolio_Date back to YYYY-MM-DD string
    df["Portfolio_Date"] = df["Portfolio_Date"].dt.strftime("%Y-%m-%d")

    # Enforce standard columns
    df = df[STANDARD_COLUMNS]

    # Write output file
    if OUTPUT_FILE.exists():
        try:
            OUTPUT_FILE.unlink()
        except Exception:
            pass

    df.to_excel(OUTPUT_FILE, index=False)

    print("\n" + "=" * 80)
    print("🎉 Cleaning Finished Successfully!")
    print("=" * 80)
    print(f"Output File  : {OUTPUT_FILE}")
    print(f"Total Rows   : {len(df):,}")
    print(f"Active Funds : {df['Fund_Name'].nunique()}")
    for fund in sorted(df["Fund_Name"].unique()):
        cnt = len(df[df["Fund_Name"] == fund])
        print(f"   - {fund:<45}: {cnt:>4} rows")
    print(f"Total Stocks : {df['ISIN'].nunique()}")
    print("=" * 80)


if __name__ == "__main__":
    main()
