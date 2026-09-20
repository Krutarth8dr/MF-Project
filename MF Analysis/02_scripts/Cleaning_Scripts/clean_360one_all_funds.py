"""
Cleaning script for 360 ONE Mutual Fund.
AMC Name: 360 One Wam MF
Historical Scope: October 2024 through August 2026 (23 months).

Target Funds:
1. 360 ONE Focused Fund
2. 360 ONE Quant Fund
3. 360 ONE ELSS Tax Saver Nifty 50 Index Fund
4. 360 ONE Flexicap Fund
5. 360 ONE Balanced Hybrid Fund
6. 360 ONE Multi Asset Allocation Fund

Output: 03_clean_data/360_ONE/360_One_Wam_MF_All_Funds_Cleaned.xlsx
"""

import io
import re
import sys
from pathlib import Path
import openpyxl
import pandas as pd

# Windows consoles default to cp1252, force UTF-8 output
try:
    sys.stdout.reconfigure(encoding="utf-8")
    sys.stderr.reconfigure(encoding="utf-8")
except Exception:
    pass

PROJECT_ROOT = Path(__file__).resolve().parents[2]
RAW_FOLDER = PROJECT_ROOT / "01_raw_files" / "360_ONE"
OUTPUT_FOLDER = PROJECT_ROOT / "03_clean_data" / "360_ONE"
OUTPUT_FOLDER.mkdir(parents=True, exist_ok=True)

OUTPUT_FILE = OUTPUT_FOLDER / "360_One_Wam_MF_All_Funds_Cleaned.xlsx"
AMC_NAME = "360 One Wam MF"

TARGET_MAP = {
    "FOCUSED": "360 ONE Focused Fund",
    "QUANT": "360 ONE Quant Fund",
    "ELSS": "360 ONE ELSS Tax Saver Nifty 50 Index Fund",
    "FLEXICAP": "360 ONE Flexicap Fund",
    "HYBRID": "360 ONE Balanced Hybrid Fund",
    "MULTI_ASSET": "360 ONE Multi Asset Allocation Fund",
}

STOP_MARKERS = [
    "sub total",
    "subtotal",
    "sub-total",
    "total",
    "grand total",
    "(b) unlisted",
    "unlisted",
    "debt instruments",
    "bonds",
    "treasury",
    "money market",
    "mutual fund",
    "fixed deposit",
    "cash",
]

MONTH_ORDER = {
    "jan": 1, "feb": 2, "mar": 3, "apr": 4, "may": 5, "jun": 6,
    "june": 6, "jul": 7, "july": 7, "aug": 8, "sep": 9, "sept": 9,
    "oct": 10, "nov": 11, "dec": 12
}

STANDARD_COLUMNS = [
    "AMC",
    "Fund_Name",
    "Portfolio_Date",
    "Month",
    "Security_Name",
    "ISIN",
    "Industry_Rating",
    "Quantity",
]


def identify_fund(sheet_name, row1_text):
    s_clean = sheet_name.strip().lower()
    r_clean = row1_text.strip().lower()

    if "multi asset" in s_clean or "yy20" in s_clean or "multi asset" in r_clean:
        return TARGET_MAP["MULTI_ASSET"]
    if "elss" in s_clean or "yy12" in s_clean or "elss" in r_clean:
        return TARGET_MAP["ELSS"]
    if "quant" in s_clean or "yy11" in s_clean or "quant" in r_clean:
        return TARGET_MAP["QUANT"]
    if "flexi" in s_clean or "yy13" in s_clean or "flexi" in r_clean:
        return TARGET_MAP["FLEXICAP"]
    if "hybrid" in s_clean or "yy14" in s_clean or "balanced hybrid" in r_clean:
        return TARGET_MAP["HYBRID"]
    if "focused" in s_clean or "yy0a" in s_clean or "focused" in r_clean:
        return TARGET_MAP["FOCUSED"]

    return None


def is_valid_isin(val):
    if not val or not isinstance(val, str):
        return False
    val = val.strip()
    return len(val) == 12 and (val.startswith("INE") or val.startswith("INF") or val.startswith("IN9"))


def clean_sec_name(s):
    if not s or pd.isna(s):
        return ""
    # Remove trailing symbols (*, @, #, $, ^, ~, %, !)
    cleaned = re.sub(r"[\s#\$\^~\!\*@%]+$", "", str(s)).strip()
    # Normalize excessive internal whitespace
    cleaned = re.sub(r"\s+", " ", cleaned)
    return cleaned


def extract_portfolio_date(row3_text, filename):
    # Parse date from row 3 text e.g. "Monthly Portfolio Statement as on August 31,2026"
    m_dt = re.search(r"as on\s+([A-Za-z]+)\s+(\d{1,2}),?\s*(\d{4})", str(row3_text), re.IGNORECASE)
    if m_dt:
        mon_str = m_dt.group(1).lower()[:3]
        yr_str = m_dt.group(3)
        m_num = MONTH_ORDER.get(mon_str, 0)
        if m_num > 0:
            return f"{yr_str}-{m_num:02d}-01"

    # Fallback to filename
    m_fn = re.search(r"(Jan|Feb|Mar|Apr|May|Jun|Jul|Aug|Sep|Sept|Oct|Nov|Dec)[a-z]*_?(\d{4})", str(filename), re.IGNORECASE)
    if m_fn:
        mon_str = m_fn.group(1).lower()[:3]
        yr_str = m_fn.group(2)
        m_num = MONTH_ORDER.get(mon_str, 0)
        if m_num > 0:
            return f"{yr_str}-{m_num:02d}-01"

    return None


def process_file(file_path):
    rows = []
    skipped_reits = []

    with open(file_path, "rb") as fp:
        wb = openpyxl.load_workbook(io.BytesIO(fp.read()), data_only=True)

    for s in wb.sheetnames:
        ws = wb[s]
        r1_vals = [str(ws.cell(row=1, column=c).value or "").strip() for c in range(1, 6)]
        r1_text = " ".join(v for v in r1_vals if v)

        fund_name = identify_fund(s, r1_text)
        if not fund_name:
            continue

        r3_text = str(ws.cell(row=3, column=2).value or "").strip()
        pdate = extract_portfolio_date(r3_text, file_path.name)
        if not pdate or pdate < "2024-10-01":
            continue

        month_str = pd.to_datetime(pdate).strftime("%b-%Y")

        in_equity = False
        for r in range(4, ws.max_row + 1):
            c2 = str(ws.cell(row=r, column=2).value or "").strip()
            c3 = str(ws.cell(row=r, column=3).value or "").strip()
            c4 = str(ws.cell(row=r, column=4).value or "").strip()
            c5 = ws.cell(row=r, column=5).value

            c2_low = c2.lower()
            if "equity & equity related" in c2_low:
                in_equity = True
                continue

            if in_equity:
                if any(c2_low.startswith(marker) for marker in STOP_MARKERS):
                    break

                if not is_valid_isin(c3):
                    continue

                sec_name = clean_sec_name(c2)
                sec_low = sec_name.lower()

                # Skip REITs / InvITs (both by 9-10th ISIN digits and keywords)
                is_reit_invit = False
                if len(c3) == 12 and c3[8:10] in ["23", "25"]:
                    is_reit_invit = True
                if any(k in sec_low for k in ["reit", "invit", "real estate investment trust", "infrastructure investment trust", "infra trust", "office parks"]):
                    is_reit_invit = True

                if is_reit_invit:
                    skipped_reits.append((fund_name, pdate, sec_name, c3))
                    continue

                # Parse Quantity
                try:
                    qty = int(float(c5)) if c5 is not None and str(c5).strip() != "" else 0
                except (ValueError, TypeError):
                    qty = 0

                if qty > 0:
                    rows.append({
                        "AMC": AMC_NAME,
                        "Fund_Name": fund_name,
                        "Portfolio_Date": pdate,
                        "Month": month_str,
                        "Security_Name": sec_name,
                        "ISIN": c3,
                        "Industry_Rating": c4,
                        "Quantity": qty,
                    })

    return rows, skipped_reits


def main():
    print("=" * 80)
    print(f"Starting Cleaning for {AMC_NAME}")
    print("=" * 80)

    raw_files = sorted(RAW_FOLDER.glob("*.*"))
    print(f"Found {len(raw_files)} files in {RAW_FOLDER}")

    all_rows = []
    total_skipped_reits = []

    for idx, f in enumerate(raw_files, 1):
        print(f"[{idx:02d}/{len(raw_files)}] Processing: {f.name}...")
        f_rows, f_skipped = process_file(f)
        all_rows.extend(f_rows)
        total_skipped_reits.extend(f_skipped)
        print(f"    Extracted {len(f_rows)} rows (Skipped {len(f_skipped)} REITs/InvITs)")

    if not all_rows:
        print("❌ ERROR: No rows extracted!")
        sys.exit(1)

    df = pd.DataFrame(all_rows)

    # Enforce standard columns & format
    df = df[STANDARD_COLUMNS].copy()
    df["Portfolio_Date"] = pd.to_datetime(df["Portfolio_Date"]).dt.date
    df["Quantity"] = df["Quantity"].astype(int)

    # Sort deterministically
    df = df.sort_values(["Fund_Name", "Portfolio_Date", "Security_Name"]).reset_index(drop=True)

    # Write output
    print("\n" + "=" * 80)
    print(f"Writing cleaned data to: {OUTPUT_FILE}")
    df.to_excel(OUTPUT_FILE, index=False, engine="openpyxl")
    print(f"File successfully saved! Size: {OUTPUT_FILE.stat().st_size:,} bytes")
    print("=" * 80)

    # Quality Checks
    print("\n" + "=" * 80)
    print("DATA QUALITY AUDIT:")
    print("=" * 80)
    print(f"Total Rows: {len(df):,}")
    print(f"Distinct Funds ({df['Fund_Name'].nunique()}):")
    for fn, count in df["Fund_Name"].value_counts().items():
        print(f"  - {fn:<45}: {count:,} rows")

    dates = sorted(df["Portfolio_Date"].unique())
    print(f"\nDistinct Dates ({len(dates)}): {dates[0]} to {dates[-1]}")
    print(f"Months: {[pd.to_datetime(d).strftime('%b-%Y') for d in dates]}")

    null_counts = df.isnull().sum()
    print(f"\nNull Counts across columns:\n{null_counts}")

    dup_count = df.duplicated(subset=["AMC", "Fund_Name", "Portfolio_Date", "ISIN"]).sum()
    print(f"\nDuplicates on (Fund, Date, ISIN): {dup_count}")

    print(f"Total REITs/InvITs filtered out: {len(total_skipped_reits):,}")
    print("=" * 80)


if __name__ == "__main__":
    main()
