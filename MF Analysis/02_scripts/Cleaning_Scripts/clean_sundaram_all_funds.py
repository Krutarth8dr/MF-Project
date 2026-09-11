import os
import re
import sys
import calendar
from datetime import datetime
from pathlib import Path
import openpyxl
import pandas as pd

# ==============================================================================
# SETTINGS & PATHS
# ==============================================================================

PROJECT_ROOT = Path(__file__).resolve().parents[2]
RAW_FOLDER = PROJECT_ROOT / "01_raw_files" / "SUNDARAM"
OUTPUT_FOLDER = PROJECT_ROOT / "03_clean_data" / "SUNDARAM"
OUTPUT_FOLDER.mkdir(parents=True, exist_ok=True)
OUTPUT_FILE = OUTPUT_FOLDER / "Sundaram_All_Funds_Cleaned.xlsx"

AMC_NAME = "Sundaram MF"
START_DATE = pd.Timestamp(year=2024, month=10, day=1)
END_DATE = pd.Timestamp(year=2026, month=8, day=31)

STANDARD_COLUMNS = [
    "AMC",
    "Fund_Name",
    "Portfolio_Date",
    "Month",
    "Security_Name",
    "ISIN",
    "Industry_Rating",
    "Quantity",
]

# Canonical 22 Target Funds (mapped from August 2026 Index)
TARGET_FUNDS_CONFIG = {
    "CAPEXG": "Sundaram Infrastructure Advantage Fund",
    "MIDCAP": "Sundaram Mid Cap Fund",
    "MULTIP": "Sundaram Large And Mid Cap Fund",
    "SLTADV3": "Sundaram Long Term Advantage Fund Series III",
    "SLTADV4": "Sundaram Long Term Advantage Fund Series IV",
    "SMILE": "Sundaram Small Cap Fund",
    "SPAHF": "Sundaram Aggressive Hybrid Fund",
    "SPBAF": "Sundaram Dynamic Asset Allocation Fund",
    "SPDYF": "Sundaram Dividend Yield Fund",
    "SPESF": "Sundaram Equity Savings Fund",
    "SPFOCUS": "Sundaram Focused Fund",
    "SPMUCF": "Sundaram Multi Cap Fund",
    "SPTAX": "Sundaram ELSS Tax Saver Fund",
    "SRURAL": "Sundaram Consumption Fund",
    "SSFUND": "Sundaram Services Fund",
    "STAX": "Sundaram Value Fund",
    "SUNBCF": "Sundaram Large Cap Fund",
    "SUNFCF": "Sundaram Flexi Cap Fund",
    "SUNFOP": "Sundaram Financial Services Opportunities Fund",
    "SUNMAF": "Sundaram Multi Asset Allocation Fund",
    "SUNCYF": "Sundaram Business Cycle Fund",
    "SUNMFF": "Sundaram Multi-Factor Fund",
}

MONTH_MAP = {
    "january": 1, "jan": 1,
    "february": 2, "feb": 2,
    "march": 3, "mar": 3,
    "april": 4, "apr": 4,
    "may": 5,
    "june": 6, "jun": 6,
    "july": 7, "jul": 7,
    "august": 8, "aug": 8,
    "september": 9, "sep": 9, "sept": 9,
    "october": 10, "oct": 10,
    "november": 11, "nov": 11,
    "december": 12, "dec": 12,
}


def clean_text(val):
    if val is None:
        return ""
    text = str(val).strip()
    return re.sub(r"\s+", " ", text)


def parse_portfolio_date_from_file(ws, file_path):
    for r in range(1, 10):
        for c in range(1, 8):
            val = ws.cell(r, c).value
            if val and any(k in str(val).lower() for k in ["month ended", "as on", "statement for"]):
                text = str(val)
                m = re.search(r"(\d{1,2})(?:st|nd|rd|th)?\s*[-_ ]?\s*([a-zA-Z]+)\s*[-_ ]?\s*(\d{4})", text, re.I)
                if m:
                    d_str, m_str, y_str = m.groups()
                    mo_num = MONTH_MAP.get(m_str.lower().strip())
                    if mo_num:
                        return datetime(int(y_str), mo_num, int(d_str))

                m2 = re.search(r"([a-zA-Z]+)\s*(\d{1,2}),?\s*(\d{4})", text, re.I)
                if m2:
                    m_str, d_str, y_str = m2.groups()
                    mo_num = MONTH_MAP.get(m_str.lower().strip())
                    if mo_num:
                        return datetime(int(y_str), mo_num, int(d_str))

    parts = file_path.parts
    try:
        y = int(parts[-3])
        m = int(parts[-2])
        last_d = calendar.monthrange(y, m)[1]
        return datetime(y, m, last_d)
    except Exception:
        return datetime(2026, 8, 31)


def clean_sundaram_data():
    print("=" * 75)
    print("SUNDARAM MUTUAL FUND - CLEANING PIPELINE")
    print(f"Raw Folder:   {RAW_FOLDER}")
    print(f"Output File:  {OUTPUT_FILE}")
    print(f"Target Date Range: {START_DATE.strftime('%b %Y')} to {END_DATE.strftime('%b %Y')}")
    print("=" * 75)

    excel_files = sorted(list(RAW_FOLDER.rglob("*.xlsx")) + list(RAW_FOLDER.rglob("*.xls")))
    if not excel_files:
        print(f"ERROR: No monthly files found in {RAW_FOLDER}")
        sys.exit(1)

    print(f"\n[Step 1/3] Found {len(excel_files)} monthly consolidated workbooks to process.")

    all_rows = []

    for file_idx, file_path in enumerate(excel_files, 1):
        try:
            wb = openpyxl.load_workbook(file_path, data_only=True)
            sheetnames = wb.sheetnames
        except Exception as e:
            print(f"  [ERROR] Failed to load {file_path.name}: {e}")
            continue

        ws_sample = wb[sheetnames[1]] if len(sheetnames) > 1 else wb[sheetnames[0]]
        portfolio_dt = parse_portfolio_date_from_file(ws_sample, file_path)
        portfolio_ts = pd.Timestamp(portfolio_dt)

        if not (START_DATE <= portfolio_ts <= END_DATE):
            continue

        month_str = portfolio_dt.strftime("%b-%Y")

        for sheet_code, canonical_name in TARGET_FUNDS_CONFIG.items():
            matched_sheet = None
            for s in sheetnames:
                if s.strip().upper() == sheet_code.upper():
                    matched_sheet = s
                    break

            if not matched_sheet:
                continue

            ws = wb[matched_sheet]

            header_row = None
            col_name = None
            col_isin = None
            col_ind = None
            col_qty = None

            for r in range(1, 20):
                row_vals = [ws.cell(r, c).value for c in range(1, 15)]
                for c_idx, v in enumerate(row_vals, 1):
                    if v is None:
                        continue
                    v_lower = str(v).strip().lower()
                    if "name of the instrument" in v_lower:
                        header_row = r
                        col_name = c_idx
                    elif "isin" in v_lower:
                        col_isin = c_idx
                    elif any(k in v_lower for k in ["rating / industry", "industry / rating", "industry", "rating"]):
                        col_ind = c_idx
                    elif "quantity" in v_lower or "qty" in v_lower:
                        col_qty = c_idx

                if header_row and col_name and col_isin and col_qty:
                    break

            if not (header_row and col_name and col_isin and col_qty):
                print(f"  [WARN] Header not detected in {file_path.name} -> sheet {matched_sheet}")
                continue

            for r in range(header_row + 1, ws.max_row + 1):
                name_v = ws.cell(r, col_name).value
                name_str = clean_text(name_v)

                # Stop scanning immediately when Sub Total is detected
                if name_str.lower() in ["sub total", "sub-total", "subtotal", "total"]:
                    break

                isin_v = ws.cell(r, col_isin).value
                isin_str = clean_text(isin_v)

                if not isin_str.startswith("INE"):
                    continue

                ind_str = clean_text(ws.cell(r, col_ind).value) if col_ind else ""

                qty_v = ws.cell(r, col_qty).value
                try:
                    if qty_v is None or str(qty_v).strip() == "":
                        continue
                    qty = float(str(qty_v).replace(",", "").strip())
                    if qty <= 0:
                        continue
                except (ValueError, TypeError):
                    continue

                all_rows.append({
                    "AMC": AMC_NAME,
                    "Fund_Name": canonical_name,
                    "Portfolio_Date": portfolio_ts.strftime("%Y-%m-%d"),
                    "Month": month_str,
                    "Security_Name": name_str,
                    "ISIN": isin_str,
                    "Industry_Rating": ind_str,
                    "Quantity": int(round(qty)),
                })

    print("\n[Step 2/3] Compiling cleaned dataset...")
    if not all_rows:
        print("ERROR: No valid equity holdings extracted!")
        sys.exit(1)

    df = pd.DataFrame(all_rows, columns=STANDARD_COLUMNS)

    # Convert types
    df["Portfolio_Date"] = pd.to_datetime(df["Portfolio_Date"])
    df["Quantity"] = df["Quantity"].astype("int64")

    # Sort dataset cleanly
    df = df.sort_values(["Portfolio_Date", "Fund_Name", "Security_Name"]).reset_index(drop=True)

    print(f"Total rows extracted: {len(df):,}")
    print(f"Unique funds:         {df['Fund_Name'].nunique()}")
    print(f"Unique ISINs:         {df['ISIN'].nunique()}")
    print(f"Unique months:        {df['Month'].nunique()}")
    print(f"Date range:           {df['Portfolio_Date'].min().strftime('%d-%b-%Y')} to {df['Portfolio_Date'].max().strftime('%d-%b-%Y')}")

    print("\n[Step 3/3] Saving cleaned data to Excel...")
    df.to_excel(OUTPUT_FILE, index=False)
    print(f"Saved successfully: {OUTPUT_FILE} ({OUTPUT_FILE.stat().st_size:,} bytes)")
    print("=" * 75)


if __name__ == "__main__":
    clean_sundaram_data()
