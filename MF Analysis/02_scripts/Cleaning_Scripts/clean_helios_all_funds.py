import os
import re
import sys
import calendar
from datetime import datetime
from pathlib import Path
import openpyxl
import pandas as pd

# ==============================================================================
# SETTINGS & PATHS
# ==============================================================================

PROJECT_ROOT = Path(__file__).resolve().parents[2]
RAW_FOLDER = PROJECT_ROOT / "01_raw_files" / "HELIOS"
OUTPUT_FOLDER = PROJECT_ROOT / "03_clean_data" / "HELIOS"
OUTPUT_FOLDER.mkdir(parents=True, exist_ok=True)
OUTPUT_FILE = OUTPUT_FOLDER / "Helios_All_Funds_Cleaned.xlsx"

AMC_NAME = "Helios MF"
START_DATE = pd.Timestamp(year=2024, month=10, day=1)
END_DATE = pd.Timestamp(year=2026, month=8, day=31)

STANDARD_COLUMNS = [
    "AMC",
    "Fund_Name",
    "Portfolio_Date",
    "Month",
    "Security_Name",
    "ISIN",
    "Industry_Rating",
    "Quantity",
]

TARGET_FUNDS_ORDERED = [
    ("Helios Large & Mid Cap Fund", ["large_and_mid", "large_mid", "large & mid", "large and mid", "large mid"]),
    ("Helios Small Cap Fund", ["smallcap", "small_cap", "small cap"]),
    ("Helios Mid Cap Fund", ["midcap", "mid_cap", "mid cap"]),
    ("Helios Flexi Cap Fund", ["flexicap", "flexi_cap", "flexi cap"]),
    ("Helios Balanced Advantage Fund", ["balanced_advantage", "balanced advantage", "baf"]),
    ("Helios Financial Services Fund", ["financial_services", "financial services", "finserv"]),
]

MONTH_MAP = {
    "january": 1, "jan": 1,
    "february": 2, "feb": 2,
    "march": 3, "mar": 3,
    "april": 4, "apr": 4,
    "may": 5,
    "june": 6, "jun": 6,
    "july": 7, "jul": 7,
    "august": 8, "aug": 8,
    "september": 9, "sep": 9, "sept": 9,
    "october": 10, "oct": 10,
    "november": 11, "nov": 11,
    "december": 12, "dec": 12,
}


def clean_text(val):
    if val is None:
        return ""
    text = str(val).strip()
    return re.sub(r"\s+", " ", text)


def match_canonical_fund_name(filename):
    norm = filename.lower().replace("-", "_").replace(" ", "_")
    for cname, patterns in TARGET_FUNDS_ORDERED:
        if any(p.replace(" ", "_") in norm for p in patterns):
            return cname
    return None


def parse_portfolio_date_from_rows(rows_data, file_path):
    for r in range(min(15, len(rows_data))):
        row = rows_data[r]
        for cell in row:
            if isinstance(cell, datetime):
                return cell
            if cell and "PORTFOLIO STATEMENT AS ON" in str(cell).upper():
                text = str(cell)
                m = re.search(r"(\d{4}-\d{2}-\d{2})", text)
                if m:
                    return datetime.strptime(m.group(1), "%Y-%m-%d")
                m2 = re.search(r"as\s+on\s*:?\s*([A-Za-z]+)\s+(\d{1,2}),?\s+(\d{4})", text, re.I)
                if m2:
                    m_str, d_str, y_str = m2.groups()
                    mo_num = MONTH_MAP.get(m_str.lower().strip())
                    if mo_num:
                        return datetime(int(y_str), mo_num, int(d_str))
                m3 = re.search(r"as\s+on\s*:?\s*(\d{1,2})(?:st|nd|rd|th)?\s+([A-Za-z]+),?\s+(\d{4})", text, re.I)
                if m3:
                    d_str, m_str, y_str = m3.groups()
                    mo_num = MONTH_MAP.get(m_str.lower().strip())
                    if mo_num:
                        return datetime(int(y_str), mo_num, int(d_str))
                m4 = re.search(r"as\s+on\s*:?\s*(\d{1,2})[.\-/](\d{1,2})[.\-/](\d{4})", text, re.I)
                if m4:
                    d, m_num, y = m4.groups()
                    return datetime(int(y), int(m_num), int(d))

    # Fallback to directory structure: 01_raw_files/HELIOS/<YYYY>/<MM>/...
    parts = file_path.parts
    try:
        y = int(parts[-3])
        m = int(parts[-2])
        last_day = calendar.monthrange(y, m)[1]
        return datetime(y, m, last_day)
    except Exception:
        return datetime(2026, 8, 31)


def parse_helios_file(file_path):
    is_xls = file_path.suffix.lower() == ".xls"
    canonical_fund_name = match_canonical_fund_name(file_path.name)
    if not canonical_fund_name:
        return []

    try:
        if is_xls:
            df_raw = pd.read_excel(file_path, sheet_name=0, header=None)
            rows_data = [row.tolist() for _, row in df_raw.iterrows()]
        else:
            wb = openpyxl.load_workbook(file_path, data_only=True)
            ws = wb.worksheets[0]
            rows_data = [[ws.cell(r, c).value for c in range(1, ws.max_column + 1)] for r in range(1, ws.max_row + 1)]
    except Exception as e:
        print(f"  [ERROR] Failed to open {file_path.name}: {e}")
        return []

    if not rows_data:
        return []

    portfolio_dt = parse_portfolio_date_from_rows(rows_data, file_path)
    date_str = portfolio_dt.strftime("%Y-%m-%d")
    month_str = portfolio_dt.strftime("%Y-%m")

    records = []
    header_found = False

    for row in rows_data:
        row_str = " ".join([clean_text(v) for v in row if v is not None]).lower()

        if "name of the instrument" in row_str and ("isin" in row_str or "quantity" in row_str):
            header_found = True
            continue

        if not header_found:
            continue

        # Stop on Total
        is_total = False
        for cell in row:
            if cell is not None and clean_text(cell).lower() in ["total", "sub total", "sub-total", "subtotal", "grand total", "grand total (aum)"]:
                is_total = True
                break
        if is_total:
            break

        # Check section boundaries (e.g. b) Unlisted, Debt Instruments, TREPS, Mutual Fund Units)
        first_val = ""
        for cell in row:
            if cell is not None and str(cell).strip():
                first_val = clean_text(cell).lower()
                break
        if first_val.startswith("b)") or first_val.startswith("c)") or "debt instruments" in first_val or "treps" in first_val or "mutual fund units" in first_val or "net current assets" in first_val:
            break

        # Find ISIN starting with INE
        isin_idx = None
        for i, cell in enumerate(row):
            if cell is not None and isinstance(cell, str) and cell.strip().startswith("INE") and len(cell.strip()) == 12:
                isin_idx = i
                break

        if isin_idx is not None:
            isin = clean_text(row[isin_idx])

            # Security Name is located before ISIN
            sec_name = None
            for i in range(isin_idx - 1, -1, -1):
                val = clean_text(row[i])
                if val and not re.match(r"^\d+(\.0)?$", val) and len(val) > 2:
                    sec_name = val
                    break

            # Industry / Rating and Quantity
            industry = ""
            qty = None
            for i in range(isin_idx + 1, len(row)):
                cell = row[i]
                if cell is None or str(cell).strip() == "":
                    continue
                try:
                    num = float(str(cell).replace(",", "").strip())
                    if qty is None:
                        qty = num
                except ValueError:
                    if not industry:
                        industry = clean_text(cell)

            if sec_name and isin and qty is not None and qty > 0:
                records.append({
                    "AMC": AMC_NAME,
                    "Fund_Name": canonical_fund_name,
                    "Portfolio_Date": date_str,
                    "Month": month_str,
                    "Security_Name": sec_name,
                    "ISIN": isin,
                    "Industry_Rating": industry,
                    "Quantity": int(round(qty))
                })

    return records


def clean_helios_data():
    print("=" * 75)
    print("HELIOS MUTUAL FUND - CLEANING PIPELINE")
    print(f"Raw Folder:   {RAW_FOLDER}")
    print(f"Output File:  {OUTPUT_FILE}")
    print(f"Target Date Range: {START_DATE.strftime('%b %Y')} to {END_DATE.strftime('%b %Y')}")
    print("=" * 75)

    excel_files = sorted(list(RAW_FOLDER.rglob("*.xlsx")) + list(RAW_FOLDER.rglob("*.xls")))
    if not excel_files:
        print(f"ERROR: No monthly files found in {RAW_FOLDER}")
        sys.exit(1)

    print(f"\n[Step 1/3] Found {len(excel_files)} monthly fund workbooks to process.")

    all_rows = []
    skipped_files = 0

    for file_idx, file_path in enumerate(excel_files, 1):
        rows = parse_helios_file(file_path)
        if rows:
            all_rows.extend(rows)
            print(f"  [{file_idx}/{len(excel_files)}] {file_path.parent.parent.name}/{file_path.parent.name}/{file_path.name[:45]}... -> {rows[0]['Fund_Name']} ({len(rows)} holdings)")
        else:
            skipped_files += 1
            print(f"  [{file_idx}/{len(excel_files)}] [SKIP/EMPTY] {file_path.name}")

    print(f"\n[Step 2/3] Raw extraction complete. Total holdings rows: {len(all_rows):,} (Skipped/Empty: {skipped_files})")

    if not all_rows:
        print("ERROR: No holdings extracted. Please check parsers.")
        sys.exit(1)

    df = pd.DataFrame(all_rows)

    # Convert Portfolio_Date to datetime for filtering and standard format
    df["Portfolio_Date"] = pd.to_datetime(df["Portfolio_Date"])

    # Filter date range
    df = df[(df["Portfolio_Date"] >= START_DATE) & (df["Portfolio_Date"] <= END_DATE)].copy()

    # Reformat Portfolio_Date as YYYY-MM-DD string
    df["Portfolio_Date"] = df["Portfolio_Date"].dt.strftime("%Y-%m-%d")

    # Filter ISIN starting with INE
    df = df[df["ISIN"].astype(str).str.startswith("INE")].copy()

    # Ensure positive Quantity
    df["Quantity"] = pd.to_numeric(df["Quantity"], errors="coerce")
    df = df[df["Quantity"] > 0].copy()
    df["Quantity"] = df["Quantity"].astype(int)

    # Ensure standard column order
    df = df[STANDARD_COLUMNS]

    # Sort
    df = df.sort_values(by=["Fund_Name", "Portfolio_Date", "Security_Name"]).reset_index(drop=True)

    print(f"\n[Step 3/3] Saving cleaned data to Excel: {OUTPUT_FILE}")
    with pd.ExcelWriter(OUTPUT_FILE, engine="openpyxl") as writer:
        df.to_excel(writer, index=False, sheet_name="Sheet1")

    print("\n" + "=" * 75)
    print("HELIOS MUTUAL FUND - CLEANING SUMMARY")
    print("=" * 75)
    print(f"Total Rows Saved:    {len(df):,}")
    print(f"Unique AMC:          {df['AMC'].unique().tolist()}")
    print(f"Unique Funds ({df['Fund_Name'].nunique()}):")
    for fn in sorted(df["Fund_Name"].unique()):
        print(f"  - {fn}: {len(df[df['Fund_Name'] == fn]):,} rows across {df[df['Fund_Name'] == fn]['Month'].nunique()} months")
    print(f"Unique Months ({df['Month'].nunique()}): {sorted(df['Month'].unique())}")
    print(f"Portfolio Date Range: {df['Portfolio_Date'].min()} to {df['Portfolio_Date'].max()}")
    print(f"Null Values per Column:\n{df.isnull().sum()}")
    print("=" * 75)


if __name__ == "__main__":
    clean_helios_data()
