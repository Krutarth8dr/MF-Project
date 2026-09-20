import os
import re
import sys
import traceback
from datetime import datetime
from pathlib import Path
import openpyxl
import pandas as pd

# Force UTF-8 output on Windows
try:
    sys.stdout.reconfigure(encoding="utf-8")
    sys.stderr.reconfigure(encoding="utf-8")
except (AttributeError, ValueError):
    pass

# ==============================================================================
# SETTINGS & PATHS
# ==============================================================================

SCRIPT_DIR = Path(__file__).resolve().parent
PROJECT_ROOT = SCRIPT_DIR.parents[1]

RAW_FOLDER = PROJECT_ROOT / "01_raw_files" / "BANDHAN_MF"
OUTPUT_FOLDER = PROJECT_ROOT / "03_clean_data" / "BANDHAN_MF"
OUTPUT_FOLDER.mkdir(parents=True, exist_ok=True)

OUTPUT_FILE = OUTPUT_FOLDER / "Bandhan_MF_All_Funds_Cleaned.xlsx"

AMC_NAME = "Bandhan MF"

STANDARD_COLUMNS = [
    "AMC",
    "Fund_Name",
    "Portfolio_Date",
    "Month",
    "Security_Name",
    "ISIN",
    "Industry_Rating",
    "Quantity",
]

STOP_MARKERS = [
    "subtotal",
    "sub total",
    "sub-total",
    "(b) unlisted",
    "(b) privately placed / unlisted",
]

EXCLUDED_FUNDS = [
    "Bandhan Short Term Fund",
    "Bandhan Medium Term Fund",
    "Bandhan Medium to Long Term Fund",
]



# ==============================================================================
# HELPER FUNCTIONS
# ==============================================================================

def clean_text(val):
    if val is None:
        return ""
    text = str(val).strip()
    return re.sub(r"\s+", " ", text)


def is_valid_isin(val):
    if not val:
        return False
    val_clean = str(val).strip().upper()
    return bool(re.fullmatch(r"INE[A-Z0-9]{9}", val_clean))


def extract_metadata_from_file(file_path: Path):
    """
    Extracts fund_name and portfolio_date from filename:
    e.g. 'Bandhan Large and Mid Cap Fund - 2026-08.xlsx' ->
         fund_name: 'Bandhan Large and Mid Cap Fund'
         portfolio_date: 2026-08-01
    """
    stem = file_path.stem
    match = re.search(r"^(.*?)\s*-\s*(\d{4})-(\d{2})$", stem)
    if match:
        fund_name = match.group(1).strip()
        year = int(match.group(2))
        month = int(match.group(3))
        portfolio_date = pd.Timestamp(year=year, month=month, day=1)
        return fund_name, portfolio_date

    # Fallback to date in filename like '31 August 2026'
    m_date = re.search(r"(\d{1,2})\s+([A-Za-z]+)\s+(\d{4})", stem)
    if m_date:
        month_str = m_date.group(2).lower()
        year = int(m_date.group(3))
        month_map = {
            "january": 1, "february": 2, "march": 3, "april": 4,
            "may": 5, "june": 6, "july": 7, "august": 8,
            "september": 9, "october": 10, "november": 11, "december": 12
        }
        month = month_map.get(month_str, 8)
        portfolio_date = pd.Timestamp(year=year, month=month, day=1)
        fund_name = re.sub(r"\s*\d{1,2}\s+[A-Za-z]+\s+\d{4}.*$", "", stem).strip()
        return fund_name, portfolio_date

    return stem, None


def clean_single_bandhan_file(file_path: Path, canonical_fund_name: str, portfolio_date: pd.Timestamp):
    if canonical_fund_name in EXCLUDED_FUNDS:
        return pd.DataFrame(columns=STANDARD_COLUMNS)

    wb = openpyxl.load_workbook(file_path, data_only=True, read_only=True)

    sheet = wb.active

    # Find header row
    name_col = None
    isin_col = None
    ind_col = None
    qty_col = None
    header_row_idx = None

    for r_idx, row in enumerate(sheet.iter_rows(max_row=20, values_only=True), start=1):
        row_str = [clean_text(v).lower() for v in row if v is not None]
        if any("name of instrument" in cell or "name of the instrument" in cell for cell in row_str):
            header_row_idx = r_idx
            for c_idx, cell_val in enumerate(row):
                c_clean = clean_text(cell_val).lower()
                if "name of instrument" in c_clean or "name of the instrument" in c_clean:
                    name_col = c_idx
                elif "isin" in c_clean:
                    isin_col = c_idx
                elif "industry" in c_clean or "rating" in c_clean:
                    ind_col = c_idx
                elif "quantity" in c_clean:
                    qty_col = c_idx
            break

    if header_row_idx is None or name_col is None or isin_col is None or qty_col is None:
        wb.close()
        raise ValueError(f"Could not locate required headers in {file_path.name}")

    records = []
    month_label = portfolio_date.strftime("%b-%Y")

    # Read data rows after header
    for row in sheet.iter_rows(min_row=header_row_idx + 1, values_only=True):
        if not any(row):
            continue

        raw_instrument = row[name_col] if name_col < len(row) else None
        instrument_str = clean_text(raw_instrument)
        instrument_lower = instrument_str.lower()
        instrument_upper = instrument_str.upper()

        # Stop condition specified by user:
        # Stop scanning after "Total", "Subtotal", "REIT", or "(b) UNLISTED" in the "Name of Instrument" column
        if (
            instrument_lower in ["sub total", "subtotal", "sub-total", "total", "grand total", "total:"]
            or instrument_upper in ["TOTAL", "GRAND TOTAL", "REIT", "REITS"]
            or instrument_upper.endswith(" REIT")
            or instrument_upper.endswith(" REITS")
            or instrument_upper.endswith("(REIT)")
            or bool(re.search(r"\(REIT\)|\bREITS?\b", instrument_upper))
            or instrument_lower.startswith("(b) unlisted")
            or instrument_lower.startswith("b) unlisted")
            or instrument_lower == "unlisted"
            or any(marker in instrument_lower for marker in STOP_MARKERS)
        ):
            break

        raw_isin = row[isin_col] if isin_col < len(row) else None
        isin_str = clean_text(raw_isin).upper()

        if not is_valid_isin(isin_str):
            continue

        raw_qty = row[qty_col] if qty_col < len(row) else 0
        try:
            qty_num = float(raw_qty) if raw_qty is not None else 0
        except (ValueError, TypeError):
            continue

        if qty_num <= 0:
            continue

        # Normalization for corrupted company name text in raw files
        if isin_str == "INE933S01016" and "total outstanding exposure" in instrument_lower:
            instrument_str = "Indiamart Intermesh Limited"

        raw_ind = row[ind_col] if (ind_col is not None and ind_col < len(row)) else ""
        ind_str = clean_text(raw_ind)

        records.append({
            "AMC": AMC_NAME,
            "Fund_Name": canonical_fund_name,
            "Portfolio_Date": portfolio_date,
            "Month": month_label,
            "Security_Name": instrument_str,
            "ISIN": isin_str,
            "Industry_Rating": ind_str,
            "Quantity": int(round(qty_num)),
        })

    wb.close()
    return pd.DataFrame(records)


# ==============================================================================
# MAIN PROCESSING LOOP
# ==============================================================================

def main():
    print("=" * 80)
    print("CLEANING BANDHAN MF MONTHLY FILES")
    print(f"Raw Folder : {RAW_FOLDER}")
    print(f"Output File: {OUTPUT_FILE}")
    print("=" * 80)

    # Collect raw files
    raw_files = sorted(list(RAW_FOLDER.glob("**/*.xlsx")))
    raw_files = [f for f in raw_files if not f.name.startswith("~") and not f.name.startswith("_") and "_temp" not in str(f) and "_test" not in str(f)]

    print(f"\nFound {len(raw_files)} candidate workbooks.")
    if not raw_files:
        print("No raw files to clean.")
        return

    # Check for existing cleaned file (unless --rebuild or --force is passed)
    force_rebuild = "--rebuild" in sys.argv or "--force" in sys.argv
    existing_df = None
    processed_keys = set()
    removed_excluded_count = 0
    if OUTPUT_FILE.exists() and not force_rebuild:
        try:
            existing_df = pd.read_excel(OUTPUT_FILE)
            if all(col in existing_df.columns for col in ["Fund_Name", "Portfolio_Date"]):
                initial_count = len(existing_df)
                existing_df = existing_df[~existing_df["Fund_Name"].isin(EXCLUDED_FUNDS)].copy()
                removed_excluded_count = initial_count - len(existing_df)
                if removed_excluded_count > 0:
                    print(f"Removed {removed_excluded_count:,} rows belonging to excluded funds: {EXCLUDED_FUNDS}")
                existing_df["Portfolio_Date"] = pd.to_datetime(existing_df["Portfolio_Date"], errors="coerce")
                processed_keys = set(zip(existing_df["Fund_Name"], existing_df["Portfolio_Date"]))
                print(f"Loaded existing cleaned file with {len(existing_df):,} rows ({len(processed_keys)} fund-month keys).")
            else:
                existing_df = None
        except Exception as e:
            print(f"Could not read existing cleaned file: {e}")
            existing_df = None
    elif force_rebuild:
        print("⚡ Force rebuild requested: Re-cleaning all candidate workbooks from scratch.")

    cleaned_dfs = []
    new_keys_processed = 0

    for file_path in raw_files:
        fund_name, portfolio_date = extract_metadata_from_file(file_path)
        if portfolio_date is None or fund_name in EXCLUDED_FUNDS:
            continue

        if (fund_name, portfolio_date) in processed_keys:
            continue

        try:
            df = clean_single_bandhan_file(file_path, fund_name, portfolio_date)
            if not df.empty:
                cleaned_dfs.append(df)
                processed_keys.add((fund_name, portfolio_date))
                new_keys_processed += 1
                print(f"  ✓ {fund_name} ({portfolio_date.strftime('%Y-%m')}) -> {len(df)} rows")
            else:
                print(f"  ⚠️ {fund_name} ({portfolio_date.strftime('%Y-%m')}) -> 0 valid equity holdings")
        except Exception as e:
            print(f"  ❌ Error cleaning {file_path.name}: {e}")

    if not cleaned_dfs and existing_df is None:
        print("\nNo cleaned data produced.")
        return

    if not cleaned_dfs and removed_excluded_count == 0:
        print("\nNo new files to add. Existing cleaned file is up-to-date.")
        return


    if existing_df is not None:
        all_dfs = [existing_df] + cleaned_dfs
    else:
        all_dfs = cleaned_dfs

    final_df = pd.concat(all_dfs, ignore_index=True)
    final_df["Portfolio_Date"] = pd.to_datetime(final_df["Portfolio_Date"], errors="coerce")
    final_df = final_df.drop_duplicates(subset=["AMC", "Fund_Name", "Portfolio_Date", "ISIN", "Security_Name"])
    final_df = final_df.sort_values(by=["Portfolio_Date", "Fund_Name", "Security_Name"]).reset_index(drop=True)
    final_df = final_df[STANDARD_COLUMNS]

    final_df.to_excel(OUTPUT_FILE, index=False)

    print("\n" + "=" * 80)
    print("BANDHAN MF CLEANING SUMMARY")
    print(f"Total Rows      : {len(final_df):,}")
    print(f"New Combinations: {new_keys_processed}")
    print(f"Unique Funds    : {final_df['Fund_Name'].nunique()}")
    print(f"Date Range      : {final_df['Portfolio_Date'].min().strftime('%Y-%m')} to {final_df['Portfolio_Date'].max().strftime('%Y-%m')}")
    print(f"Output Path     : {OUTPUT_FILE}")
    print("=" * 80)


if __name__ == "__main__":
    main()
