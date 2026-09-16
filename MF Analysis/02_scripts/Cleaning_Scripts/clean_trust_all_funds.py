import os
import re
import sys
import io
from datetime import datetime
from pathlib import Path
import openpyxl
import xlrd
import pandas as pd

# Windows console cp1252 handling
try:
    sys.stdout.reconfigure(encoding="utf-8")
    sys.stderr.reconfigure(encoding="utf-8")
except (AttributeError, ValueError):
    pass

# ==============================================================================
# SETTINGS & PATHS
# ==============================================================================

PROJECT_ROOT = Path(__file__).resolve().parents[2]
RAW_FOLDER = PROJECT_ROOT / "01_raw_files" / "TRUST_MF"
OUTPUT_FOLDER = PROJECT_ROOT / "03_clean_data" / "TRUST_MF"
OUTPUT_FOLDER.mkdir(parents=True, exist_ok=True)
OUTPUT_FILE = OUTPUT_FOLDER / "Trust_MF_All_Funds_Cleaned.xlsx"

AMC_NAME = "Trust MF"
START_DATE = pd.Timestamp(year=2024, month=10, day=1)
END_DATE = pd.Timestamp(year=2026, month=8, day=31)

STANDARD_COLUMNS = [
    "AMC",
    "Fund_Name",
    "Portfolio_Date",
    "Month",
    "Security_Name",
    "ISIN",
    "Industry_Rating",
    "Quantity",
]

# Canonical Target 5 Equity Funds
TARGET_FUNDS_CONFIG = {
    "TMFFLEXI": "TRUSTMF Flexi Cap Fund",
    "TMFSCAP": "TRUSTMF Small Cap Fund",
    "TMFMCAP": "TRUSTMF Multi Cap Fund",
    "TMFMID": "TRUSTMF Mid Cap Fund",
    "TMFLRMCF": "TRUSTMF Large & Mid Cap Fund",
    # Full sheet name aliases in older files
    "TRUSTMF FLEXI CAP FUND": "TRUSTMF Flexi Cap Fund",
    "TRUSTMF SMALL CAP FUND": "TRUSTMF Small Cap Fund",
    "TRUSTMF MULTI CAP FUND": "TRUSTMF Multi Cap Fund",
    "TRUSTMF MID CAP FUND": "TRUSTMF Mid Cap Fund",
    "TRUSTMF LARGE & MID CAP FUND": "TRUSTMF Large & Mid Cap Fund",
    "TRUSTMF LARGE AND MID CAP FUND": "TRUSTMF Large & Mid Cap Fund",
}


def clean_text(val):
    if val is None:
        return ""
    text = str(val).strip()
    return re.sub(r"\s+", " ", text)


def normalize_sheet_name(name):
    clean = clean_text(name).upper()
    return clean


def parse_row_data(row_vals):
    """
    Given a list of cell values for a row, determine if it represents an equity holding.
    Equity holdings have an ISIN starting with 'INE' of length 12.
    """
    isin_val = None
    isin_idx = -1

    for idx, val in enumerate(row_vals):
        v_str = str(val).strip().upper() if val is not None else ""
        if v_str.startswith("INE") and len(v_str) == 12:
            isin_val = v_str
            isin_idx = idx
            break

    if not isin_val or isin_idx < 0:
        return None

    # In Trust MF files:
    # Pattern A (newer xlsx with code):
    # Col 0: Ticker/Code (optional)
    # Col 1: Security Name
    # Col 2: ISIN
    # Col 3: Industry
    # Col 4: Quantity
    #
    # Pattern B (older xls):
    # Col 0: Security Name
    # Col 1: ISIN
    # Col 2: Industry / Rating
    # Col 3: Quantity

    sec_name = ""
    industry = ""
    qty = 0

    if isin_idx > 0:
        sec_name = clean_text(row_vals[isin_idx - 1])

    if isin_idx + 1 < len(row_vals):
        industry = clean_text(row_vals[isin_idx + 1])

    if isin_idx + 2 < len(row_vals):
        q_raw = row_vals[isin_idx + 2]
        try:
            q_clean = str(q_raw).replace(",", "").strip()
            qty = float(q_clean)
        except (ValueError, TypeError):
            qty = 0

    if qty <= 0:
        # Check if quantity was in another adjacent column
        for offset in [3, 4, -2]:
            if 0 <= isin_idx + offset < len(row_vals):
                try:
                    q_clean = str(row_vals[isin_idx + offset]).replace(",", "").strip()
                    val_float = float(q_clean)
                    if val_float > 0:
                        qty = val_float
                        break
                except (ValueError, TypeError):
                    pass

    # Basic validity checks
    if not sec_name or "subtotal" in sec_name.lower() or "total" in sec_name.lower():
        return None
    if qty <= 0:
        return None

    return {
        "security_name": sec_name,
        "isin": isin_val,
        "industry": industry or "Others",
        "quantity": qty,
    }


def clean_trust_data():
    print("=" * 75)
    print("TRUST MF - CLEANING PIPELINE")
    print(f"Raw Folder:        {RAW_FOLDER}")
    print(f"Output File:       {OUTPUT_FILE}")
    print(f"Target Date Range: {START_DATE.strftime('%b %Y')} to {END_DATE.strftime('%b %Y')}")
    print("=" * 75)

    monthly_dirs = sorted(list(RAW_FOLDER.glob("*/*")))
    if not monthly_dirs:
        print(f"ERROR: No monthly files found in {RAW_FOLDER}")
        sys.exit(1)

    print(f"\n[Step 1/3] Found {len(monthly_dirs)} monthly folders to scan.")

    all_rows = []

    for mdir in monthly_dirs:
        if not mdir.is_dir() or mdir.name.startswith("_"):
            continue

        try:
            yr = int(mdir.parent.name)
            mo = int(mdir.name)
        except ValueError:
            continue

        # Calculate exact month-end date
        portfolio_date = pd.Period(f"{yr}-{mo:02d}", freq="M").end_time.date()
        month_str = portfolio_date.strftime("%b-%Y")
        portfolio_dt = pd.to_datetime(portfolio_date)

        if not (START_DATE <= portfolio_dt <= END_DATE):
            continue

        files = list(mdir.glob("*.*"))
        if not files:
            continue

        fpath = files[0]
        print(f"  Processing {month_str} ({fpath.name})...")

        # Try openpyxl first, then xlrd
        parsed_sheets = False

        # Attempt 1: openpyxl
        try:
            wb = openpyxl.load_workbook(fpath, data_only=True)
            for sname in wb.sheetnames:
                norm_s = normalize_sheet_name(sname)
                canonical_name = TARGET_FUNDS_CONFIG.get(norm_s)
                if not canonical_name:
                    continue

                ws = wb[sname]
                for r in range(1, ws.max_row + 1):
                    row_vals = [ws.cell(r, c).value for c in range(1, ws.max_column + 1)]
                    item = parse_row_data(row_vals)
                    if item:
                        all_rows.append({
                            "AMC": AMC_NAME,
                            "Fund_Name": canonical_name,
                            "Portfolio_Date": portfolio_date,
                            "Month": month_str,
                            "Security_Name": item["security_name"],
                            "ISIN": item["isin"],
                            "Industry_Rating": item["industry"],
                            "Quantity": item["quantity"],
                        })
            parsed_sheets = True
        except Exception:
            parsed_sheets = False

        # Attempt 2: xlrd
        if not parsed_sheets:
            try:
                wb = xlrd.open_workbook(fpath)
                for sname in wb.sheet_names():
                    norm_s = normalize_sheet_name(sname)
                    canonical_name = TARGET_FUNDS_CONFIG.get(norm_s)
                    if not canonical_name:
                        continue

                    sh = wb.sheet_by_name(sname)
                    for r in range(sh.nrows):
                        row_vals = [sh.cell_value(r, c) for c in range(sh.ncols)]
                        item = parse_row_data(row_vals)
                        if item:
                            all_rows.append({
                                "AMC": AMC_NAME,
                                "Fund_Name": canonical_name,
                                "Portfolio_Date": portfolio_date,
                                "Month": month_str,
                                "Security_Name": item["security_name"],
                                "ISIN": item["isin"],
                                "Industry_Rating": item["industry"],
                                "Quantity": item["quantity"],
                            })
                parsed_sheets = True
            except Exception as e:
                print(f"    [ERROR] Failed to parse {fpath.name}: {e}")

    print(f"\n[Step 2/3] Extracted {len(all_rows):,} holding records.")
    if not all_rows:
        print("ERROR: No rows extracted! Check file formats or sheet configs.")
        sys.exit(1)

    df = pd.DataFrame(all_rows)

    # Clean duplicates and sort
    df.drop_duplicates(subset=["AMC", "Fund_Name", "Portfolio_Date", "ISIN"], inplace=True)
    df.sort_values(by=["Fund_Name", "Portfolio_Date", "Security_Name"], inplace=True)

    print(f"\n[Step 3/3] Saving cleaned output to {OUTPUT_FILE.name}...")
    df.to_excel(OUTPUT_FILE, index=False)

    print("\n" + "=" * 75)
    print("CLEANING COMPLETED SUCCESSFULLY!")
    print(f"Total Rows:     {len(df):,}")
    print(f"Funds Covered:  {df['Fund_Name'].nunique()} ({list(df['Fund_Name'].unique())})")
    print(f"Months Covered: {df['Month'].nunique()} ({df['Month'].min()} to {df['Month'].max()})")
    print(f"ISINs Covered:  {df['ISIN'].nunique():,}")
    print(f"Saved Path:     {OUTPUT_FILE}")
    print("=" * 75)


if __name__ == "__main__":
    clean_trust_data()
