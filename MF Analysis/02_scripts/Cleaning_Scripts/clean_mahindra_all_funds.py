"""
Clean all downloaded monthly portfolio files for Mahindra Manulife Mutual Fund (Mahindra Manulife MF).

Target Schemes:
    MMF02, MMF03, MMF05, MMF07, MMF09, MMF10 (MF10), MMF11, MMF14,
    MMF16, MMF18, MMF20, MMF21, MMF22, MMF23, MMF24, MMF25,
    MMF26, MMF28 (18 funds).

Process:
    - Reads all files in 01_raw_files/Mahindra_Manulife_MF/<YYYY>/<MM>/
    - Loops through the target scheme sheets
    - Stops immediately upon encountering 'Sub Total', '(b)Unlisted', '(b) Unlisted', or 'Total'
    - Filters for valid listed Indian equities (ISIN starts with INE, length 12, quantity > 0)
    - Generates 03_clean_data/Mahindra_Manulife_MF/Mahindra_Manulife_MF_All_Funds_Cleaned.xlsx
"""

import sys
import os
import calendar
from datetime import datetime
from pathlib import Path
import openpyxl
import pandas as pd

# Windows consoles default to cp1252, force UTF-8 output
try:
    sys.stdout.reconfigure(encoding="utf-8")
    sys.stderr.reconfigure(encoding="utf-8")
except (AttributeError, ValueError):
    pass

PROJECT_ROOT = Path(__file__).resolve().parents[2]
RAW_DIR = PROJECT_ROOT / "01_raw_files" / "Mahindra_Manulife_MF"
CLEAN_DIR = PROJECT_ROOT / "03_clean_data" / "Mahindra_Manulife_MF"
CLEAN_DIR.mkdir(parents=True, exist_ok=True)

OUTPUT_FILE = CLEAN_DIR / "Mahindra_Manulife_MF_All_Funds_Cleaned.xlsx"

TARGET_SHEETS = [
    "MMF02", "MMF03", "MMF05", "MMF07", "MMF09", "MMF10", "MMF11", "MMF14",
    "MMF16", "MMF18", "MMF20", "MMF21", "MMF22", "MMF23", "MMF24", "MMF25",
    "MMF26", "MMF28"
]

FUND_NAME_FALLBACK = {
    "MMF02": "Mahindra Manulife ELSS Tax Saver Fund",
    "MMF03": "Mahindra Manulife Equity Savings Fund",
    "MMF05": "Mahindra Manulife Multi Cap Fund",
    "MMF07": "Mahindra Manulife Mid Cap Fund",
    "MMF09": "Mahindra Manulife Consumption Fund",
    "MMF10": "Mahindra Manulife Large Cap Fund",
    "MMF11": "Mahindra Manulife Aggressive Hybrid Fund",
    "MMF14": "Mahindra Manulife Large & Mid Cap Fund",
    "MMF16": "Mahindra Manulife Focused Fund",
    "MMF18": "Mahindra Manulife Flexi Cap Fund",
    "MMF20": "Mahindra Manulife Balanced Advantage Fund",
    "MMF21": "Mahindra Manulife Small Cap Fund",
    "MMF22": "Mahindra Manulife Business Cycle Fund",
    "MMF23": "Mahindra Manulife Multi Asset Allocation Fund",
    "MMF24": "Mahindra Manulife Manufacturing Fund",
    "MMF25": "Mahindra Manulife Value Fund",
    "MMF26": "Mahindra Manulife Banking & Financial Services Fund",
    "MMF28": "Mahindra Manulife Innovation Opportunities Fund",
}

STOP_TERMS = ["sub total", "(b)unlisted", "(b) unlisted", "total"]

STANDARD_COLUMNS = [
    "AMC",
    "Fund_Name",
    "Portfolio_Date",
    "Month",
    "Security_Name",
    "ISIN",
    "Industry_Rating",
    "Quantity",
]


def clean_file(file_path):
    """Parses a monthly multi-scheme workbook and extracts listed equity holdings."""
    holdings = []
    try:
        month = int(file_path.parent.name)
        year = int(file_path.parent.parent.name)

        last_day = calendar.monthrange(year, month)[1]
        portfolio_date = f"{year}-{month:02d}-{last_day:02d}"
        month_str = datetime.strptime(f"{year}-{month:02d}-01", "%Y-%m-%d").strftime("%b-%Y")

        wb = openpyxl.load_workbook(file_path, data_only=True)
        sheet_map = {s.strip().upper(): s for s in wb.sheetnames}

        for code in TARGET_SHEETS:
            actual_sheet = None
            if code in sheet_map:
                actual_sheet = sheet_map[code]
            elif code == "MMF10" and "MF10" in sheet_map:
                actual_sheet = sheet_map["MF10"]

            if not actual_sheet:
                continue

            ws = wb[actual_sheet]

            # Fund name from row 3 column 2 or fallback
            fund_name = str(ws.cell(row=3, column=2).value or "").strip()
            if not fund_name or len(fund_name) < 5:
                fund_name = FUND_NAME_FALLBACK.get(code, code)

            # Locate header row
            header_row = None
            col_map = {}
            for r in range(1, 15):
                row_vals = [str(ws.cell(row=r, column=c).value or "").strip() for c in range(1, ws.max_column + 1)]
                lower_vals = [v.lower() for v in row_vals]
                if any("isin" in v for v in lower_vals) and any("instrument" in v or "name" in v for v in lower_vals):
                    header_row = r
                    for c_idx, val in enumerate(lower_vals, 1):
                        if "isin" in val:
                            col_map["isin"] = c_idx
                        elif ("instrument" in val or "name" in val) and "yield" not in val:
                            col_map["name"] = c_idx
                        elif "rating" in val or "industry" in val:
                            col_map["rating"] = c_idx
                        elif "quantity" in val or "qty" in val:
                            col_map["qty"] = c_idx
                    break

            if not header_row or "isin" not in col_map or "name" not in col_map or "qty" not in col_map:
                continue

            for r in range(header_row + 1, ws.max_row + 1):
                name_val = str(ws.cell(row=r, column=col_map["name"]).value or "").strip()
                isin_val = str(ws.cell(row=r, column=col_map["isin"]).value or "").strip()

                name_lower = name_val.lower()
                isin_lower = isin_val.lower()

                # Stop condition
                if any(term in name_lower or term in isin_lower for term in STOP_TERMS):
                    break

                # Filter for listed equity
                if not (isin_val.startswith("INE") and len(isin_val) == 12):
                    continue

                qty_val = ws.cell(row=r, column=col_map["qty"]).value
                try:
                    qty = float(qty_val)
                except (ValueError, TypeError):
                    continue

                if qty <= 0:
                    continue

                rating_val = ""
                if "rating" in col_map:
                    rating_val = str(ws.cell(row=r, column=col_map["rating"]).value or "").strip()

                holdings.append({
                    "AMC": "Mahindra Manulife MF",
                    "Fund_Name": fund_name,
                    "Portfolio_Date": portfolio_date,
                    "Month": month_str,
                    "Security_Name": name_val,
                    "ISIN": isin_val,
                    "Industry_Rating": rating_val,
                    "Quantity": qty,
                })

        wb.close()

    except Exception as e:
        print(f"  [ERROR] {file_path.name}: {e}")

    return holdings


def main():
    print("=" * 80)
    print("Mahindra Manulife MF - Portfolio Cleaning Pipeline")
    print("=" * 80)

    files = sorted(list(RAW_DIR.rglob("*.xlsx")))
    print(f"Found {len(files)} monthly portfolio files in {RAW_DIR}")

    all_holdings = []
    for idx, f in enumerate(files, 1):
        file_holdings = clean_file(f)
        all_holdings.extend(file_holdings)
        print(f"[{idx:2d}/{len(files)}] Processed {f.name} ({len(file_holdings)} holdings, {len(all_holdings)} total)")

    if not all_holdings:
        print("❌ No holdings extracted. Please check input raw files.")
        return

    df = pd.DataFrame(all_holdings)[STANDARD_COLUMNS]

    initial_len = len(df)
    df.drop_duplicates(subset=["AMC", "Fund_Name", "Portfolio_Date", "ISIN", "Security_Name"], inplace=True)
    if len(df) < initial_len:
        print(f"Deduplicated {initial_len - len(df)} duplicate records.")

    df.sort_values(by=["Fund_Name", "Portfolio_Date", "Security_Name"], inplace=True)

    df.to_excel(OUTPUT_FILE, index=False)

    print("\n" + "=" * 80)
    print("Cleaned Data Summary")
    print("=" * 80)
    print(f"Total Records : {len(df)}")
    print(f"Unique AMCs   : {df['AMC'].nunique()} ({df['AMC'].unique().tolist()})")
    print(f"Unique Funds  : {df['Fund_Name'].nunique()}")
    for fund, cnt in df["Fund_Name"].value_counts().items():
        print(f"  - {fund:52} : {cnt:5d} rows")
    print(f"Unique ISINs  : {df['ISIN'].nunique()}")
    print(f"Date Range    : {df['Portfolio_Date'].min()} to {df['Portfolio_Date'].max()} ({df['Portfolio_Date'].nunique()} months)")
    print(f"Output File   : {OUTPUT_FILE}")
    print("=" * 80)


if __name__ == "__main__":
    main()
