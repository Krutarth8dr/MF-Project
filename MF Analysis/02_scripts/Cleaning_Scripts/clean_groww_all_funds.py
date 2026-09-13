import os
import re
import sys
from datetime import datetime
from pathlib import Path
import warnings

import openpyxl
import pandas as pd
import xlrd

# Suppress openpyxl user warnings for custom metadata
warnings.filterwarnings("ignore", category=UserWarning, module="openpyxl")

# ==============================================================================
# SETTINGS & PATHS
# ==============================================================================

PROJECT_ROOT = Path(__file__).resolve().parents[2]
RAW_FOLDER = PROJECT_ROOT / "01_raw_files" / "GROWW"
CLEAN_FOLDER = PROJECT_ROOT / "03_clean_data" / "GROWW"
CLEAN_FOLDER.mkdir(parents=True, exist_ok=True)

OUTPUT_FILE = CLEAN_FOLDER / "Groww_All_Funds_Cleaned.xlsx"

AMC_NAME = "Groww MF"

TARGET_CONFIG = {
    # 2-letter codes (used in 2026 workbooks)
    "BC": ("BC", "Groww Large Cap Fund"),
    "EH": ("EH", "Groww Aggressive Hybrid Fund"),
    "TS": ("TS", "Groww ELSS Tax Saver Fund"),
    "VD": ("VD", "Groww Value Fund"),
    "BS": ("BS", "Groww Banking & Financial Services Fund"),
    "NC": ("NC", "Groww Nifty Non-Cyclical Consumer Index Fund"),
    "MU": ("MU", "Groww Multicap Fund"),
    "MA": ("MA", "Groww Multi Asset Allocation Fund"),
    "SC": ("SC", "Groww Small Cap Fund"),
    # IB codes (used in 2024/2025 workbooks and some 2026 sheets)
    "IB01": ("BC", "Groww Large Cap Fund"),
    "IB03": ("EH", "Groww Aggressive Hybrid Fund"),
    "IB11": ("TS", "Groww ELSS Tax Saver Fund"),
    "IB13": ("VD", "Groww Value Fund"),
    "IB19": ("BS", "Groww Banking & Financial Services Fund"),
    "IB21": ("NC", "Groww Nifty Non-Cyclical Consumer Index Fund"),
    "IB29": ("MU", "Groww Multicap Fund"),
    "IB49": ("MA", "Groww Multi Asset Allocation Fund"),
    "IB60": ("SC", "Groww Small Cap Fund"),
}

MONTH_MAP = {
    "jan": 1, "january": 1, "feb": 2, "february": 2, "mar": 3, "march": 3,
    "apr": 4, "april": 4, "may": 5, "jun": 6, "june": 6, "jul": 7, "july": 7,
    "aug": 8, "august": 8, "sep": 9, "sept": 9, "september": 9,
    "oct": 10, "october": 10, "nov": 11, "november": 11, "dec": 12, "december": 12
}

STOP_TRIGGERS = [
    "sub total",
    "(b) unlisted",
    "b) unlisted",
    "unlisted",
    "total",
    "debt instruments",
    "money market",
    "treps",
    "net current assets",
    "total net assets",
    "grand total"
]

# ==============================================================================
# HELPER FUNCTIONS
# ==============================================================================

def parse_date_robust(text):
    """
    Parses date string into a datetime object and returns exact month-end date.
    """
    if not text:
        return None
    s = str(text).replace(",", " ").replace("_", " ")
    # Format: 31-AUG-2026 or 31 Aug 2026
    m = re.search(r'(\d{1,2})\s*[-/ ]\s*([A-Za-z]+)\s*[-/ ]\s*(\d{4})', s)
    if m:
        mon_name, year = m.group(2).lower(), int(m.group(3))
        mon = MONTH_MAP.get(mon_name[:3])
        if mon:
            return pd.Period(f"{year}-{mon:02d}", freq="M").end_time.date()
    # Format: Aug 31 2026 or August 31 2026
    m2 = re.search(r'([A-Za-z]+)\s*[-/ ]\s*(\d{1,2})\s*[-/ ]\s*(\d{4})', s)
    if m2:
        mon_name, year = m2.group(1).lower(), int(m2.group(3))
        mon = MONTH_MAP.get(mon_name[:3])
        if mon:
            return pd.Period(f"{year}-{mon:02d}", freq="M").end_time.date()
    # Format: Oct 2024
    m3 = re.search(r'([A-Za-z]+)\s*[-/ ]\s*(\d{4})', s)
    if m3:
        mon_name, year = m3.group(1).lower(), int(m3.group(2))
        mon = MONTH_MAP.get(mon_name[:3])
        if mon:
            return pd.Period(f"{year}-{mon:02d}", freq="M").end_time.date()
    return None


def is_valid_isin(value):
    """
    Strict equity ISIN: starts with INE, length 12.
    """
    if pd.isna(value):
        return False
    val = str(value).strip().upper()
    return val.startswith("INE") and len(val) == 12


def parse_xlsx_sheet(ws, fund_canonical, fallback_date):
    """
    Extracts holdings from an openpyxl sheet.
    """
    port_date = fallback_date
    for r in range(1, 20):
        for c in range(1, min(10, ws.max_column + 1)):
            val = ws.cell(r, c).value
            if val and "portfolio as on" in str(val).lower():
                dt = parse_date_robust(val)
                if dt:
                    port_date = dt
                    break
        if port_date != fallback_date:
            break

    header_row = None
    col_isin = None
    col_sec = None
    col_ind = None
    col_qty = None

    for r in range(1, min(25, ws.max_row + 1)):
        row_vals = [str(ws.cell(r, c).value or "").strip().lower() for c in range(1, ws.max_column + 1)]
        if "isin" in row_vals:
            header_row = r
            for c_idx, val in enumerate(row_vals, start=1):
                if val == "isin":
                    col_isin = c_idx
                elif "instrument" in val or "security" in val or "company" in val:
                    col_sec = c_idx
                elif "industry" in val or "rating" in val:
                    col_ind = c_idx
                elif "quantity" in val or "shares" in val:
                    col_qty = c_idx
            break

    if not header_row or not col_isin or not col_sec or not col_qty:
        return []

    rows = []
    equity_started = False

    for r in range(header_row + 1, ws.max_row + 1):
        isin_val = ws.cell(r, col_isin).value
        sec_val = ws.cell(r, col_sec).value
        ind_val = ws.cell(r, col_ind).value if col_ind else ""
        qty_val = ws.cell(r, col_qty).value

        isin_str = str(isin_val or "").strip().upper()
        sec_str = str(sec_val or "").strip()
        ind_str = str(ind_val or "").strip()
        sec_lower = sec_str.lower()

        # Preamble gating
        if not equity_started:
            if is_valid_isin(isin_str):
                equity_started = True
            else:
                continue

        # Stop triggers (Sub Total, (b) Unlisted, Total, etc.)
        if any(trigger in sec_lower for trigger in STOP_TRIGGERS):
            break

        if is_valid_isin(isin_str):
            try:
                qty_clean = str(qty_val).replace(",", "").strip()
                qty_num = float(qty_clean)
                if qty_num > 0:
                    rows.append({
                        "AMC": AMC_NAME,
                        "Fund_Name": fund_canonical,
                        "Portfolio_Date": port_date.strftime("%Y-%m-%d"),
                        "Month": port_date.strftime("%b-%Y"),
                        "Security_Name": sec_str,
                        "ISIN": isin_str,
                        "Industry_Rating": ind_str,
                        "Quantity": int(round(qty_num))
                    })
            except Exception:
                pass

    return rows


def parse_xls_sheet(ws, fund_canonical, fallback_date):
    """
    Extracts holdings from an xlrd sheet.
    """
    port_date = fallback_date
    for r in range(min(20, ws.nrows)):
        for c in range(min(10, ws.ncols)):
            val = ws.cell_value(r, c)
            if val and "portfolio as on" in str(val).lower():
                dt = parse_date_robust(val)
                if dt:
                    port_date = dt
                    break
        if port_date != fallback_date:
            break

    header_row = None
    col_isin = None
    col_sec = None
    col_ind = None
    col_qty = None

    for r in range(min(25, ws.nrows)):
        row_vals = [str(ws.cell_value(r, c) or "").strip().lower() for c in range(ws.ncols)]
        if "isin" in row_vals:
            header_row = r
            for c_idx, val in enumerate(row_vals):
                if val == "isin":
                    col_isin = c_idx
                elif "instrument" in val or "security" in val or "company" in val:
                    col_sec = c_idx
                elif "industry" in val or "rating" in val:
                    col_ind = c_idx
                elif "quantity" in val or "shares" in val:
                    col_qty = c_idx
            break

    if header_row is None or col_isin is None or col_sec is None or col_qty is None:
        return []

    rows = []
    equity_started = False

    for r in range(header_row + 1, ws.nrows):
        isin_val = ws.cell_value(r, col_isin)
        sec_val = ws.cell_value(r, col_sec)
        ind_val = ws.cell_value(r, col_ind) if col_ind is not None else ""
        qty_val = ws.cell_value(r, col_qty)

        isin_str = str(isin_val or "").strip().upper()
        sec_str = str(sec_val or "").strip()
        ind_str = str(ind_val or "").strip()
        sec_lower = sec_str.lower()

        if not equity_started:
            if is_valid_isin(isin_str):
                equity_started = True
            else:
                continue

        if any(trigger in sec_lower for trigger in STOP_TRIGGERS):
            break

        if is_valid_isin(isin_str):
            try:
                qty_clean = str(qty_val).replace(",", "").strip()
                qty_num = float(qty_clean)
                if qty_num > 0:
                    rows.append({
                        "AMC": AMC_NAME,
                        "Fund_Name": fund_canonical,
                        "Portfolio_Date": port_date.strftime("%Y-%m-%d"),
                        "Month": port_date.strftime("%b-%Y"),
                        "Security_Name": sec_str,
                        "ISIN": isin_str,
                        "Industry_Rating": ind_str,
                        "Quantity": int(round(qty_num))
                    })
            except Exception:
                pass

    return rows


def parse_groww_file(file_path):
    """
    Parses a single Groww monthly file (.xlsx or .xls).
    """
    # Parse date from folder structure or filename
    try:
        yr = int(file_path.parent.parent.name)
        mo = int(file_path.parent.name)
        fallback_dt = pd.Period(f"{yr}-{mo:02d}", freq="M").end_time.date()
    except Exception:
        fallback_dt = parse_date_robust(file_path.name) or pd.Period("2024-10", freq="M").end_time.date()

    file_rows = []

    if file_path.suffix.lower() == ".xlsx":
        wb = openpyxl.load_workbook(file_path, data_only=True)
        for sname in wb.sheetnames:
            code_norm = sname.strip().upper()
            if code_norm in TARGET_CONFIG:
                fund_code, fund_canonical = TARGET_CONFIG[code_norm]
                rows = parse_xlsx_sheet(wb[sname], fund_canonical, fallback_dt)
                file_rows.extend(rows)
        wb.close()
    elif file_path.suffix.lower() == ".xls":
        wb = xlrd.open_workbook(file_path)
        for sname in wb.sheet_names():
            code_norm = sname.strip().upper()
            if code_norm in TARGET_CONFIG:
                fund_code, fund_canonical = TARGET_CONFIG[code_norm]
                rows = parse_xls_sheet(wb.sheet_by_name(sname), fund_canonical, fallback_dt)
                file_rows.extend(rows)

    return pd.DataFrame(file_rows)


# ==============================================================================
# MAIN PIPELINE
# ==============================================================================

def main():
    print("=" * 80)
    print("GROWW MUTUAL FUND - HOLDINGS CLEANING PIPELINE")
    print(f"Raw Input Folder : {RAW_FOLDER}")
    print(f"Output File      : {OUTPUT_FILE}")
    print("=" * 80)

    raw_files = sorted(
        [
            f for f in RAW_FOLDER.rglob("*.*")
            if f.suffix.lower() in [".xlsx", ".xls"]
            and not f.name.startswith("~$")
        ]
    )

    print(f"\nFound {len(raw_files)} raw files to process.\n")

    frames = []
    error_count = 0

    for idx, fpath in enumerate(raw_files, start=1):
        try:
            df_file = parse_groww_file(fpath)
            if not df_file.empty:
                frames.append(df_file)
                print(f"[{idx:02d}/{len(raw_files):02d}] {fpath.name[:45]:45} -> {len(df_file):4d} rows")
            else:
                print(f"[{idx:02d}/{len(raw_files):02d}] {fpath.name[:45]:45} -> 0 rows (WARNING)")
        except Exception as e:
            print(f"[{idx:02d}/{len(raw_files):02d}] {fpath.name[:45]:45} -> ERROR: {e}")
            error_count += 1

    if not frames:
        print("\nERROR: No data extracted from any file!")
        return

    df_master = pd.concat(frames, ignore_index=True)

    # Reorder columns to standard format
    cols_order = [
        "AMC",
        "Fund_Name",
        "Portfolio_Date",
        "Month",
        "Security_Name",
        "ISIN",
        "Industry_Rating",
        "Quantity",
    ]
    df_master = df_master[cols_order]

    # Save to Excel
    print("\n" + "-" * 80)
    print("Writing cleaned dataset...")
    df_master.to_excel(OUTPUT_FILE, index=False)

    print("\n" + "=" * 80)
    print("CLEANING COMPLETED SUCCESSFULLY")
    print(f"Output File      : {OUTPUT_FILE}")
    print(f"Total Rows       : {len(df_master):,}")
    print(f"Total Funds      : {df_master['Fund_Name'].nunique()}")
    print(f"Total Months     : {df_master['Portfolio_Date'].nunique()}")
    print(f"Unique ISINs     : {df_master['ISIN'].nunique()}")
    print(f"Processing Errors: {error_count}")
    print("=" * 80)


if __name__ == "__main__":
    main()
