import re
import sys
from datetime import datetime
from pathlib import Path
import openpyxl
import pandas as pd

# ==============================================================================
# SETTINGS & PATHS
# ==============================================================================

PROJECT_ROOT = Path(__file__).resolve().parents[2]
RAW_FOLDER = PROJECT_ROOT / "01_raw_files" / "HSBC"
OUTPUT_FOLDER = PROJECT_ROOT / "03_clean_data" / "HSBC"
OUTPUT_FOLDER.mkdir(parents=True, exist_ok=True)
OUTPUT_FILE = OUTPUT_FOLDER / "HSBC_All_Funds_Cleaned.xlsx"

AMC_NAME = "HSBC Mutual Fund"
START_DATE = pd.Timestamp(year=2024, month=10, day=1)
END_DATE = pd.Timestamp(year=2026, month=8, day=31)

STANDARD_COLUMNS = [
    "AMC",
    "Fund_Name",
    "Portfolio_Date",
    "Month",
    "Security_Name",
    "ISIN",
    "Industry_Rating",
    "Quantity",
]

# Canonical 14 Target Funds
CANONICAL_FUNDS = [
    "HSBC Aggressive Hybrid Fund",
    "HSBC Large Cap Fund",
    "HSBC Large & Mid Cap Fund",
    "HSBC Midcap Fund",
    "HSBC Flexi Cap Fund",
    "HSBC Multi Cap Fund",
    "HSBC Small Cap Fund",
    "HSBC Infrastructure Fund",
    "HSBC Value Fund",
    "HSBC Business Cycles Fund",
    "HSBC ELSS Tax saver Fund",
    "HSBC Consumption Fund",
    "HSBC Balanced Advantage Fund",
    "HSBC Financial Services Fund",
]

MONTH_NAME_MAP = {
    "jan": 1, "january": 1,
    "feb": 2, "february": 2,
    "mar": 3, "march": 3,
    "apr": 4, "april": 4,
    "may": 5,
    "jun": 6, "june": 6,
    "jul": 7, "july": 7,
    "aug": 8, "august": 8,
    "sep": 9, "september": 9,
    "oct": 10, "october": 10,
    "nov": 11, "november": 11,
    "dec": 12, "december": 12,
}


# ==============================================================================
# HELPER FUNCTIONS
# ==============================================================================

def is_valid_isin(value):
    """
    Strict equity ISIN: starts with INE, length 12.
    """
    if pd.isna(value):
        return False
    val = str(value).strip().upper()
    return val.startswith("INE") and len(val) == 12


def match_canonical_fund(text):
    """
    Maps fund name / file stem to canonical fund name.
    """
    combined = text.lower()

    if "aggressive" in combined and "hybrid" in combined:
        return "HSBC Aggressive Hybrid Fund"
    if "balanced" in combined and "advantage" in combined:
        return "HSBC Balanced Advantage Fund"
    if "large" in combined and "mid" in combined:
        return "HSBC Large & Mid Cap Fund"
    if "large" in combined and "cap" in combined and "mid" not in combined:
        return "HSBC Large Cap Fund"
    if ("midcap" in combined or "mid cap" in combined or "mid-cap" in combined) and "large" not in combined:
        return "HSBC Midcap Fund"
    if "flexi" in combined:
        return "HSBC Flexi Cap Fund"
    if "multi" in combined and "cap" in combined:
        return "HSBC Multi Cap Fund"
    if "small" in combined and "cap" in combined:
        return "HSBC Small Cap Fund"
    if "infrastructure" in combined:
        return "HSBC Infrastructure Fund"
    if "value" in combined:
        return "HSBC Value Fund"
    if "business" in combined:
        return "HSBC Business Cycles Fund"
    if "elss" in combined or "tax saver" in combined or "tax-saver" in combined or "tax saver equity" in combined:
        return "HSBC ELSS Tax saver Fund"
    if "consumption" in combined:
        return "HSBC Consumption Fund"
    if "financial" in combined and "services" in combined:
        return "HSBC Financial Services Fund"

    return None


def parse_date_from_workbook(ws, file_path):
    """
    Extracts portfolio date from sheet header rows or fallback to filename.
    """
    title_text = ""
    for r in range(1, 7):
        val = ws.cell(r, 1).value
        if val:
            title_text += "\n" + str(val)

    # 1. Regex for "as of Month DD, YYYY" or "as on Month DD, YYYY"
    m_date = re.search(r"as (?:of|on)\s+([A-Za-z]+)\s+(\d{1,2}),?\s+(\d{4})", title_text, re.I)
    if m_date:
        mon_str, day_str, yr_str = m_date.group(1).lower(), m_date.group(2), m_date.group(3)
        if mon_str in MONTH_NAME_MAP:
            return pd.Timestamp(year=int(yr_str), month=MONTH_NAME_MAP[mon_str], day=int(day_str))

    # 2. Regex for "as of DD Month YYYY" or "as on DD Month YYYY"
    m_date2 = re.search(r"as (?:of|on)\s+(\d{1,2})\s+([A-Za-z]+)\s+(\d{4})", title_text, re.I)
    if m_date2:
        day_str, mon_str, yr_str = m_date2.group(1), m_date2.group(2).lower(), m_date2.group(3)
        if mon_str in MONTH_NAME_MAP:
            return pd.Timestamp(year=int(yr_str), month=MONTH_NAME_MAP[mon_str], day=int(day_str))

    # 3. Fallback to filename
    file_stem = file_path.stem.lower()
    m_file = re.search(r"(\d{1,2})-([a-z]+)-(\d{4})", file_stem)
    if m_file:
        day_str, mon_str, yr_str = m_file.group(1), m_file.group(2).lower(), m_file.group(3)
        if mon_str in MONTH_NAME_MAP:
            return pd.Timestamp(year=int(yr_str), month=MONTH_NAME_MAP[mon_str], day=int(day_str))

    # 4. Fallback to parent directory YYYY/MM
    try:
        y = int(file_path.parent.parent.name)
        m = int(file_path.parent.name)
        # Last day of month
        dt = pd.Timestamp(year=y, month=m, day=1) + pd.offsets.MonthEnd(1)
        return dt
    except Exception:
        pass

    raise ValueError(f"Could not determine portfolio date for {file_path}")


STOP_TRIGGERS = [
    "total",
    "debt instruments",
    "fixed rates bonds - corporate",
    "fixed rate bonds - corporate",
    "fixed rates bonds",
    "fixed rate bonds",
    "securitised debt",
    "government securities",
    "money market instruments",
    "commercial papers",
    "certificate of deposit",
    "treps",
    "net current assets",
    "total net assets",
]


def parse_hsbc_workbook(file_path):
    """
    Parses an individual HSBC workbook and extracts equity holdings.
    """
    wb = openpyxl.load_workbook(file_path, data_only=True)
    ws = wb.worksheets[0]

    # Portfolio Date
    portfolio_date = parse_date_from_workbook(ws, file_path)
    month_date = pd.Timestamp(year=portfolio_date.year, month=portfolio_date.month, day=1)

    # Fund Name
    fund_name = match_canonical_fund(file_path.stem)
    if not fund_name:
        # Check sheet or header
        fund_name = match_canonical_fund(str(ws.cell(1, 1).value or ""))
    if not fund_name:
        wb.close()
        return pd.DataFrame()

    # Find header row
    header_row_idx = None
    col_map = {}
    for r in range(1, 15):
        row_vals = [str(ws.cell(r, c).value or "").strip().lower() for c in range(1, ws.max_column + 1)]
        if any("isin" in v for v in row_vals):
            header_row_idx = r
            for c_idx, val in enumerate(row_vals, start=1):
                if "name of the instrument" in val or "security name" in val:
                    col_map["security_name"] = c_idx
                elif "isin" in val:
                    col_map["isin"] = c_idx
                elif "industry" in val or "rating" in val:
                    col_map["industry"] = c_idx
                elif "quantity" in val or "qty" in val:
                    col_map["quantity"] = c_idx
            break

    if not header_row_idx:
        wb.close()
        return pd.DataFrame()

    col_sec = col_map.get("security_name", 1)
    col_isin = col_map.get("isin", 2)
    col_ind = col_map.get("industry", 3)
    col_qty = col_map.get("quantity", 4)

    rows = []
    equity_section_started = False

    for r in range(header_row_idx + 1, ws.max_row + 1):
        name_val = ws.cell(r, col_sec).value
        isin_val = ws.cell(r, col_isin).value
        ind_val = ws.cell(r, col_ind).value
        qty_val = ws.cell(r, col_qty).value

        name_str = str(name_val or "").strip()
        isin_str = str(isin_val or "").strip().upper()
        ind_str = str(ind_val or "").strip()
        name_lower = name_str.lower()

        # Skip preamble headers before first equity row
        if not equity_section_started:
            if is_valid_isin(isin_str):
                equity_section_started = True
            else:
                continue

        # Stop conditions as soon as Total, Debt Instruments, Fixed rates bonds, or subsequent Listed / Awaiting listing is reached
        if any(trigger in name_lower for trigger in STOP_TRIGGERS):
            break

        if "listed / awaiting listing on stock exchanges" in name_lower:
            break

        if is_valid_isin(isin_str):
            try:
                qty_num = float(qty_val) if qty_val is not None else None
            except Exception:
                qty_num = None

            if qty_num is not None and qty_num > 0:
                rows.append({
                    "AMC": AMC_NAME,
                    "Fund_Name": fund_name,
                    "Portfolio_Date": portfolio_date.strftime("%Y-%m-%d"),
                    "Month": month_date.strftime("%Y-%m-%d"),
                    "Security_Name": name_str,
                    "ISIN": isin_str,
                    "Industry_Rating": ind_str,
                    "Quantity": int(round(qty_num)),
                })

    wb.close()
    return pd.DataFrame(rows)


# ==============================================================================
# MAIN CLEANING PIPELINE
# ==============================================================================

def main():
    print("\n" + "=" * 90)
    print("CLEANING HSBC MUTUAL FUND PORTFOLIOS")
    print("=" * 90)

    raw_files = sorted(list(RAW_FOLDER.rglob("*.xlsx")) + list(RAW_FOLDER.rglob("*.xls")))
    print(f"Found {len(raw_files)} raw files in {RAW_FOLDER}")

    if not raw_files:
        print("No raw files found. Run download_hsbc_monthly_files.py first.")
        return

    frames = []
    processed_count = 0
    error_count = 0

    for idx, fpath in enumerate(raw_files, start=1):
        if fpath.name.startswith("~$"):
            continue
        try:
            df_file = parse_hsbc_workbook(fpath)
            if not df_file.empty:
                frames.append(df_file)
                processed_count += 1
                if idx % 25 == 0 or idx == len(raw_files):
                    print(f"  Processed {idx}/{len(raw_files)} files... (Rows so far: {sum(len(f) for f in frames):,})")
            else:
                print(f"  [WARN] No equity data extracted from {fpath.name}")
        except Exception as e:
            error_count += 1
            print(f"  [ERROR] Failed to process {fpath.name}: {e}")

    if not frames:
        print("No data extracted. Exiting.")
        return

    combined_df = pd.concat(frames, ignore_index=True)

    # Filter date range
    combined_df["Month_DT"] = pd.to_datetime(combined_df["Month"])
    combined_df = combined_df[(combined_df["Month_DT"] >= START_DATE) & (combined_df["Month_DT"] <= END_DATE)].copy()
    combined_df.drop(columns=["Month_DT"], inplace=True)

    # Deduplicate in case of duplicate file runs
    combined_df = combined_df.drop_duplicates(
        subset=["AMC", "Fund_Name", "Portfolio_Date", "ISIN"]
    )

    # Sort
    combined_df = combined_df.sort_values(
        by=["Fund_Name", "Portfolio_Date", "Security_Name"]
    ).reset_index(drop=True)

    # Ensure standard columns
    combined_df = combined_df[STANDARD_COLUMNS]

    # Save to Excel
    print(f"\nSaving cleaned data to: {OUTPUT_FILE}")
    combined_df.to_excel(OUTPUT_FILE, index=False)

    print("\n" + "=" * 90)
    print("HSBC CLEANING SUMMARY")
    print("=" * 90)
    print(f"Total Files Processed: {processed_count}")
    print(f"Errors Encountered:    {error_count}")
    print(f"Total Clean Rows:      {len(combined_df):,}")
    print(f"Unique Funds:          {combined_df['Fund_Name'].nunique()}")
    print(f"Unique Months:         {combined_df['Month'].nunique()}")
    print(f"Date Range:            {combined_df['Month'].min()} to {combined_df['Month'].max()}")
    print(f"Unique ISINs:          {combined_df['ISIN'].nunique()}")

    print("\nRows per Month:")
    month_counts = combined_df.groupby("Month").size()
    for m, c in month_counts.items():
        print(f"  {m}: {c:,} rows")

    print("\nFunds Covered:")
    for fn in sorted(combined_df["Fund_Name"].unique()):
        cnt = combined_df[combined_df["Fund_Name"] == fn]["Month"].nunique()
        print(f"  - {fn:<35}: {cnt} months")


if __name__ == "__main__":
    main()
