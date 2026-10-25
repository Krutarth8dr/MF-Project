"""
Cleaning script for BNP MF (Baroda BNP Paribas Mutual Fund).
Extracts listed equity holdings from consolidated monthly files for 23 target funds:
- T0ME02: Baroda BNP Paribas Mid Cap Fund
- T0ME04: Baroda BNP Paribas Large Cap Fund
- T0ME05: Baroda BNP Paribas ELSS Tax Saver Fund
- T0ME08: Baroda BNP Paribas Conservative Hybrid Fund
- T0ME19: Baroda BNP Paribas Aggressive Hybrid Fund
- T0ME20: Baroda BNP Paribas Focused Fund
- T0ME21: Baroda BNP Paribas India Consumption Fund
- T0ME25: Baroda BNP Paribas Flexi Cap Fund
- T0ME26: Baroda BNP Paribas Multi Asset Fund
- T0ME30: Baroda BNP Paribas Value Fund
- T0ME32: Baroda BNP Paribas Small Cap Fund
- T0ME34: Baroda BNP Paribas Innovation Fund
- T0ME37: Baroda BNP Paribas Manufacturing Fund
- T0ME42: Baroda BNP Paribas Energy Opportunities Fund
- T0ME45: Baroda BNP Paribas Health and Wellness Fund
- T0ME47: Baroda BNP Paribas Business Conglomerates Fund
- T0ME49: Baroda BNP Paribas Services Fund
- YR04: Baroda BNP Paribas Multi Cap Fund
- YR29: Baroda BNP Paribas Banking & Financial Services Fund
- YR48: Baroda BNP Paribas Balanced Advantage Fund
- YR53: Baroda BNP Paribas Equity Savings Fund
- YR54: Baroda BNP Paribas Large & Mid Cap Fund
- YR56: Baroda BNP Paribas Business Cycle Fund

Stops cleaning as soon as "Sub Total", "(b)Unlisted", or "Total" are detected
in the "Name of the Instrument" column.
Output:
03_clean_data/BNP_MF/BNP_MF_All_Funds_Cleaned.xlsx
"""

import datetime
import glob
import io
import os
import re
import sys
from pathlib import Path

import openpyxl
import pandas as pd
import xlrd

# ==============================================================================
# SETTINGS & PATHS
# ==============================================================================

PROJECT_ROOT = Path(__file__).resolve().parents[2]
RAW_FOLDER = PROJECT_ROOT / "01_raw_files" / "BNP_MF"
OUTPUT_FOLDER = PROJECT_ROOT / "03_clean_data" / "BNP_MF"
OUTPUT_FOLDER.mkdir(parents=True, exist_ok=True)
OUTPUT_FILE = OUTPUT_FOLDER / "BNP_MF_All_Funds_Cleaned.xlsx"

AMC_NAME = "BNP MF"

STANDARD_COLUMNS = [
    "AMC",
    "Fund_Name",
    "Portfolio_Date",
    "Month",
    "Security_Name",
    "ISIN",
    "Industry_Rating",
    "Quantity",
]

# Target fund mapping taken from September 2026 sheet
TARGET_FUNDS = {
    "T0ME02": "Baroda BNP Paribas Mid Cap Fund",
    "T0ME04": "Baroda BNP Paribas Large Cap Fund",
    "T0ME05": "Baroda BNP Paribas ELSS Tax Saver Fund",
    "T0ME08": "Baroda BNP Paribas Conservative Hybrid Fund",
    "T0ME19": "Baroda BNP Paribas Aggressive Hybrid Fund",
    "T0ME20": "Baroda BNP Paribas Focused Fund",
    "T0ME21": "Baroda BNP Paribas India Consumption Fund",
    "T0ME25": "Baroda BNP Paribas Flexi Cap Fund",
    "T0ME26": "Baroda BNP Paribas Multi Asset Fund",
    "T0ME30": "Baroda BNP Paribas Value Fund",
    "T0ME32": "Baroda BNP Paribas Small Cap Fund",
    "T0ME34": "Baroda BNP Paribas Innovation Fund",
    "T0ME37": "Baroda BNP Paribas Manufacturing Fund",
    "T0ME42": "Baroda BNP Paribas Energy Opportunities Fund",
    "T0ME45": "Baroda BNP Paribas Health and Wellness Fund",
    "T0ME47": "Baroda BNP Paribas Business Conglomerates Fund",
    "T0ME49": "Baroda BNP Paribas Services Fund",
    "YR04": "Baroda BNP Paribas Multi Cap Fund",
    "YR29": "Baroda BNP Paribas Banking & Financial Services Fund",
    "YR48": "Baroda BNP Paribas Balanced Advantage Fund",
    "YR53": "Baroda BNP Paribas Equity Savings Fund",
    "YR54": "Baroda BNP Paribas Large & Mid Cap Fund",
    "YR56": "Baroda BNP Paribas Business Cycle Fund",
}


def clean_text(val):
    if val is None:
        return ""
    text = str(val).strip()
    return re.sub(r"\s+", " ", text)


def is_stop_instrument(name_val):
    """
    User instruction:
    'stop cleaning as soon as you detect "Sub Total", "(b)Unlisted" "Total" cells containing these things in the "Name of the Instrument" column.'
    """
    if not name_val:
        return False
    lower = clean_text(name_val).lower()
    if not lower:
        return False
    
    # Specific substring checks requested by user
    if any(m in lower for m in [
        "sub total",
        "subtotal",
        "sub-total",
        "sub_total",
        "(b)unlisted",
        "(b) unlisted",
        "b) unlisted",
        "b)unlisted",
        "(b)",
        "grand total",
    ]):
        return True
    
    # Total exact or startswith check (to avoid matching stocks like 'Adani Total Gas Limited')
    if lower in ["total", "total:", "total (a)", "total(a)", "total (b)", "total(b)", "unlisted"]:
        return True
    if lower.startswith(("total (", "total(", "total -", "total:", "grand total")):
        return True
        
    # Non-equity asset classes
    if any(t in lower for t in [
        "debt instruments",
        "money market",
        "treasury bill",
        "reverse repo",
        "treps",
        "net current assets",
        "total net assets",
        "derivatives",
    ]):
        return True

    return False


def parse_date_from_text(text):
    if not text:
        return None, None
    m = re.search(r"as on\s+([A-Za-z]+)\s+(\d{1,2}),?\s*(\d{4})", str(text), re.I)
    if m:
        month_name, day, year = m.group(1), m.group(2), m.group(3)
        for fmt in ("%B %d %Y", "%b %d %Y"):
            try:
                dt = datetime.datetime.strptime(f"{month_name} {day} {year}", fmt)
                return dt.strftime("%Y-%m-%d"), dt.strftime("%b-%Y")
            except Exception:
                continue
    return None, None


def clean_quantity(val):
    if val is None:
        return None
    if isinstance(val, (int, float)):
        return float(val) if val > 0 else None
    s = str(val).strip().replace(",", "")
    try:
        f = float(s)
        return f if f > 0 else None
    except ValueError:
        return None


def clean_openpyxl_sheet(ws, fund_code, fund_name, default_date, default_month):
    rows = list(ws.iter_rows(values_only=True))
    if not rows:
        return []
    
    # 1. Detect portfolio date from top 6 rows
    portfolio_date, month_str = None, None
    for r in rows[:6]:
        for cell in r:
            if cell and isinstance(cell, str) and "as on" in cell.lower():
                p_date, p_month = parse_date_from_text(cell)
                if p_date:
                    portfolio_date, month_str = p_date, p_month
                    break
        if portfolio_date:
            break
            
    if not portfolio_date:
        portfolio_date, month_str = default_date, default_month

    # 2. Find header row and column indices
    inst_col, isin_col, ind_col, qty_col = None, None, None, None
    header_idx = None
    for r_idx, r in enumerate(rows[:10]):
        r_str = [str(x or "").strip().lower() for x in r]
        for c_idx, cell in enumerate(r_str):
            if "name of" in cell or "instrument" in cell:
                inst_col = c_idx
            elif "isin" in cell:
                isin_col = c_idx
            elif "industry" in cell or "rating" in cell:
                ind_col = c_idx
            elif "quantity" in cell or "qty" in cell:
                qty_col = c_idx
        if inst_col is not None and isin_col is not None and qty_col is not None:
            header_idx = r_idx
            break

    if header_idx is None:
        return []

    # 3. Iterate data rows until stop condition
    extracted = []
    for r in rows[header_idx + 1:]:
        inst_val = r[inst_col] if inst_col < len(r) else None
        
        # Check stop condition
        if is_stop_instrument(inst_val):
            break
            
        isin_val = clean_text(r[isin_col]) if isin_col < len(r) and r[isin_col] is not None else ""
        if not (isin_val.startswith("INE") and len(isin_val) == 12):
            continue
            
        qty_val = clean_quantity(r[qty_col]) if qty_col < len(r) else None
        if qty_val is None or qty_val <= 0:
            continue
            
        sec_name = clean_text(inst_val)
        ind_val = clean_text(r[ind_col]) if ind_col is not None and ind_col < len(r) and r[ind_col] is not None else ""
        
        extracted.append({
            "AMC": AMC_NAME,
            "Fund_Name": fund_name,
            "Portfolio_Date": portfolio_date,
            "Month": month_str,
            "Security_Name": sec_name,
            "ISIN": isin_val,
            "Industry_Rating": ind_val,
            "Quantity": qty_val,
        })
        
    return extracted


def clean_xlrd_sheet(ws, fund_code, fund_name, default_date, default_month):
    if ws.nrows == 0:
        return []

    # 1. Detect portfolio date from top 6 rows
    portfolio_date, month_str = None, None
    for r in range(min(6, ws.nrows)):
        for c in range(ws.ncols):
            val = str(ws.cell_value(r, c) or "")
            if "as on" in val.lower():
                p_date, p_month = parse_date_from_text(val)
                if p_date:
                    portfolio_date, month_str = p_date, p_month
                    break
        if portfolio_date:
            break
            
    if not portfolio_date:
        portfolio_date, month_str = default_date, default_month

    # 2. Find header row and column indices
    inst_col, isin_col, ind_col, qty_col = None, None, None, None
    header_idx = None
    for r in range(min(10, ws.nrows)):
        r_str = [str(ws.cell_value(r, c) or "").strip().lower() for c in range(ws.ncols)]
        for c_idx, cell in enumerate(r_str):
            if "name of" in cell or "instrument" in cell:
                inst_col = c_idx
            elif "isin" in cell:
                isin_col = c_idx
            elif "industry" in cell or "rating" in cell:
                ind_col = c_idx
            elif "quantity" in cell or "qty" in cell:
                qty_col = c_idx
        if inst_col is not None and isin_col is not None and qty_col is not None:
            header_idx = r
            break

    if header_idx is None:
        return []

    # 3. Iterate data rows until stop condition
    extracted = []
    for r in range(header_idx + 1, ws.nrows):
        inst_val = ws.cell_value(r, inst_col) if inst_col < ws.ncols else None
        
        # Check stop condition
        if is_stop_instrument(inst_val):
            break
            
        isin_val = clean_text(ws.cell_value(r, isin_col)) if isin_col < ws.ncols else ""
        if not (isin_val.startswith("INE") and len(isin_val) == 12):
            continue
            
        qty_val = clean_quantity(ws.cell_value(r, qty_col)) if qty_col < ws.ncols else None
        if qty_val is None or qty_val <= 0:
            continue
            
        sec_name = clean_text(inst_val)
        ind_val = clean_text(ws.cell_value(r, ind_col)) if ind_col is not None and ind_col < ws.ncols else ""
        
        extracted.append({
            "AMC": AMC_NAME,
            "Fund_Name": fund_name,
            "Portfolio_Date": portfolio_date,
            "Month": month_str,
            "Security_Name": sec_name,
            "ISIN": isin_val,
            "Industry_Rating": ind_val,
            "Quantity": qty_val,
        })
        
    return extracted


def process_all_files():
    files = sorted(glob.glob(str(RAW_FOLDER / "*" / "*" / "*.*")))
    print(f"Found {len(files)} files in {RAW_FOLDER}")
    
    all_data = []

    for fpath in files:
        file_path = Path(fpath)
        # Determine fallback date from directory structure: e.g. 2026/09
        try:
            year_dir = file_path.parent.parent.name
            month_dir = file_path.parent.name
            default_date = f"{year_dir}-{month_dir}-28"
            dt_fallback = datetime.datetime.strptime(f"{year_dir}-{month_dir}", "%Y-%m")
            default_month = dt_fallback.strftime("%b-%Y")
        except Exception:
            default_date = "2026-09-30"
            default_month = "Sep-2026"

        with open(file_path, "rb") as fp:
            sig = fp.read(8)

        is_xlsx = sig.startswith(b"PK\x03\x04")
        is_ole = sig.startswith(b"\xd0\xcf\x11\xe0")

        file_rows = []
        if is_xlsx:
            with open(file_path, "rb") as fp:
                wb = openpyxl.load_workbook(io.BytesIO(fp.read()), data_only=True, read_only=True)
            for code, fname in TARGET_FUNDS.items():
                if code in wb.sheetnames:
                    ws = wb[code]
                    rows = clean_openpyxl_sheet(ws, code, fname, default_date, default_month)
                    file_rows.extend(rows)
            wb.close()
        elif is_ole:
            wb = xlrd.open_workbook(file_path)
            for code, fname in TARGET_FUNDS.items():
                if code in wb.sheet_names():
                    ws = wb.sheet_by_name(code)
                    rows = clean_xlrd_sheet(ws, code, fname, default_date, default_month)
                    file_rows.extend(rows)
        else:
            print(f"[SKIP] Unknown file signature: {file_path}")
            continue

        date_sample = file_rows[0]["Portfolio_Date"] if file_rows else default_date
        print(f"[{file_path.parent.name}/{file_path.name[:35]}] -> {len(file_rows):>5} holdings (Date: {date_sample})")
        all_data.extend(file_rows)

    df = pd.DataFrame(all_data, columns=STANDARD_COLUMNS)
    print(f"\nTotal extracted holdings: {len(df):,}")
    print("Holdings by month:")
    print(df.groupby("Month")["Fund_Name"].count())
    print("\nHoldings by Fund:")
    print(df.groupby("Fund_Name")["ISIN"].count())

    # Sort and save
    df.sort_values(by=["Fund_Name", "Portfolio_Date", "Security_Name"], inplace=True)
    df.to_excel(OUTPUT_FILE, index=False)
    print(f"\n[SUCCESS] Cleaned data saved to: {OUTPUT_FILE}")
    print(f"File size: {OUTPUT_FILE.stat().st_size:,} bytes")
    return df


if __name__ == "__main__":
    process_all_files()
