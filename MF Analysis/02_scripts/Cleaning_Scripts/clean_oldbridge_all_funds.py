"""
Clean all downloaded monthly portfolio files for Old Bridge Mutual Fund (Old Bridge MF).

Process:
    - Reads all files in 01_raw_files/Old_Bridge_MF/<YYYY>/<MM>/
    - Processes target schemes:
        1. Old Bridge Focused Fund (24 months: Oct 2024 to Sep 2026)
        2. Old Bridge Flexi Cap Fund (7 months: Mar 2026 to Sep 2026)
    - Dynamic header detection for Name, ISIN, Industry/Rating, and Quantity
    - Stops immediately upon encountering 'Sub Total', '(b)Unlisted', '(b) Unlisted', or 'Total'
    - Filters for valid listed Indian equities (ISIN starts with INE, length 12, quantity > 0)
    - Generates 03_clean_data/Old_Bridge_MF/Old_Bridge_MF_All_Funds_Cleaned.xlsx
"""

import sys
import os
import calendar
import warnings
from datetime import datetime
from pathlib import Path
import openpyxl
import pandas as pd

warnings.filterwarnings("ignore", category=UserWarning, module="openpyxl")

# Force UTF-8 output
try:
    sys.stdout.reconfigure(encoding="utf-8")
    sys.stderr.reconfigure(encoding="utf-8")
except (AttributeError, ValueError):
    pass

PROJECT_ROOT = Path(__file__).resolve().parents[2]
RAW_DIR = PROJECT_ROOT / "01_raw_files" / "Old_Bridge_MF"
CLEAN_DIR = PROJECT_ROOT / "03_clean_data" / "Old_Bridge_MF"
CLEAN_DIR.mkdir(parents=True, exist_ok=True)

OUTPUT_FILE = CLEAN_DIR / "Old_Bridge_MF_All_Funds_Cleaned.xlsx"

STOP_TERMS = ["sub total", "(b)unlisted", "(b) unlisted", "total", "subtotal"]

STANDARD_COLUMNS = [
    "AMC",
    "Fund_Name",
    "Portfolio_Date",
    "Month",
    "Security_Name",
    "ISIN",
    "Industry_Rating",
    "Quantity",
]


def clean_sheet(ws, fund_name, portfolio_date, month_str):
    """Parses a single worksheet of an Old Bridge portfolio workbook."""
    holdings = []

    # Locate header row and column mapping
    header_row = None
    col_map = {}
    for r in range(1, 16):
        row_vals = [str(ws.cell(row=r, column=c).value or "").strip() for c in range(1, min(15, ws.max_column + 1))]
        lower_vals = [v.lower() for v in row_vals]
        if any("isin" in v for v in lower_vals) and any("instrument" in v or "name" in v for v in lower_vals):
            header_row = r
            for c_idx, val in enumerate(lower_vals, 1):
                if "isin" in val and "isin" not in col_map:
                    col_map["isin"] = c_idx
                elif ("instrument" in val or "name" in val) and "yield" not in val and "name" not in col_map:
                    col_map["name"] = c_idx
                elif ("rating" in val or "industry" in val) and "industry" not in col_map:
                    col_map["industry"] = c_idx
                elif ("quantity" in val or "qty" in val) and "qty" not in col_map:
                    col_map["qty"] = c_idx
            break

    if not header_row or "isin" not in col_map or "name" not in col_map or "qty" not in col_map:
        return []

    for r in range(header_row + 1, ws.max_row + 1):
        name_val = str(ws.cell(row=r, column=col_map["name"]).value or "").strip()
        isin_val = str(ws.cell(row=r, column=col_map["isin"]).value or "").strip()

        name_lower = name_val.lower()
        isin_lower = isin_val.lower()
        combined = f"{name_lower} {isin_lower}"

        # Immediate stop condition
        if any(term in combined for term in STOP_TERMS):
            break

        # Filter for listed Indian equity
        if not (isin_val.startswith("INE") and len(isin_val) == 12):
            continue

        raw_qty = ws.cell(row=r, column=col_map["qty"]).value
        try:
            qty = float(raw_qty)
            if qty <= 0:
                continue
        except (ValueError, TypeError):
            continue

        industry_val = ""
        if "industry" in col_map:
            industry_val = str(ws.cell(row=r, column=col_map["industry"]).value or "").strip()

        holdings.append({
            "AMC": "Old Bridge MF",
            "Fund_Name": fund_name,
            "Portfolio_Date": portfolio_date,
            "Month": month_str,
            "Security_Name": name_val,
            "ISIN": isin_val,
            "Industry_Rating": industry_val,
            "Quantity": qty,
        })

    return holdings


def process_all_files():
    """Iterate over all raw files in Old_Bridge_MF folder, clean and save to single Excel file."""
    print("=" * 80)
    print("Starting Old Bridge MF Portfolio Cleaner")
    print(f"Reading from: {RAW_DIR}")
    print(f"Output target: {OUTPUT_FILE}")
    print("=" * 80)

    all_records = []
    file_summary = []

    # Traverse YYYY / MM
    for year_dir in sorted(RAW_DIR.iterdir()):
        if not year_dir.is_dir() or not year_dir.name.isdigit():
            continue
        year = int(year_dir.name)

        for month_dir in sorted(year_dir.iterdir()):
            if not month_dir.is_dir() or not month_dir.name.isdigit():
                continue
            month = int(month_dir.name)

            last_day = calendar.monthrange(year, month)[1]
            portfolio_date = datetime(year, month, last_day)
            month_str = f"{year}-{month:02d}-01"

            for file_path in sorted(month_dir.glob("*.xls*")):
                file_name = file_path.name
                
                # Determine fund name from prefix or filename
                if file_name.startswith("FlexiCap_"):
                    fund_name = "Old Bridge Flexi Cap Fund"
                elif file_name.startswith("Focused_"):
                    fund_name = "Old Bridge Focused Fund"
                else:
                    # Fallback check
                    fn_lower = file_name.lower()
                    if "flexi" in fn_lower or "obfx" in fn_lower or "obflx" in fn_lower:
                        fund_name = "Old Bridge Flexi Cap Fund"
                    else:
                        fund_name = "Old Bridge Focused Fund"

                try:
                    wb = openpyxl.load_workbook(file_path, data_only=True)
                    sheet_name = wb.sheetnames[0]
                    ws = wb[sheet_name]
                    holdings = clean_sheet(ws, fund_name, portfolio_date, month_str)

                    all_records.extend(holdings)
                    file_summary.append({
                        "file": file_name,
                        "year_month": f"{year}-{month:02d}",
                        "fund": fund_name,
                        "records": len(holdings),
                        "status": "SUCCESS" if len(holdings) > 0 else "ZERO_RECORDS"
                    })
                    print(f"Processed {year}-{month:02d} | {fund_name:28} | {len(holdings):3d} records | {file_name}")
                except Exception as e:
                    file_summary.append({
                        "file": file_name,
                        "year_month": f"{year}-{month:02d}",
                        "fund": fund_name,
                        "records": 0,
                        "status": f"ERROR: {str(e)}"
                    })
                    print(f"ERROR processing {file_name}: {e}")

    df = pd.DataFrame(all_records, columns=STANDARD_COLUMNS)

    print("\n" + "=" * 80)
    print("CLEANING COMPLETED - SUMMARY")
    print("=" * 80)
    print(f"Total Rows Extracted: {len(df):,}")
    print(f"Unique AMCs: {df['AMC'].nunique()} -> {list(df['AMC'].unique())}")
    print(f"Unique Funds ({df['Fund_Name'].nunique()}):")
    for f in df["Fund_Name"].unique():
        f_df = df[df["Fund_Name"] == f]
        print(f"  - {f:30}: {len(f_df):5,d} rows, {f_df['Month'].nunique():2d} months ({f_df['Month'].min()} to {f_df['Month'].max()})")
    print(f"Unique ISINs: {df['ISIN'].nunique():,}")
    print(f"Date Range: {df['Month'].min()} to {df['Month'].max()}")
    print("=" * 80)

    # Save to Excel
    print(f"Writing clean dataset to {OUTPUT_FILE}...")
    df.to_excel(OUTPUT_FILE, index=False, engine="openpyxl")
    file_size_kb = OUTPUT_FILE.stat().st_size / 1024
    print(f"Successfully saved {OUTPUT_FILE.name} ({file_size_kb:.1f} KB)")


if __name__ == "__main__":
    process_all_files()
