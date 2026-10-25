"""
Cleaning script for AlphaGrep MF.
Extracts listed equity holdings from monthly portfolio workbooks for:
- AlphaGrep Flexi Cap Fund
- AlphaGrep Multi Asset Allocation Fund

Stops cleaning as soon as "Sub Total", "(b) Unlisted", "Total", etc. are detected.
Output:
03_clean_data/ALPHAGREP_MF/AlphaGrep_MF_All_Funds_Cleaned.xlsx
"""

import os
import re
import sys
from pathlib import Path

import openpyxl
import pandas as pd
import xlrd

# ==============================================================================
# SETTINGS & PATHS
# ==============================================================================

PROJECT_ROOT = Path(__file__).resolve().parents[2]
RAW_FOLDER = PROJECT_ROOT / "01_raw_files" / "ALPHAGREP_MF"
OUTPUT_FOLDER = PROJECT_ROOT / "03_clean_data" / "ALPHAGREP_MF"
OUTPUT_FOLDER.mkdir(parents=True, exist_ok=True)
OUTPUT_FILE = OUTPUT_FOLDER / "AlphaGrep_MF_All_Funds_Cleaned.xlsx"

AMC_NAME = "AlphaGrep MF"

STANDARD_COLUMNS = [
    "AMC",
    "Fund_Name",
    "Portfolio_Date",
    "Month",
    "Security_Name",
    "ISIN",
    "Industry_Rating",
    "Quantity",
]

TARGET_FUNDS = {
    "AG001": "AlphaGrep Multi Asset Allocation Fund",
    "AG002": "AlphaGrep Flexi Cap Fund",
}

STOP_MARKERS = [
    "sub total",
    "subtotal",
    "sub-total",
    "sub_total",
    "(b) unlisted",
    "(b)unlisted",
    "b) unlisted",
    "b)unlisted",
    "(b)",
    "unlisted",
    "total",
    "total (a)",
    "total(a)",
    "total (b)",
    "grand total",
    "debt instruments",
    "money market",
    "treps",
    "net current assets",
    "total net assets",
    "derivatives",
]


def clean_text(val):
    if val is None:
        return ""
    text = str(val).strip()
    return re.sub(r"\s+", " ", text)


def is_stop_instrument(name_val):
    if not name_val:
        return False
    lower = clean_text(name_val).lower()
    if lower in [
        "total", "sub total", "subtotal", "sub-total", "sub_total",
        "grand total", "total:", "total (a)", "total(a)", "total (b)", "total(b)", "unlisted"
    ]:
        return True
    if lower.startswith((
        "sub total", "subtotal", "sub-total", "sub_total",
        "total (", "total(", "total -", "total:", "grand total",
        "(b) unlisted", "(b)unlisted", "b) unlisted", "b)unlisted", "(b)", "b)"
    )):
        return True
    if any(t in lower for t in ["debt instruments", "money market", "treps", "net current assets", "total net assets", "derivatives"]):
        return True
    return False


def clean_alphagrep_data():
    print("=" * 80)
    print("ALPHAGREP MF - CLEANING PIPELINE")
    print(f"Raw Folder:   {RAW_FOLDER}")
    print(f"Output File:  {OUTPUT_FILE}")
    print("=" * 80)

    files = sorted(
        [
            f for f in list(RAW_FOLDER.glob("*/*/*.*"))
            if not f.name.startswith("~$")
        ]
    )

    if not files:
        print(f"ERROR: No monthly files found in {RAW_FOLDER}")
        sys.exit(1)

    print(f"\n[Step 1/3] Found {len(files)} raw files to process.")

    all_records = []

    for file_path in files:
        # Determine year/month from folder structure
        parts = file_path.parts
        yr = int(parts[-3])
        mo = int(parts[-2])
        pdate = pd.Timestamp(year=yr, month=mo, day=1)
        month_str = pdate.strftime("%b-%Y")

        # Open workbook (all AlphaGrep files are BIFF8 OLE .xls)
        try:
            wb = xlrd.open_workbook(file_path)
        except Exception as e:
            # Fallback to openpyxl if OpenXML
            try:
                wb = openpyxl.load_workbook(file_path, data_only=True)
            except Exception as e2:
                print(f"  [ERROR] Cannot open {file_path.name}: {e} / {e2}")
                continue

        # Check sheets
        if isinstance(wb, xlrd.Book):
            sheet_names = wb.sheet_names()
            for code, canonical_name in TARGET_FUNDS.items():
                if code not in sheet_names:
                    continue
                ws = wb.sheet_by_name(code)

                # Detect header columns
                col_inst = None
                col_isin = None
                col_ind = None
                col_qty = None
                header_r = None

                for r in range(min(15, ws.nrows)):
                    row_vals = [clean_text(ws.cell_value(r, c)).lower() for c in range(ws.ncols)]
                    for c_idx, val in enumerate(row_vals):
                        if "name of the instrument" in val or "instrument" in val:
                            col_inst = c_idx
                            header_r = r
                        elif val == "isin":
                            col_isin = c_idx
                        elif "industry" in val or "rating" in val:
                            col_ind = c_idx
                        elif "quantity" in val or "qty" in val:
                            col_qty = c_idx

                    if col_inst is not None and col_isin is not None and col_qty is not None:
                        break

                if col_inst is None or col_isin is None or col_qty is None:
                    print(f"  [WARN] Header not detected for {code} in {file_path.name}")
                    continue

                file_records = 0
                for r in range(header_r + 1, ws.nrows):
                    inst_val = clean_text(ws.cell_value(r, col_inst))
                    if is_stop_instrument(inst_val):
                        break

                    isin_val = clean_text(ws.cell_value(r, col_isin)).upper()
                    if not (isin_val.startswith("INE") and len(isin_val) == 12):
                        continue

                    ind_val = clean_text(ws.cell_value(r, col_ind)) if col_ind is not None else ""

                    try:
                        qty_raw = ws.cell_value(r, col_qty)
                        if qty_raw is None or str(qty_raw).strip() == "":
                            continue
                        qty = float(str(qty_raw).replace(",", "").strip())
                        if qty <= 0:
                            continue
                    except Exception:
                        continue

                    all_records.append({
                        "AMC": AMC_NAME,
                        "Fund_Name": canonical_name,
                        "Portfolio_Date": pdate.strftime("%Y-%m-%d"),
                        "Month": month_str,
                        "Security_Name": inst_val,
                        "ISIN": isin_val,
                        "Industry_Rating": ind_val,
                        "Quantity": qty,
                    })
                    file_records += 1

                print(f"  {month_str} | {file_path.name} -> {canonical_name}: {file_records} holdings")
        else:
            # openpyxl workbook branch
            sheet_names = wb.sheetnames
            for code, canonical_name in TARGET_FUNDS.items():
                if code not in sheet_names:
                    continue
                ws = wb[code]

                col_inst = None
                col_isin = None
                col_ind = None
                col_qty = None
                header_r = None

                for r in range(1, min(15, ws.max_row + 1)):
                    row_vals = [clean_text(ws.cell(r, c).value).lower() for c in range(1, ws.max_column + 1)]
                    for c_idx, val in enumerate(row_vals, 1):
                        if "name of the instrument" in val or "instrument" in val:
                            col_inst = c_idx
                            header_r = r
                        elif val == "isin":
                            col_isin = c_idx
                        elif "industry" in val or "rating" in val:
                            col_ind = c_idx
                        elif "quantity" in val or "qty" in val:
                            col_qty = c_idx

                    if col_inst and col_isin and col_qty:
                        break

                if not (col_inst and col_isin and col_qty):
                    print(f"  [WARN] Header not detected for {code} in {file_path.name}")
                    continue

                file_records = 0
                for r in range(header_r + 1, ws.max_row + 1):
                    inst_val = clean_text(ws.cell(r, col_inst).value)
                    if is_stop_instrument(inst_val):
                        break

                    isin_val = clean_text(ws.cell(r, col_isin).value).upper()
                    if not (isin_val.startswith("INE") and len(isin_val) == 12):
                        continue

                    ind_val = clean_text(ws.cell(r, col_ind).value) if col_ind else ""

                    try:
                        qty_raw = ws.cell(r, col_qty).value
                        if qty_raw is None or str(qty_raw).strip() == "":
                            continue
                        qty = float(str(qty_raw).replace(",", "").strip())
                        if qty <= 0:
                            continue
                    except Exception:
                        continue

                    all_records.append({
                        "AMC": AMC_NAME,
                        "Fund_Name": canonical_name,
                        "Portfolio_Date": pdate.strftime("%Y-%m-%d"),
                        "Month": month_str,
                        "Security_Name": inst_val,
                        "ISIN": isin_val,
                        "Industry_Rating": ind_val,
                        "Quantity": qty,
                    })
                    file_records += 1

                print(f"  {month_str} | {file_path.name} -> {canonical_name}: {file_records} holdings")

    df = pd.DataFrame(all_records)
    print("\n" + "=" * 80)
    print(f"[Step 2/3] Total extracted holdings: {len(df)}")
    print("Holdings by Month & Fund:")
    print(df.groupby(["Month", "Fund_Name"]).size())
    print("=" * 80)

    # Sort deterministically
    df["Month_Sort"] = pd.to_datetime(df["Portfolio_Date"])
    df = df.sort_values(by=["Month_Sort", "Fund_Name", "Security_Name"]).reset_index(drop=True)
    df = df.drop(columns=["Month_Sort"])
    df = df[STANDARD_COLUMNS]

    print(f"\n[Step 3/3] Saving cleaned data to: {OUTPUT_FILE}")
    with pd.ExcelWriter(OUTPUT_FILE, engine="openpyxl") as writer:
        df.to_excel(writer, index=False, sheet_name="Cleaned_Data")

    print(f"[SUCCESS] {AMC_NAME} data cleaning completed: {len(df)} records saved.\n")
    return df


if __name__ == "__main__":
    clean_alphagrep_data()
