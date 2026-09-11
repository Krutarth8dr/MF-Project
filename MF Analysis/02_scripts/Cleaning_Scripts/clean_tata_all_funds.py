import os
import re
import sys
import io
from datetime import datetime
from pathlib import Path
import openpyxl
import xlrd
import pandas as pd

# ==============================================================================
# SETTINGS & PATHS
# ==============================================================================

PROJECT_ROOT = Path(__file__).resolve().parents[2]
RAW_FOLDER = PROJECT_ROOT / "01_raw_files" / "TATA"
OUTPUT_FOLDER = PROJECT_ROOT / "03_clean_data" / "TATA"
OUTPUT_FOLDER.mkdir(parents=True, exist_ok=True)
OUTPUT_FILE = OUTPUT_FOLDER / "Tata_All_Funds_Cleaned.xlsx"

AMC_NAME = "TATA MF"
START_DATE = pd.Timestamp(year=2024, month=10, day=1)
END_DATE = pd.Timestamp(year=2026, month=8, day=31)

STANDARD_COLUMNS = [
    "AMC",
    "Fund_Name",
    "Portfolio_Date",
    "Month",
    "Security_Name",
    "ISIN",
    "Industry_Rating",
    "Quantity",
]

# Canonical 24 Target Funds (mapped from August 2026 Index)
TARGET_FUNDS_CONFIG = {
    "TEGF": "TATA AGGRESSIVE HYBRID FUND",
    "TBAF": "TATA BALANCED ADVANTAGE FUND",
    "TBFSF": "TATA BANKING & FINANCIAL SERVICES FUND",
    "TBCF": "TATA BUSINESS CYCLE FUND",
    "TYCF": "TATA CHILDRENS FUND",
    "TDIF": "TATA DIGITAL INDIA FUND",
    "TTSF96": "TATA ELSS - TAX SAVER FUND",
    "TCS": "TATA ETHICAL FUND",
    "TMCAPF": "TATA FLEXI CAP FUND",
    "TFEF": "TATA FOCUSED FUND",
    "THOF": "TATA HOUSING OPPORTUNITIES FUND",
    "TICF": "TATA INDIA CONSUMER FUND",
    "TIIF": "TATA INDIA INNOVATION FUND",
    "TIPHF": "TATA INDIA PHARMA & HEALTHCARE FUND",
    "TISF": "TATA INFRASTRUCTURE FUND",
    "TEOF": "TATA LARGE & MID CAP FUND",
    "TTOFE": "TATA LARGE CAP FUND",
    "TFRSTF": "TATA LIQUID FUND",
    "TINR": "TATA MID CAP FUND",
    "TMAOF": "TATA MULTI ASSET ALLOCATION FUND",
    "TMULTICF": "TATA MULTICAP FUND",
    "TREF": "TATA RESOURCES & ENERGY FUND",
    "TSCAPF": "TATA SMALL CAP FUND",
    "TEQPEF": "TATA VALUE FUND",
}


def clean_text(val):
    if val is None:
        return ""
    text = str(val).strip()
    return re.sub(r"\s+", " ", text)


def get_sheet_rows_openpyxl(wb, sheet_name):
    ws = wb[sheet_name]
    return list(ws.iter_rows(values_only=True))


def get_sheet_rows_xlrd(wb, sheet_name):
    ws = wb.sheet_by_name(sheet_name)
    rows = []
    for r in range(ws.nrows):
        rows.append([ws.cell_value(r, c) for c in range(ws.ncols)])
    return rows


def clean_tata_data():
    print("=" * 70)
    print("TATA MUTUAL FUND - CLEANING PIPELINE")
    print(f"Raw Folder:   {RAW_FOLDER}")
    print(f"Output File:  {OUTPUT_FILE}")
    print(f"Target Date Range: {START_DATE.strftime('%b %Y')} to {END_DATE.strftime('%b %Y')}")
    print("=" * 70)

    excel_files = sorted(list(RAW_FOLDER.glob("*/*/*.*")))
    if not excel_files:
        print(f"ERROR: No monthly files found in {RAW_FOLDER}")
        sys.exit(1)

    print(f"\n[Step 1/3] Found {len(excel_files)} monthly workbooks to process.")

    all_rows = []

    for fpath in excel_files:
        yr = int(fpath.parent.parent.name)
        mo = int(fpath.parent.name)

        portfolio_date = pd.Period(f"{yr}-{mo:02d}", freq="M").end_time.date()
        month_str = portfolio_date.strftime("%b-%Y")
        portfolio_dt = pd.to_datetime(portfolio_date)

        if not (START_DATE <= portfolio_dt <= END_DATE):
            continue

        print(f"  Processing {month_str} ({fpath.name})...")
        with open(fpath, "rb") as f:
            content = f.read()

        is_openxml = content[:4] == b'PK\x03\x04'

        if is_openxml:
            wb = openpyxl.load_workbook(io.BytesIO(content), data_only=True)
            sheet_names = set(wb.sheetnames)
            get_rows_fn = lambda sname: get_sheet_rows_openpyxl(wb, sname)
        else:
            wb = xlrd.open_workbook(file_contents=content)
            sheet_names = set(wb.sheet_names())
            get_rows_fn = lambda sname: get_sheet_rows_xlrd(wb, sname)

        for code, canonical_name in TARGET_FUNDS_CONFIG.items():
            if code not in sheet_names:
                continue

            rows_data = get_rows_fn(code)

            header_row = -1
            col_instrument = 1
            col_isin = 4
            col_industry = 3
            col_qty = 5

            for r_idx, r_vals in enumerate(rows_data[:20]):
                r_str = [str(v).lower() if v is not None else "" for v in r_vals]
                if any("instrument" in v for v in r_str) and (any("isin" in v for v in r_str) or any("quantity" in v for v in r_str)):
                    header_row = r_idx
                    for c_idx, val in enumerate(r_str):
                        if "instrument" in val or "name of" in val:
                            col_instrument = c_idx
                        elif "isin" in val:
                            col_isin = c_idx
                        elif "industry" in val or "rating" in val:
                            col_industry = c_idx
                        elif "quantity" in val or "qty" in val:
                            col_qty = c_idx
                    break

            if header_row == -1:
                header_row = 11

            for row in rows_data[header_row + 1:]:
                if not row or not any(row):
                    continue

                inst_val = clean_text(row[col_instrument]) if col_instrument < len(row) else ""

                # Stop scanning as soon as "UNLISTED" is seen in Name of Instrument column
                if "unlisted" in inst_val.lower():
                    break

                isin_val = clean_text(row[col_isin]) if col_isin < len(row) else ""
                if not isin_val.startswith("INE"):
                    continue

                industry_val = clean_text(row[col_industry]) if col_industry < len(row) else ""

                qty_raw = row[col_qty] if col_qty < len(row) else None
                try:
                    qty = float(qty_raw) if qty_raw is not None else 0.0
                except (ValueError, TypeError):
                    continue

                if qty <= 0:
                    continue

                all_rows.append({
                    "AMC": AMC_NAME,
                    "Fund_Name": canonical_name,
                    "Portfolio_Date": portfolio_date.strftime("%Y-%m-%d"),
                    "Month": month_str,
                    "Security_Name": inst_val,
                    "ISIN": isin_val,
                    "Industry_Rating": industry_val,
                    "Quantity": int(round(qty)),
                })

    print("\n[Step 2/3] Compiling cleaned dataset...")
    if not all_rows:
        print("ERROR: No valid equity holdings extracted!")
        sys.exit(1)

    df = pd.DataFrame(all_rows, columns=STANDARD_COLUMNS)

    # Convert types
    df["Portfolio_Date"] = pd.to_datetime(df["Portfolio_Date"])
    df["Quantity"] = df["Quantity"].astype("int64")

    # Sort dataset cleanly
    df = df.sort_values(["Portfolio_Date", "Fund_Name", "Security_Name"]).reset_index(drop=True)

    print(f"Total rows extracted: {len(df):,}")
    print(f"Unique funds:         {df['Fund_Name'].nunique()}")
    print(f"Unique ISINs:         {df['ISIN'].nunique()}")
    print(f"Unique months:        {df['Month'].nunique()}")

    print("\n[Step 3/3] Saving cleaned data to Excel...")
    df.to_excel(OUTPUT_FILE, index=False)
    print(f"Saved successfully: {OUTPUT_FILE} ({OUTPUT_FILE.stat().st_size:,} bytes)")
    print("=" * 70)


if __name__ == "__main__":
    clean_tata_data()
