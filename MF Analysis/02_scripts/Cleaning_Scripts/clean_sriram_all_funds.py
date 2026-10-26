"""
Clean all downloaded monthly portfolio files for Sriram Mutual Fund (Sriram MF).

Process:
    - Reads all files in 01_raw_files/Sriram_MF/<YYYY>/<MM>/
    - Parses target sheets for the 6 target funds:
        1. Shriram Aggressive Hybrid Fund
        2. Shriram Flexi Cap Fund
        3. Shriram ELSS Tax Saver Fund
        4. Shriram Balanced Advantage Fund
        5. Shriram Multi Asset Allocation Fund
        6. Shriram Multi Sector Rotation Fund
    - Dynamic header detection for Name of Instrument, ISIN, Industry/Rating, and Quantity
    - Stops immediately upon encountering 'Sub Total', '(b) Unlisted', '(b)unlisted', 'Total', etc.
    - Filters for valid listed Indian equities (ISIN starts with INE, length 12, quantity > 0)
    - Generates 03_clean_data/Sriram_MF/Sriram_MF_All_Funds_Cleaned.xlsx
"""

import sys
import os
import calendar
import warnings
from datetime import datetime
from pathlib import Path
import xlrd
import openpyxl
import pandas as pd

# Force UTF-8 output
try:
    sys.stdout.reconfigure(encoding="utf-8")
    sys.stderr.reconfigure(encoding="utf-8")
except (AttributeError, ValueError):
    pass

PROJECT_ROOT = Path(__file__).resolve().parents[2]
RAW_DIR = PROJECT_ROOT / "01_raw_files" / "Sriram_MF"
CLEAN_DIR = PROJECT_ROOT / "03_clean_data" / "Sriram_MF"
CLEAN_DIR.mkdir(parents=True, exist_ok=True)

OUTPUT_FILE = CLEAN_DIR / "Sriram_MF_All_Funds_Cleaned.xlsx"

AMC_NAME = "Sriram MF"

TARGET_SHEET_MAP = {
    "shriram aggressive hybrid fund": "Shriram Aggressive Hybrid Fund",
    "shriram flexi cap fund": "Shriram Flexi Cap Fund",
    "shriram elss tax saver fund": "Shriram ELSS Tax Saver Fund",
    "shriram balanced advantage fund": "Shriram Balanced Advantage Fund",
    "shriram multi asset allocation": "Shriram Multi Asset Allocation Fund",
    "shriram multi sector rotation": "Shriram Multi Sector Rotation Fund",
}

STOP_TERMS = [
    "sub total",
    "subtotal",
    "(b) unlisted",
    "(b)unlisted",
    "total",
    "debt instruments",
    "money market",
    "treasury bill",
    "cash & cash equivalents",
    "net current assets",
]

STANDARD_COLUMNS = [
    "AMC",
    "Fund_Name",
    "Portfolio_Date",
    "Month",
    "Security_Name",
    "ISIN",
    "Industry_Rating",
    "Quantity",
]


def clean_xls_sheet(sheet, fund_name, portfolio_date, month_str):
    """Parses a single sheet from an xlrd .xls workbook."""
    holdings = []
    header_row = None
    col_map = {}

    for r in range(min(sheet.nrows, 35)):
        row_vals = [str(sheet.cell_value(r, c)).strip().lower() for c in range(sheet.ncols)]
        if any("isin" in v for v in row_vals) and any("name" in v or "instrument" in v for v in row_vals):
            header_row = r
            for c, v in enumerate(row_vals):
                if "isin" in v and "isin" not in col_map:
                    col_map["isin"] = c
                elif ("name" in v or "instrument" in v) and "name" not in col_map:
                    col_map["name"] = c
                elif ("industry" in v or "rating" in v) and "industry" not in col_map:
                    col_map["industry"] = c
                elif ("quantity" in v or "qty" in v) and "qty" not in col_map:
                    col_map["qty"] = c
            break

    if header_row is None or "isin" not in col_map or "qty" not in col_map:
        return []

    if "name" not in col_map:
        col_map["name"] = 1
    if "industry" not in col_map:
        col_map["industry"] = 3

    in_listed_equity = False

    for r in range(header_row + 1, sheet.nrows):
        c0 = str(sheet.cell_value(r, 0)).strip().lower()
        c_name = str(sheet.cell_value(r, col_map["name"])).strip().lower()
        combined_text = f"{c0} {c_name}"

        if "(a)listed" in combined_text or "(a) listed" in combined_text or "listed / awaiting" in combined_text:
            in_listed_equity = True
            continue

        if in_listed_equity:
            if any(term in combined_text for term in STOP_TERMS):
                in_listed_equity = False
                break

            isin_val = str(sheet.cell_value(r, col_map["isin"])).strip().upper()
            name_val = str(sheet.cell_value(r, col_map["name"])).strip()
            ind_val = str(sheet.cell_value(r, col_map["industry"])).strip()
            raw_qty = sheet.cell_value(r, col_map["qty"])

            try:
                qty_num = float(raw_qty)
                if qty_num <= 0:
                    continue
                qty_clean = int(qty_num) if qty_num.is_integer() else qty_num
            except (ValueError, TypeError):
                continue

            if isin_val.startswith("INE") and len(isin_val) == 12:
                holdings.append({
                    "AMC": AMC_NAME,
                    "Fund_Name": fund_name,
                    "Portfolio_Date": portfolio_date,
                    "Month": month_str,
                    "Security_Name": name_val,
                    "ISIN": isin_val,
                    "Industry_Rating": ind_val,
                    "Quantity": qty_clean,
                })

    return holdings


def clean_xlsx_sheet(ws, fund_name, portfolio_date, month_str):
    """Parses a single sheet from an openpyxl .xlsx workbook."""
    holdings = []
    header_row = None
    col_map = {}

    for r in range(1, min(ws.max_row + 1, 35)):
        row_vals = [str(ws.cell(r, c).value or "").strip().lower() for c in range(1, min(15, ws.max_column + 1))]
        if any("isin" in v for v in row_vals) and any("name" in v or "instrument" in v for v in row_vals):
            header_row = r
            for c, v in enumerate(row_vals, 1):
                if "isin" in v and "isin" not in col_map:
                    col_map["isin"] = c
                elif ("name" in v or "instrument" in v) and "name" not in col_map:
                    col_map["name"] = c
                elif ("industry" in v or "rating" in v) and "industry" not in col_map:
                    col_map["industry"] = c
                elif ("quantity" in v or "qty" in v) and "qty" not in col_map:
                    col_map["qty"] = c
            break

    if header_row is None or "isin" not in col_map or "qty" not in col_map:
        return []

    if "name" not in col_map:
        col_map["name"] = 1
    if "industry" not in col_map:
        col_map["industry"] = 3

    in_listed_equity = False

    for r in range(header_row + 1, ws.max_row + 1):
        c0 = str(ws.cell(r, 1).value or "").strip().lower()
        c_name = str(ws.cell(r, col_map["name"]).value or "").strip().lower()
        combined_text = f"{c0} {c_name}"

        if "(a)listed" in combined_text or "(a) listed" in combined_text or "listed / awaiting" in combined_text:
            in_listed_equity = True
            continue

        if in_listed_equity:
            if any(term in combined_text for term in STOP_TERMS):
                in_listed_equity = False
                break

            isin_val = str(ws.cell(r, col_map["isin"]).value or "").strip().upper()
            name_val = str(ws.cell(r, col_map["name"]).value or "").strip()
            ind_val = str(ws.cell(r, col_map["industry"]).value or "").strip()
            raw_qty = ws.cell(r, col_map["qty"]).value

            try:
                qty_num = float(raw_qty)
                if qty_num <= 0:
                    continue
                qty_clean = int(qty_num) if qty_num.is_integer() else qty_num
            except (ValueError, TypeError):
                continue

            if isin_val.startswith("INE") and len(isin_val) == 12:
                holdings.append({
                    "AMC": AMC_NAME,
                    "Fund_Name": fund_name,
                    "Portfolio_Date": portfolio_date,
                    "Month": month_str,
                    "Security_Name": name_val,
                    "ISIN": isin_val,
                    "Industry_Rating": ind_val,
                    "Quantity": qty_clean,
                })

    return holdings


def process_all_files():
    """Iterate over all raw files in Sriram_MF, extract holdings for target funds, and write clean Excel."""
    print("=" * 80)
    print("Starting Sriram MF Portfolio Cleaner")
    print(f"Reading from: {RAW_DIR}")
    print(f"Output target: {OUTPUT_FILE}")
    print("=" * 80)

    all_records = []
    file_count = 0

    for year_dir in sorted(RAW_DIR.iterdir()):
        if not year_dir.is_dir() or not year_dir.name.isdigit():
            continue
        year = int(year_dir.name)

        for month_dir in sorted(year_dir.iterdir()):
            if not month_dir.is_dir() or not month_dir.name.isdigit():
                continue
            month = int(month_dir.name)

            last_day = calendar.monthrange(year, month)[1]
            portfolio_date = datetime(year, month, last_day)
            month_str = f"{year}-{month:02d}-01"

            for file_path in sorted(list(month_dir.glob("*.xls")) + list(month_dir.glob("*.xlsx"))):
                file_count += 1
                month_records = 0

                if file_path.suffix.lower() == ".xls":
                    try:
                        wb = xlrd.open_workbook(file_path)
                        for sheet_name in wb.sheet_names():
                            s_clean = sheet_name.strip().lower()
                            canon = None
                            for k, v in TARGET_SHEET_MAP.items():
                                if k in s_clean:
                                    canon = v
                                    break
                            if not canon:
                                continue

                            sheet = wb.sheet_by_name(sheet_name)
                            h = clean_xls_sheet(sheet, canon, portfolio_date, month_str)
                            all_records.extend(h)
                            month_records += len(h)
                    except Exception as e:
                        print(f"Error parsing .xls file {file_path.name}: {e}")

                elif file_path.suffix.lower() == ".xlsx":
                    try:
                        wb = openpyxl.load_workbook(file_path, data_only=True)
                        for sheet_name in wb.sheetnames:
                            s_clean = sheet_name.strip().lower()
                            canon = None
                            for k, v in TARGET_SHEET_MAP.items():
                                if k in s_clean:
                                    canon = v
                                    break
                            if not canon:
                                continue

                            ws = wb[sheet_name]
                            h = clean_xlsx_sheet(ws, canon, portfolio_date, month_str)
                            all_records.extend(h)
                            month_records += len(h)
                    except Exception as e:
                        print(f"Error parsing .xlsx file {file_path.name}: {e}")

                print(f"[{year}-{month:02d}] {file_path.name}: {month_records} holdings extracted.")

    print("=" * 80)
    print(f"Total processed files: {file_count}")
    print(f"Total holdings extracted: {len(all_records)}")

    if not all_records:
        print("ERROR: No holdings extracted! Aborting.")
        return False

    df = pd.DataFrame(all_records, columns=STANDARD_COLUMNS)

    # Validate dataframe
    print("\nSummary by Fund:")
    fund_summary = df.groupby("Fund_Name")["Quantity"].agg(["count", "nunique"]).reset_index()
    fund_summary.columns = ["Fund Name", "Total Records", "Unique Quantities"]
    print(fund_summary.to_string(index=False))

    print("\nSummary by Month:")
    month_summary = df.groupby("Month")["Fund_Name"].agg(["count", "nunique"]).reset_index()
    month_summary.columns = ["Month", "Total Records", "Funds Present"]
    print(month_summary.to_string(index=False))

    # Save to Excel
    print(f"\nWriting cleaned dataset to: {OUTPUT_FILE}")
    with pd.ExcelWriter(OUTPUT_FILE, engine="openpyxl") as writer:
        df.to_excel(writer, index=False, sheet_name="Cleaned_Holdings")

    print(f"Successfully generated: {OUTPUT_FILE} ({OUTPUT_FILE.stat().st_size / 1024:.1f} KB)")
    print("=" * 80)
    return True


if __name__ == "__main__":
    success = process_all_files()
    sys.exit(0 if success else 1)
