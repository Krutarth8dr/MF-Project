"""
Clean all downloaded monthly portfolio files for PGIM Mutual Fund (PGIM MF).

Process:
    - Reads all files in 01_raw_files/PGIM_MF/<YYYY>/<MM>/
    - Supports both .xlsx and .xlsb workbooks
    - Target Schemes (12 funds):
        1. PGIM India Healthcare Fund
        2. PGIM India Small Cap Fund
        3. PGIM India Multi Cap Fund
        4. PGIM India Multi Asset Allocation Fund
        5. PGIM India Midcap Fund
        6. PGIM India Large Cap Fund
        7. PGIM India Large and Midcap Fund
        8. PGIM India Flexi Cap Fund
        9. PGIM India Equity Savings Fund
        10. PGIM India ELSS Tax Saver Fund
        11. PGIM India Balanced Advantage Fund
        12. PGIM India Aggressive Hybrid Equity Fund
    - Dynamic header detection for Name, ISIN, Industry/Rating, and Quantity
    - Stops immediately upon encountering 'Sub Total', '(b)Unlisted', '(b) Unlisted', or 'Total'
    - Filters for valid listed Indian equities (ISIN starts with INE, length 12, quantity > 0)
    - Generates 03_clean_data/PGIM_MF/PGIM_MF_All_Funds_Cleaned.xlsx
"""

import sys
import os
import calendar
import warnings
from datetime import datetime
from pathlib import Path
import openpyxl
from pyxlsb import open_workbook
import pandas as pd

warnings.filterwarnings("ignore", category=UserWarning, module="openpyxl")

# Force UTF-8 output
try:
    sys.stdout.reconfigure(encoding="utf-8")
    sys.stderr.reconfigure(encoding="utf-8")
except (AttributeError, ValueError):
    pass

PROJECT_ROOT = Path(__file__).resolve().parents[2]
RAW_DIR = PROJECT_ROOT / "01_raw_files" / "PGIM_MF"
CLEAN_DIR = PROJECT_ROOT / "03_clean_data" / "PGIM_MF"
CLEAN_DIR.mkdir(parents=True, exist_ok=True)

OUTPUT_FILE = CLEAN_DIR / "PGIM_MF_All_Funds_Cleaned.xlsx"

SLUG_TO_FUND = {
    "Healthcare": "PGIM India Healthcare Fund",
    "SmallCap": "PGIM India Small Cap Fund",
    "MultiCap": "PGIM India Multi Cap Fund",
    "MultiAsset": "PGIM India Multi Asset Allocation Fund",
    "Midcap": "PGIM India Midcap Fund",
    "LargeCap": "PGIM India Large Cap Fund",
    "LargeAndMidcap": "PGIM India Large and Midcap Fund",
    "FlexiCap": "PGIM India Flexi Cap Fund",
    "EquitySavings": "PGIM India Equity Savings Fund",
    "ELSS": "PGIM India ELSS Tax Saver Fund",
    "BalancedAdvantage": "PGIM India Balanced Advantage Fund",
    "AggressiveHybrid": "PGIM India Aggressive Hybrid Equity Fund",
}

STOP_TERMS = ["sub total", "(b)unlisted", "(b) unlisted", "total", "subtotal"]

STANDARD_COLUMNS = [
    "AMC",
    "Fund_Name",
    "Portfolio_Date",
    "Month",
    "Security_Name",
    "ISIN",
    "Industry_Rating",
    "Quantity",
]


def clean_xlsx_sheet(ws, fund_name, portfolio_date, month_str):
    """Parses an openpyxl worksheet."""
    header_row = None
    col_map = {}
    for r in range(1, 20):
        row_vals = [str(ws.cell(row=r, column=c).value or "").strip() for c in range(1, min(15, ws.max_column + 1))]
        lower = [v.lower() for v in row_vals]
        if any("isin" in v for v in lower) and any("instrument" in v or "name" in v for v in lower):
            header_row = r
            for c_idx, val in enumerate(lower, 1):
                if "isin" in val and "isin" not in col_map:
                    col_map["isin"] = c_idx
                elif ("instrument" in val or "name" in val) and "yield" not in val and "name" not in col_map:
                    col_map["name"] = c_idx
                elif ("rating" in val or "industry" in val) and "industry" not in col_map:
                    col_map["industry"] = c_idx
                elif ("quantity" in val or "qty" in val) and "qty" not in col_map:
                    col_map["qty"] = c_idx
            break

    if not header_row or "isin" not in col_map or "name" not in col_map or "qty" not in col_map:
        return []

    holdings = []
    for r in range(header_row + 1, ws.max_row + 1):
        name_val = str(ws.cell(row=r, column=col_map["name"]).value or "").strip()
        isin_val = str(ws.cell(row=r, column=col_map["isin"]).value or "").strip()

        combined = f"{name_val.lower()} {isin_val.lower()}"
        if any(term in combined for term in STOP_TERMS):
            break

        if not (isin_val.startswith("INE") and len(isin_val) == 12):
            continue

        raw_qty = ws.cell(row=r, column=col_map["qty"]).value
        try:
            qty = float(raw_qty)
            if qty <= 0:
                continue
        except (ValueError, TypeError):
            continue

        industry_val = ""
        if "industry" in col_map:
            industry_val = str(ws.cell(row=r, column=col_map["industry"]).value or "").strip()

        holdings.append({
            "AMC": "PGIM MF",
            "Fund_Name": fund_name,
            "Portfolio_Date": portfolio_date,
            "Month": month_str,
            "Security_Name": name_val,
            "ISIN": isin_val,
            "Industry_Rating": industry_val,
            "Quantity": qty,
        })

    return holdings


def clean_xlsb_file(file_path, fund_name, portfolio_date, month_str):
    """Parses a pyxlsb binary workbook."""
    holdings = []
    with open_workbook(file_path) as wb:
        with wb.get_sheet(wb.sheets[0]) as ws:
            rows = list(ws.rows())
            header_row = None
            col_map = {}
            for r_idx, row in enumerate(rows[:20]):
                vals = [str(cell.v or "").strip() for cell in row]
                lower = [v.lower() for v in vals]
                if any("isin" in v for v in lower) and any("instrument" in v or "name" in v for v in lower):
                    header_row = r_idx + 1
                    for c_idx, val in enumerate(lower, 1):
                        if "isin" in val and "isin" not in col_map:
                            col_map["isin"] = c_idx
                        elif ("instrument" in val or "name" in val) and "yield" not in val and "name" not in col_map:
                            col_map["name"] = c_idx
                        elif ("rating" in val or "industry" in val) and "industry" not in col_map:
                            col_map["industry"] = c_idx
                        elif ("quantity" in val or "qty" in val) and "qty" not in col_map:
                            col_map["qty"] = c_idx
                    break

            if not header_row or "isin" not in col_map or "name" not in col_map or "qty" not in col_map:
                return []

            for r_idx, row in enumerate(rows[header_row:], header_row + 1):
                vals = [str(cell.v or "").strip() for cell in row]
                name_val = vals[col_map["name"] - 1] if col_map["name"] <= len(vals) else ""
                isin_val = vals[col_map["isin"] - 1] if col_map["isin"] <= len(vals) else ""

                combined = f"{name_val.lower()} {isin_val.lower()}"
                if any(term in combined for term in STOP_TERMS):
                    break

                if not (isin_val.startswith("INE") and len(isin_val) == 12):
                    continue

                raw_qty = vals[col_map["qty"] - 1] if col_map["qty"] <= len(vals) else 0
                try:
                    qty = float(raw_qty)
                    if qty <= 0:
                        continue
                except (ValueError, TypeError):
                    continue

                industry_val = ""
                if "industry" in col_map and col_map["industry"] <= len(vals):
                    industry_val = vals[col_map["industry"] - 1]

                holdings.append({
                    "AMC": "PGIM MF",
                    "Fund_Name": fund_name,
                    "Portfolio_Date": portfolio_date,
                    "Month": month_str,
                    "Security_Name": name_val,
                    "ISIN": isin_val,
                    "Industry_Rating": industry_val,
                    "Quantity": qty,
                })

    return holdings


def process_all_files():
    """Iterate over all raw files in PGIM_MF folder, clean and save to single Excel file."""
    print("=" * 80)
    print("Starting PGIM MF Portfolio Cleaner")
    print(f"Reading from: {RAW_DIR}")
    print(f"Output target: {OUTPUT_FILE}")
    print("=" * 80)

    all_records = []
    processed_count = 0

    for year_dir in sorted(RAW_DIR.iterdir()):
        if not year_dir.is_dir() or not year_dir.name.isdigit():
            continue
        year = int(year_dir.name)

        for month_dir in sorted(year_dir.iterdir()):
            if not month_dir.is_dir() or not month_dir.name.isdigit():
                continue
            month = int(month_dir.name)

            last_day = calendar.monthrange(year, month)[1]
            portfolio_date = datetime(year, month, last_day)
            month_str = f"{year}-{month:02d}-01"

            for file_path in sorted(month_dir.glob("*.*")):
                file_name = file_path.name
                ext = file_path.suffix.lower()
                if ext not in (".xlsx", ".xlsb", ".xls"):
                    continue

                # Determine fund name from slug prefix
                fund_name = None
                for slug, full_name in SLUG_TO_FUND.items():
                    if file_name.startswith(f"{slug}_"):
                        fund_name = full_name
                        break

                if not fund_name:
                    print(f"WARNING: Unknown fund slug for {file_name}")
                    continue

                try:
                    if ext == ".xlsx":
                        wb = openpyxl.load_workbook(file_path, data_only=True)
                        ws = wb.active
                        holdings = clean_xlsx_sheet(ws, fund_name, portfolio_date, month_str)
                    elif ext == ".xlsb":
                        holdings = clean_xlsb_file(file_path, fund_name, portfolio_date, month_str)
                    else:
                        continue

                    all_records.extend(holdings)
                    processed_count += 1
                except Exception as e:
                    print(f"ERROR processing {file_name}: {e}")

    df = pd.DataFrame(all_records, columns=STANDARD_COLUMNS)

    print("\n" + "=" * 80)
    print("CLEANING COMPLETED - SUMMARY")
    print("=" * 80)
    print(f"Total Files Processed: {processed_count}")
    print(f"Total Rows Extracted: {len(df):,}")
    print(f"Unique AMCs: {df['AMC'].nunique()} -> {list(df['AMC'].unique())}")
    print(f"Unique Funds ({df['Fund_Name'].nunique()}):")
    for f in df["Fund_Name"].unique():
        f_df = df[df["Fund_Name"] == f]
        print(f"  - {f:42}: {len(f_df):5,d} rows, {f_df['Month'].nunique():2d} months ({f_df['Month'].min()} to {f_df['Month'].max()})")
    print(f"Unique ISINs: {df['ISIN'].nunique():,}")
    print(f"Date Range: {df['Month'].min()} to {df['Month'].max()}")
    print("=" * 80)

    # Save to Excel
    print(f"Writing clean dataset to {OUTPUT_FILE}...")
    df.to_excel(OUTPUT_FILE, index=False, engine="openpyxl")
    file_size_kb = OUTPUT_FILE.stat().st_size / 1024
    print(f"Successfully saved {OUTPUT_FILE.name} ({file_size_kb:.1f} KB)")


if __name__ == "__main__":
    process_all_files()
