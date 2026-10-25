"""
Cleaning script for JM Financial Mutual Fund (JM Financial MF).
Extracts listed equity holdings from monthly files for 10 target funds:
- JM Aggressive Hybrid Fund
- JM Small Cap Fund
- JM Multi Asset Allocation Fund
- JM Mid Cap Fund
- JM Large Cap Fund
- JM Large and Mid Cap Fund
- JM Focused Fund
- JM Flexicap Fund
- JM ELSS Tax Saver Fund
- JM Value Fund

Stops cleaning as soon as "Sub Total", "(b)Unlisted", "(b) Unlisted", or "Total"
are detected in the "Name of Instrument" or "ISIN" column.

Output:
03_clean_data/JM_Financial_MF/JM_Financial_MF_All_Funds_Cleaned.xlsx
"""

import datetime
import io
import os
import re
import sys
from pathlib import Path

import openpyxl
import pandas as pd
import xlrd

# ==============================================================================
# SETTINGS & PATHS
# ==============================================================================

PROJECT_ROOT = Path(__file__).resolve().parents[2]
RAW_FOLDER = PROJECT_ROOT / "01_raw_files" / "JM_Financial_MF"
OUTPUT_FOLDER = PROJECT_ROOT / "03_clean_data" / "JM_Financial_MF"
OUTPUT_FOLDER.mkdir(parents=True, exist_ok=True)
OUTPUT_FILE = OUTPUT_FOLDER / "JM_Financial_MF_All_Funds_Cleaned.xlsx"

AMC_NAME = "JM Financial MF"

STANDARD_COLUMNS = [
    "AMC",
    "Fund_Name",
    "Portfolio_Date",
    "Month",
    "Security_Name",
    "ISIN",
    "Industry_Rating",
    "Quantity",
]

TARGET_FUNDS = {
    "JM Aggressive Hybrid Fund": [r"aggressive\s*hybrid"],
    "JM Small Cap Fund": [r"small\s*cap"],
    "JM Multi Asset Allocation Fund": [r"multi\s*asset"],
    "JM Mid Cap Fund": [r"mid\s*cap"],
    "JM Large Cap Fund": [r"large\s*cap"],
    "JM Large and Mid Cap Fund": [r"large\s*(?:and|&)\s*mid"],
    "JM Focused Fund": [r"focused"],
    "JM Flexicap Fund": [r"flexi\s*cap"],
    "JM ELSS Tax Saver Fund": [r"elss", r"tax\s*saver"],
    "JM Value Fund": [r"value\s*fund"],
}

STOP_TERMS = ["sub total", "(b)unlisted", "(b) unlisted", "total"]

MONTHS_PREFIX = {
    "jan": 1, "feb": 2, "mar": 3, "apr": 4, "may": 5, "jun": 6,
    "jul": 7, "aug": 8, "sep": 9, "oct": 10, "nov": 11, "dec": 12,
}

def identify_fund(filename: str):
    """
    Identifies target fund name from the raw filename.
    """
    name_lower = filename.lower()
    if re.search(r"large\s*(?:and|&)\s*mid", name_lower):
        return "JM Large and Mid Cap Fund"
    for fund_name, patterns in TARGET_FUNDS.items():
        if fund_name == "JM Large and Mid Cap Fund":
            continue
        for pat in patterns:
            if re.search(pat, name_lower):
                return fund_name
    return None

def extract_date_from_text(text: str):
    """
    Extracts date string 'YYYY-MM-DD' from text.
    """
    if not text:
        return None

    # E.g. 30.09.2026 or 31/10/2024 or 31-10-2024
    m_num = re.search(r'(\d{1,2})[./-](\d{1,2})[./-](\d{4})', text)
    if m_num:
        d = int(m_num.group(1))
        m = int(m_num.group(2))
        y = int(m_num.group(3))
        try:
            return datetime.date(y, m, d).strftime("%Y-%m-%d")
        except Exception:
            pass

    # E.g. September 30, 2026 or October 31 2024
    m_word = re.search(r'([A-Za-z]+)\s*(\d{1,2})?,?\s*(\d{4})', text)
    if m_word:
        w = m_word.group(1).lower()
        y = int(m_word.group(3))
        d = int(m_word.group(2)) if m_word.group(2) else 28
        for pfx, m_val in MONTHS_PREFIX.items():
            if w.startswith(pfx):
                try:
                    return datetime.date(y, m_val, d).strftime("%Y-%m-%d")
                except Exception:
                    pass
    return None

def process_openpyxl_sheet(content: bytes, fund_name: str, default_date: str):
    """
    Processes an openpyxl sheet for a single fund file.
    """
    wb = openpyxl.load_workbook(io.BytesIO(content), data_only=True)
    ws = wb.active

    portfolio_date = None
    for r in range(1, 6):
        for c in range(1, 10):
            val = ws.cell(row=r, column=c).value
            if val and isinstance(val, str) and ("statement" in val.lower() or "period ended" in val.lower() or "portfolio" in val.lower()):
                portfolio_date = extract_date_from_text(val)
                if portfolio_date:
                    break
        if portfolio_date:
            break

    if not portfolio_date:
        portfolio_date = default_date

    header_row = None
    col_map = {}
    for r in range(1, 15):
        row_vals = [ws.cell(row=r, column=c).value for c in range(1, 15)]
        str_vals = [str(v).strip().lower() if v is not None else "" for v in row_vals]
        if any("isin" in s for s in str_vals) and any("instrument" in s or "name" in s for s in str_vals):
            header_row = r
            for c_idx, s in enumerate(str_vals, start=1):
                if "isin" in s:
                    col_map["isin"] = c_idx
                elif ("instrument" in s or "name" in s) and "yield" not in s:
                    if "name" not in col_map:
                        col_map["name"] = c_idx
                elif "industry" in s or "rating" in s:
                    if "industry" not in col_map:
                        col_map["industry"] = c_idx
                elif "quantity" in s or s == "qty":
                    col_map["qty"] = c_idx
            break

    if not header_row or "name" not in col_map or "isin" not in col_map or "qty" not in col_map:
        return []

    records = []
    for r in range(header_row + 1, ws.max_row + 1):
        name_val = ws.cell(row=r, column=col_map["name"]).value
        isin_val = ws.cell(row=r, column=col_map["isin"]).value
        qty_val = ws.cell(row=r, column=col_map["qty"]).value
        ind_val = ws.cell(row=r, column=col_map["industry"]).value if "industry" in col_map else None

        name_str = str(name_val).strip() if name_val is not None else ""
        if any(term in name_str.lower() for term in STOP_TERMS):
            break

        isin_str = str(isin_val).strip() if isin_val is not None else ""
        if any(term in isin_str.lower() for term in STOP_TERMS):
            break

        if not (isin_str.startswith("INE") and len(isin_str) == 12):
            continue

        if not name_str:
            continue

        try:
            qty = float(qty_val)
            if qty <= 0:
                continue
        except (ValueError, TypeError):
            continue

        records.append({
            "AMC": AMC_NAME,
            "Fund_Name": fund_name,
            "Portfolio_Date": portfolio_date,
            "Security_Name": name_str,
            "ISIN": isin_str,
            "Industry_Rating": str(ind_val).strip() if ind_val is not None else "",
            "Quantity": qty,
        })

    return records

def process_xlrd_sheet(content: bytes, fund_name: str, default_date: str):
    """
    Processes an xlrd sheet for legacy .xls files.
    """
    wb = xlrd.open_workbook(file_contents=content)
    ws = wb.sheet_by_index(0)

    portfolio_date = None
    for r in range(min(6, ws.nrows)):
        for c in range(min(10, ws.ncols)):
            val = ws.cell_value(r, c)
            if val and isinstance(val, str) and ("statement" in val.lower() or "period ended" in val.lower() or "portfolio" in val.lower()):
                portfolio_date = extract_date_from_text(val)
                if portfolio_date:
                    break
        if portfolio_date:
            break

    if not portfolio_date:
        portfolio_date = default_date

    header_row = None
    col_map = {}
    for r in range(min(15, ws.nrows)):
        row_vals = ws.row_values(r)
        str_vals = [str(v).strip().lower() if v is not None else "" for v in row_vals]
        if any("isin" in s for s in str_vals) and any("instrument" in s or "name" in s for s in str_vals):
            header_row = r
            for c_idx, s in enumerate(str_vals):
                if "isin" in s:
                    col_map["isin"] = c_idx
                elif ("instrument" in s or "name" in s) and "yield" not in s:
                    if "name" not in col_map:
                        col_map["name"] = c_idx
                elif "industry" in s or "rating" in s:
                    if "industry" not in col_map:
                        col_map["industry"] = c_idx
                elif "quantity" in s or s == "qty":
                    col_map["qty"] = c_idx
            break

    if header_row is None or "name" not in col_map or "isin" not in col_map or "qty" not in col_map:
        return []

    records = []
    for r in range(header_row + 1, ws.nrows):
        name_val = ws.cell_value(r, col_map["name"])
        isin_val = ws.cell_value(r, col_map["isin"])
        qty_val = ws.cell_value(r, col_map["qty"])
        ind_val = ws.cell_value(r, col_map["industry"]) if "industry" in col_map else None

        name_str = str(name_val).strip() if name_val is not None else ""
        if any(term in name_str.lower() for term in STOP_TERMS):
            break

        isin_str = str(isin_val).strip() if isin_val is not None else ""
        if any(term in isin_str.lower() for term in STOP_TERMS):
            break

        if not (isin_str.startswith("INE") and len(isin_str) == 12):
            continue

        if not name_str:
            continue

        try:
            qty = float(qty_val)
            if qty <= 0:
                continue
        except (ValueError, TypeError):
            continue

        records.append({
            "AMC": AMC_NAME,
            "Fund_Name": fund_name,
            "Portfolio_Date": portfolio_date,
            "Security_Name": name_str,
            "ISIN": isin_str,
            "Industry_Rating": str(ind_val).strip() if ind_val is not None else "",
            "Quantity": qty,
        })

    return records

def clean_all_jm_files():
    print("=" * 80)
    print("JM Financial MF - Portfolio Cleaning Pipeline")
    print("=" * 80)

    files = sorted(RAW_FOLDER.rglob("*.xlsx")) + sorted(RAW_FOLDER.rglob("*.xls"))
    print(f"Found {len(files)} portfolio files in {RAW_FOLDER}")

    all_records = []
    success_count = 0

    for idx, f in enumerate(files, start=1):
        rel = f.relative_to(RAW_FOLDER)
        year_str, month_str = rel.parts[0], rel.parts[1]
        last_day = 31 if int(month_str) in [1, 3, 5, 7, 8, 10, 12] else (30 if int(month_str) in [4, 6, 9, 11] else (29 if int(year_str) == 2024 else 28))
        default_date = f"{year_str}-{month_str}-{last_day:02d}"

        fund_name = identify_fund(f.name)
        if not fund_name:
            print(f"[{idx:3d}/{len(files)}] [WARN] Could not identify fund for: {f.name}")
            continue

        content = f.read_bytes()
        if content.startswith(b"PK"):
            recs = process_openpyxl_sheet(content, fund_name, default_date)
        elif content.startswith(b"\xd0\xcf"):
            recs = process_xlrd_sheet(content, fund_name, default_date)
        else:
            recs = []

        if recs:
            all_records.extend(recs)
            success_count += 1

        if idx % 25 == 0 or idx == len(files):
            print(f"[{idx:3d}/{len(files)}] Processed... ({success_count} files valid, {len(all_records)} holdings accumulated)")

    if not all_records:
        print("❌ No records extracted!")
        return

    df = pd.DataFrame(all_records)
    df["Month"] = pd.to_datetime(df["Portfolio_Date"]).dt.strftime("%b-%Y")
    df = df[STANDARD_COLUMNS]

    df.sort_values(
        by=["Fund_Name", "Portfolio_Date", "Security_Name"],
        ascending=True,
        inplace=True,
    )
    df.reset_index(drop=True, inplace=True)

    df.to_excel(OUTPUT_FILE, index=False)

    print("\n" + "=" * 80)
    print("Cleaned Data Summary")
    print("=" * 80)
    print(f"Total Records : {len(df)}")
    print(f"Unique AMCs   : {df['AMC'].nunique()} ({list(df['AMC'].unique())})")
    print(f"Unique Funds  : {df['Fund_Name'].nunique()}")
    for fname, count in df["Fund_Name"].value_counts().items():
        print(f"  - {fname:<35}: {count:5d} rows")
    print(f"Unique ISINs  : {df['ISIN'].nunique()}")
    print(f"Date Range    : {df['Portfolio_Date'].min()} to {df['Portfolio_Date'].max()} ({df['Portfolio_Date'].nunique()} months)")
    print(f"Output File   : {OUTPUT_FILE}")
    print("=" * 80)

if __name__ == "__main__":
    clean_all_jm_files()
