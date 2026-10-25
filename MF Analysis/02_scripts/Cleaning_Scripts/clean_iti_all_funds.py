"""
Cleaning script for ITI Mutual Fund (ITI MF).
Extracts listed equity holdings from monthly files for 10 target funds:
- ITIBAF: ITI Balanced Advantage Fund
- ITIBCF: ITI Bharat Consumption Fund
- ITIBCL: ITI Business Cycle Fund
- ITIBFS: ITI Banking and Financial Services Fund
- ITIFCF: ITI Flexi Cap Fund
- ITIFF:  ITI Focused Fund
- ITILCF: ITI Large Cap Fund
- ITILMF: ITI Large & Mid Cap Fund
- ITILTE: ITI ELSS Tax Saver Fund
- ITIMAF: ITI Multi Asset Allocation Fund

Stops cleaning as soon as "Sub Total", "(b)Unlisted", "(b) Unlisted", or "Total"
are detected in the "Name of the Instrument" column.

Output:
03_clean_data/ITI_MF/ITI_MF_All_Funds_Cleaned.xlsx
"""

import datetime
import io
import os
import re
import sys
from pathlib import Path

import openpyxl
import pandas as pd
import xlrd

# ==============================================================================
# SETTINGS & PATHS
# ==============================================================================

PROJECT_ROOT = Path(__file__).resolve().parents[2]
RAW_FOLDER = PROJECT_ROOT / "01_raw_files" / "ITI_MF"
OUTPUT_FOLDER = PROJECT_ROOT / "03_clean_data" / "ITI_MF"
OUTPUT_FOLDER.mkdir(parents=True, exist_ok=True)
OUTPUT_FILE = OUTPUT_FOLDER / "ITI_MF_All_Funds_Cleaned.xlsx"

AMC_NAME = "ITI MF"

STANDARD_COLUMNS = [
    "AMC",
    "Fund_Name",
    "Portfolio_Date",
    "Month",
    "Security_Name",
    "ISIN",
    "Industry_Rating",
    "Quantity",
]

TARGET_FUNDS = {
    "ITIBAF": "ITI Balanced Advantage Fund",
    "ITIBCF": "ITI Bharat Consumption Fund",
    "ITIBCL": "ITI Business Cycle Fund",
    "ITIBFS": "ITI Banking and Financial Services Fund",
    "ITIFCF": "ITI Flexi Cap Fund",
    "ITIFF":  "ITI Focused Fund",
    "ITILCF": "ITI Large Cap Fund",
    "ITILMF": "ITI Large & Mid Cap Fund",
    "ITILTE": "ITI ELSS Tax Saver Fund",
    "ITIMAF": "ITI Multi Asset Allocation Fund",
}

STOP_TERMS = ["sub total", "(b)unlisted", "(b) unlisted", "total"]

MONTHS_MAP = {
    "january": 1, "february": 2, "march": 3, "april": 4, "may": 5, "june": 6,
    "july": 7, "august": 8, "september": 9, "october": 10, "november": 11, "december": 12,
    "jan": 1, "feb": 2, "mar": 3, "apr": 4, "jun": 6, "jul": 7, "aug": 8, "sep": 9, "oct": 10, "nov": 11, "dec": 12,
}

def extract_date_from_text(text: str):
    """
    Extracts date string 'YYYY-MM-DD' from portfolio header text.
    """
    if not text:
        return None
    m = re.search(r'([A-Za-z]+)\s+(\d{1,2})\s*,?\s*(\d{4})', text)
    if m:
        month_name = m.group(1).lower()
        day = int(m.group(2))
        year = int(m.group(3))
        month = MONTHS_MAP.get(month_name)
        if month:
            try:
                return datetime.date(year, month, day).strftime("%Y-%m-%d")
            except Exception:
                pass
    return None

def process_openpyxl_sheet(ws, fund_code, fund_name, default_date):
    """
    Processes an openpyxl sheet for a single fund.
    """
    portfolio_date = None
    for r in range(1, 9):
        for c in range(1, 10):
            val = ws.cell(row=r, column=c).value
            if val and isinstance(val, str) and "statement as on" in val.lower():
                portfolio_date = extract_date_from_text(val)
                if portfolio_date:
                    break
        if portfolio_date:
            break

    if not portfolio_date:
        portfolio_date = default_date

    header_row = None
    col_map = {}
    for r in range(1, 15):
        row_vals = [ws.cell(row=r, column=c).value for c in range(1, 15)]
        str_vals = [str(v).strip().lower() if v is not None else "" for v in row_vals]
        if any("instrument" in s for s in str_vals) and any("isin" in s for s in str_vals):
            header_row = r
            for idx, s in enumerate(str_vals, start=1):
                if ("name of" in s or "instrument" in s) and "yield" not in s:
                    if "name" not in col_map:
                        col_map["name"] = idx
                elif "isin" in s:
                    col_map["isin"] = idx
                elif ("industry" in s or "rating" in s) and "industry" not in col_map:
                    col_map["industry"] = idx
                elif "quantity" in s or s == "qty":
                    col_map["qty"] = idx
            break

    if not header_row or "name" not in col_map or "isin" not in col_map or "qty" not in col_map:
        return []

    records = []
    for r in range(header_row + 1, ws.max_row + 1):
        name_val = ws.cell(row=r, column=col_map["name"]).value
        isin_val = ws.cell(row=r, column=col_map["isin"]).value
        qty_val = ws.cell(row=r, column=col_map["qty"]).value
        ind_val = ws.cell(row=r, column=col_map["industry"]).value if "industry" in col_map else None

        name_str = str(name_val).strip() if name_val is not None else ""
        if any(term in name_str.lower() for term in STOP_TERMS):
            break

        isin_str = str(isin_val).strip() if isin_val is not None else ""
        if not (isin_str.startswith("INE") and len(isin_str) == 12):
            continue

        if not name_str:
            continue

        try:
            qty = float(qty_val)
            if qty <= 0:
                continue
        except (ValueError, TypeError):
            continue

        records.append({
            "AMC": AMC_NAME,
            "Fund_Name": fund_name,
            "Portfolio_Date": portfolio_date,
            "Security_Name": name_str,
            "ISIN": isin_str,
            "Industry_Rating": str(ind_val).strip() if ind_val is not None else "",
            "Quantity": qty,
        })

    return records

def process_xlrd_sheet(ws, fund_code, fund_name, default_date):
    """
    Processes an xlrd sheet for a single fund (for legacy BIFF8 .xls files).
    """
    portfolio_date = None
    for r in range(min(8, ws.nrows)):
        for c in range(min(10, ws.ncols)):
            val = ws.cell_value(r, c)
            if val and isinstance(val, str) and "statement as on" in val.lower():
                portfolio_date = extract_date_from_text(val)
                if portfolio_date:
                    break
        if portfolio_date:
            break

    if not portfolio_date:
        portfolio_date = default_date

    header_row = None
    col_map = {}
    for r in range(min(15, ws.nrows)):
        row_vals = ws.row_values(r)
        str_vals = [str(v).strip().lower() if v is not None else "" for v in row_vals]
        if any("instrument" in s for s in str_vals) and any("isin" in s for s in str_vals):
            header_row = r
            for idx, s in enumerate(str_vals):
                if ("name of" in s or "instrument" in s) and "yield" not in s:
                    if "name" not in col_map:
                        col_map["name"] = idx
                elif "isin" in s:
                    col_map["isin"] = idx
                elif ("industry" in s or "rating" in s) and "industry" not in col_map:
                    col_map["industry"] = idx
                elif "quantity" in s or s == "qty":
                    col_map["qty"] = idx
            break

    if header_row is None or "name" not in col_map or "isin" not in col_map or "qty" not in col_map:
        return []

    records = []
    for r in range(header_row + 1, ws.nrows):
        name_val = ws.cell_value(r, col_map["name"])
        isin_val = ws.cell_value(r, col_map["isin"])
        qty_val = ws.cell_value(r, col_map["qty"])
        ind_val = ws.cell_value(r, col_map["industry"]) if "industry" in col_map else None

        name_str = str(name_val).strip() if name_val is not None else ""
        if any(term in name_str.lower() for term in STOP_TERMS):
            break

        isin_str = str(isin_val).strip() if isin_val is not None else ""
        if not (isin_str.startswith("INE") and len(isin_str) == 12):
            continue

        if not name_str:
            continue

        try:
            qty = float(qty_val)
            if qty <= 0:
                continue
        except (ValueError, TypeError):
            continue

        records.append({
            "AMC": AMC_NAME,
            "Fund_Name": fund_name,
            "Portfolio_Date": portfolio_date,
            "Security_Name": name_str,
            "ISIN": isin_str,
            "Industry_Rating": str(ind_val).strip() if ind_val is not None else "",
            "Quantity": qty,
        })

    return records

def clean_all_iti_files():
    print("=" * 80)
    print("ITI MF - Portfolio Cleaning Pipeline")
    print("=" * 80)

    files = sorted(RAW_FOLDER.rglob("*.*"))
    data_files = [f for f in files if f.suffix.lower() in [".xlsx", ".xls"]]
    print(f"Found {len(data_files)} portfolio files in {RAW_FOLDER}")

    all_records = []

    for idx, f in enumerate(data_files, start=1):
        rel = f.relative_to(RAW_FOLDER)
        year_str, month_str = rel.parts[0], rel.parts[1]
        last_day = 31 if int(month_str) in [1, 3, 5, 7, 8, 10, 12] else (30 if int(month_str) in [4, 6, 9, 11] else (29 if int(year_str) == 2024 else 28))
        default_date = f"{year_str}-{month_str}-{last_day:02d}"

        content = f.read_bytes()
        file_recs = 0

        # Detect format via magic bytes
        if content.startswith(b"PK"):
            wb = openpyxl.load_workbook(io.BytesIO(content), data_only=True)
            sheet_map = {s.strip().upper(): s for s in wb.sheetnames}
            for code, fund_name in TARGET_FUNDS.items():
                if code in sheet_map:
                    ws = wb[sheet_map[code]]
                    recs = process_openpyxl_sheet(ws, code, fund_name, default_date)
                    file_recs += len(recs)
                    all_records.extend(recs)
        elif content.startswith(b"\xd0\xcf"):
            wb = xlrd.open_workbook(file_contents=content)
            sheet_map = {s.strip().upper(): s for s in wb.sheet_names()}
            for code, fund_name in TARGET_FUNDS.items():
                if code in sheet_map:
                    ws = wb.sheet_by_name(sheet_map[code])
                    recs = process_xlrd_sheet(ws, code, fund_name, default_date)
                    file_recs += len(recs)
                    all_records.extend(recs)
        else:
            print(f"  [WARN] Unknown format for {f.name}")

        print(f"[{idx:2d}/{len(data_files)}] {year_str}-{month_str}: {f.name} -> {file_recs} records")

    if not all_records:
        print("❌ No records extracted!")
        return

    df = pd.DataFrame(all_records)
    df["Month"] = pd.to_datetime(df["Portfolio_Date"]).dt.strftime("%b-%Y")
    df = df[STANDARD_COLUMNS]

    df.sort_values(
        by=["Fund_Name", "Portfolio_Date", "Security_Name"],
        ascending=True,
        inplace=True,
    )
    df.reset_index(drop=True, inplace=True)

    df.to_excel(OUTPUT_FILE, index=False)

    print("\n" + "=" * 80)
    print("Cleaned Data Summary")
    print("=" * 80)
    print(f"Total Records : {len(df)}")
    print(f"Unique AMCs   : {df['AMC'].nunique()} ({list(df['AMC'].unique())})")
    print(f"Unique Funds  : {df['Fund_Name'].nunique()}")
    for fname, count in df["Fund_Name"].value_counts().items():
        print(f"  - {fname:<42}: {count:5d} rows")
    print(f"Unique ISINs  : {df['ISIN'].nunique()}")
    print(f"Date Range    : {df['Portfolio_Date'].min()} to {df['Portfolio_Date'].max()} ({df['Portfolio_Date'].nunique()} months)")
    print(f"Output File   : {OUTPUT_FILE}")
    print("=" * 80)

if __name__ == "__main__":
    clean_all_iti_files()
