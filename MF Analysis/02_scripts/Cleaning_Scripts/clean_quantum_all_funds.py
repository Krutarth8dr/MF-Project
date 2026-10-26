"""
Clean all downloaded monthly portfolio files for Quantum Mutual Fund (Quantum MF).

Process:
    - Reads all files in 01_raw_files/Quantum_MF/<YYYY>/<MM>/
    - Parses target sheets for the 6 target funds:
        1. QLTEVF   -> Quantum Value Fund
        2. QTSF     -> Quantum ELSS Tax Saver Fund
        3. QSCAPF   -> Quantum Small Cap Fund
        4. QMULTI   -> Quantum Multi Asset Allocation Fund
        5. QETHICAL -> Quantum Ethical Fund
        6. QFLEXI   -> Quantum Flexi Cap Fund
    - Dynamic header detection for Name, ISIN, Industry/Rating, and Quantity
    - Stops immediately upon encountering 'Sub Total', 'Total of all Equity', '(b)Unlisted', '(b) Unlisted', or 'Total'
    - Filters for valid listed Indian equities (ISIN starts with INE, length 12, quantity > 0)
    - Generates 03_clean_data/Quantum_MF/Quantum_MF_All_Funds_Cleaned.xlsx
"""

import sys
import os
import calendar
import warnings
from datetime import datetime
from pathlib import Path
import openpyxl
import pandas as pd

warnings.filterwarnings("ignore", category=UserWarning, module="openpyxl")

# Force UTF-8 output
try:
    sys.stdout.reconfigure(encoding="utf-8")
    sys.stderr.reconfigure(encoding="utf-8")
except (AttributeError, ValueError):
    pass

PROJECT_ROOT = Path(__file__).resolve().parents[2]
RAW_DIR = PROJECT_ROOT / "01_raw_files" / "Quantum_MF"
CLEAN_DIR = PROJECT_ROOT / "03_clean_data" / "Quantum_MF"
CLEAN_DIR.mkdir(parents=True, exist_ok=True)

OUTPUT_FILE = CLEAN_DIR / "Quantum_MF_All_Funds_Cleaned.xlsx"

CODE_TO_FUND = {
    "QLTEVF": "Quantum Value Fund",
    "QTSF": "Quantum ELSS Tax Saver Fund",
    "QSCAPF": "Quantum Small Cap Fund",
    "QMULTI": "Quantum Multi Asset Allocation Fund",
    "QETHICAL": "Quantum Ethical Fund",
    "QFLEXI": "Quantum Flexi Cap Fund",
}

STOP_TERMS = [
    "sub total",
    "total of all equity",
    "(b) unlisted",
    "(b)unlisted",
    "total",
    "subtotal"
]

STANDARD_COLUMNS = [
    "AMC",
    "Fund_Name",
    "Portfolio_Date",
    "Month",
    "Security_Name",
    "ISIN",
    "Industry_Rating",
    "Quantity",
]


def clean_sheet(ws, fund_name, portfolio_date, month_str):
    """Parses a single sheet for a Quantum fund."""
    holdings = []

    header_row = None
    col_map = {}
    for r_idx in range(1, 20):
        row_vals = [str(ws.cell(r_idx, c).value or "").strip() for c in range(1, min(15, ws.max_column + 1))]
        lower = [v.lower() for v in row_vals]
        if any("isin" in v for v in lower) and any("instrument" in v or "name" in v for v in lower):
            header_row = r_idx
            for c_idx, v in enumerate(lower, 1):
                if "isin" in v and "isin" not in col_map:
                    col_map["isin"] = c_idx
                elif ("instrument" in v or "name" in v) and "yield" not in v and "name" not in col_map:
                    col_map["name"] = c_idx
                elif ("rating" in v or "industry" in v) and "industry" not in col_map:
                    col_map["industry"] = c_idx
                elif ("quantity" in v or "qty" in v) and "qty" not in col_map:
                    col_map["qty"] = c_idx
            break

    if not header_row or "isin" not in col_map or "qty" not in col_map:
        return []

    for r_idx in range(header_row + 1, ws.max_row + 1):
        name_val = str(ws.cell(r_idx, col_map.get("name", 1)).value or "").strip()
        isin_val = str(ws.cell(r_idx, col_map["isin"]).value or "").strip()

        combined = f"{name_val.lower()} {isin_val.lower()}"
        if any(term in combined for term in STOP_TERMS):
            break

        if not (isin_val.startswith("INE") and len(isin_val) == 12):
            continue

        raw_qty = ws.cell(r_idx, col_map["qty"]).value
        try:
            qty = float(raw_qty)
            if qty <= 0:
                continue
        except (ValueError, TypeError):
            continue

        industry_val = ""
        if "industry" in col_map:
            industry_val = str(ws.cell(r_idx, col_map["industry"]).value or "").strip()

        holdings.append({
            "AMC": "Quantum MF",
            "Fund_Name": fund_name,
            "Portfolio_Date": portfolio_date,
            "Month": month_str,
            "Security_Name": name_val,
            "ISIN": isin_val,
            "Industry_Rating": industry_val,
            "Quantity": qty,
        })

    return holdings


def process_all_files():
    """Iterate over all raw combined workbooks in Quantum_MF, clean and save to single Excel file."""
    print("=" * 80)
    print("Starting Quantum MF Portfolio Cleaner")
    print(f"Reading from: {RAW_DIR}")
    print(f"Output target: {OUTPUT_FILE}")
    print("=" * 80)

    all_records = []
    file_count = 0

    for year_dir in sorted(RAW_DIR.iterdir()):
        if not year_dir.is_dir() or not year_dir.name.isdigit():
            continue
        year = int(year_dir.name)

        for month_dir in sorted(year_dir.iterdir()):
            if not month_dir.is_dir() or not month_dir.name.isdigit():
                continue
            month = int(month_dir.name)

            last_day = calendar.monthrange(year, month)[1]
            portfolio_date = datetime(year, month, last_day)
            month_str = f"{year}-{month:02d}-01"

            for file_path in sorted(month_dir.glob("*.xlsx")):
                file_count += 1
                try:
                    wb = openpyxl.load_workbook(file_path, data_only=True)
                    month_records = 0

                    for code, fund_name in CODE_TO_FUND.items():
                        matched_sheet = None
                        for s in wb.sheetnames:
                            if s.strip().upper() == code.upper():
                                matched_sheet = s
                                break
                        if not matched_sheet:
                            continue

                        ws = wb[matched_sheet]
                        holdings = clean_sheet(ws, fund_name, portfolio_date, month_str)
                        all_records.extend(holdings)
                        month_records += len(holdings)

                    print(f"Processed {year}-{month:02d} | {month_records:4d} holdings across active target funds | {file_path.name}")
                except Exception as e:
                    print(f"ERROR processing {file_path.name}: {e}")

    df = pd.DataFrame(all_records, columns=STANDARD_COLUMNS)

    print("\n" + "=" * 80)
    print("CLEANING COMPLETED - SUMMARY")
    print("=" * 80)
    print(f"Total Combined Workbooks Processed: {file_count}")
    print(f"Total Rows Extracted: {len(df):,}")
    print(f"Unique AMCs: {df['AMC'].nunique()} -> {list(df['AMC'].unique())}")
    print(f"Unique Funds ({df['Fund_Name'].nunique()}):")
    for f in df["Fund_Name"].unique():
        f_df = df[df["Fund_Name"] == f]
        print(f"  - {f:36}: {len(f_df):5,d} rows, {f_df['Month'].nunique():2d} months ({f_df['Month'].min()} to {f_df['Month'].max()})")
    print(f"Unique ISINs: {df['ISIN'].nunique():,}")
    print(f"Date Range: {df['Month'].min()} to {df['Month'].max()}")
    print("=" * 80)

    # Save to Excel
    print(f"Writing clean dataset to {OUTPUT_FILE}...")
    df.to_excel(OUTPUT_FILE, index=False, engine="openpyxl")
    file_size_kb = OUTPUT_FILE.stat().st_size / 1024
    print(f"Successfully saved {OUTPUT_FILE.name} ({file_size_kb:.1f} KB)")


if __name__ == "__main__":
    process_all_files()
